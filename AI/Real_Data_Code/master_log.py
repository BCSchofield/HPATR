"""master_log.xlsx: its layout, its migrations, one run's measured numbers, and the update of existing rows.

Shared by GUI_Clean (a row per capture, written at the end of a run and again when its AI chain finishes) and the
analysis app (after an analysis it fills the MEASURED columns of the rows of the runs it analysed: Ben, 2026-10-17,
"it autofills after CCD analysis"). Columns are always found by their header NAME, never by letter.

Layout since sizer 3.1.0 (plan 2026-10-17, Phase 4; Ben chose the columns): 'Droplets/s' stays (the throughput being
optimised) and becomes the 30-60 um band by rates.py; immediately right of it the concentration in five size bands
(droplets per cm^3 in the measurement volume, adding up to the total), the in-band liquid flow, the total
concentration, the atomised fraction inside the volume and the sizer version. The 'Sizer 3.0.0 ...' columns appended
on 2026-10-16 go (their numbers are in the backup of that day and the CCD comparison workbook).

numpy only at import; openpyxl is imported where a workbook is opened (the lab PC has it: GUI_Clean uses it)."""
from __future__ import annotations

import csv
import json
import math
import os
import re
import shutil
import subprocess
import sys
import time
from pathlib import Path

HERE = Path(__file__).resolve().parent
if str(HERE) not in sys.path:
    sys.path.insert(0, str(HERE))
import rates  # noqa: E402

RATE_BAND_UM = rates.DEFAULT_BAND_UM
RATE_HEADER = f"Droplets/s ({RATE_BAND_UM[0]:g}-{RATE_BAND_UM[1]:g} um)"
OLD_RATE_HEADER = "Droplets/s (20-60 um)"           # 2026-10-13 .. 2026-10-17: the 2.2.0 formula, 20-60 um
BANDS = ((None, 30.0, "<30 um (/cm3)"), (30.0, 40.0, "30-40 um (/cm3)"), (40.0, 60.0, "40-60 um (/cm3)"),
         (60.0, 100.0, "60-100 um (/cm3)"), (100.0, None, "100+ um (/cm3)"))
NEW_COLUMNS = tuple(b[2] for b in BANDS) + ("In-band flow (mL/min)", "Concentration (/cm3)",
                                            "Atomised in volume (%)", "Sizer")
MASTER_HEADERS = ('Timestamp', 'Orifice', 'Orifice Length (mm)', 'Bubbler Height (mm)', 'Bubbler RPM',
                  'Flow Range (sccm)', 'GLR', 'Pressure Range (barA)', 'Speed (steps/s)', 'Motor Travel (mm)',
                  'D32 (um)', 'Mean Diameter (um)', 'Atomised (%)', RATE_HEADER) + NEW_COLUMNS + (
                  'Avg Lamella Thickness', 'Notes', 'Cone Image', 'Shadowgraph', 'Pressure Graph', 'Mass Flow Graph')
MASTER_COL_WIDTHS = (20, 8, 10, 16, 12, 18, 10, 18, 14, 12, 11, 12, 12, 14,
                     11, 11, 11, 11, 11, 12, 13, 12, 8,
                     12, 35, 36.5, 43, 56, 56)
MEASURED = ('D32 (um)', 'Mean Diameter (um)', 'Atomised (%)', RATE_HEADER) + NEW_COLUMNS
RATE_NOTE = (f"Since sizer 3.1.0 (2026-10-17): in-band droplets produced per second, {RATE_BAND_UM[0]:g}-"
             f"{RATE_BAND_UM[1]:g} um, by rates.py. Rows with no Sizer were measured before: 20-60 um by the 2.2.0 "
             f"mass-balance formula (the backups hold them as written).")
NO_SPRAY_ATOM_PCT = 2.0         # as the app (stats.NO_SPRAY_ATOM_PCT): a dribble's size is not reported ('/')


# ---- migrations (the older ones were GUI_Clean's; module-level since 2026-10-13 so a back-fill uses them) ---------

def _column_index(letters: str) -> int:
    n = 0
    for ch in letters:
        n = n * 26 + (ord(ch) - 64)
    return n


def _column_letters(n: int) -> str:
    s = ""
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def shift_images(ws, from_col0: int, by: int) -> None:
    """Move every image anchored at or right of 0-indexed column `from_col0` by `by` columns (insert_cols /
    delete_cols move cell values, never images)."""
    for img in ws._images:
        anc = img.anchor
        if isinstance(anc, str):
            m = re.match(r'^([A-Z]+)(\d+)$', anc)
            if m and _column_index(m.group(1)) - 1 >= from_col0:
                img.anchor = f'{_column_letters(_column_index(m.group(1)) + by)}{m.group(2)}'
        elif hasattr(anc, '_from'):
            if anc._from.col >= from_col0:
                anc._from.col += by
            if hasattr(anc, 'to') and anc.to and anc.to.col >= from_col0:
                anc.to.col += by


def migrate_master_mean_diameter(ws) -> bool:
    """Insert 'Mean Diameter (um)' between 'D32 (um)' and 'Atomised (%)' (17-column layout -> 18). Cannot fire
    twice: once in, K reads 'Mean Diameter (um)'."""
    if not (ws.cell(row=1, column=10).value == 'D32 (um)' and ws.cell(row=1, column=11).value == 'Atomised (%)'):
        return False
    ws.insert_cols(11)
    ws.cell(row=1, column=11).value = 'Mean Diameter (um)'
    shift_images(ws, 10, 1)
    return True


def migrate_master_orifice_length(ws) -> bool:
    """Insert 'Orifice Length (mm)' after 'Orifice' (18-column layout -> 19). Runs after the mean-diameter one."""
    if not (ws.cell(row=1, column=2).value == 'Orifice' and ws.cell(row=1, column=3).value == 'Bubbler Height (mm)'):
        return False
    ws.insert_cols(3)
    ws.cell(row=1, column=3).value = 'Orifice Length (mm)'
    shift_images(ws, 2, 1)
    return True


def migrate_master_droplet_rate(ws) -> bool:
    """Insert the droplets/s column after 'Atomised (%)' (19-column layout -> 20). Runs after the orifice-length one.
    It is the 2026-10-13 column, so it arrives under its old name; migrate_master_sizer31 renames it."""
    if not (ws.cell(row=1, column=13).value == 'Atomised (%)'
            and ws.cell(row=1, column=14).value == 'Avg Lamella Thickness'):
        return False
    ws.insert_cols(14)
    ws.cell(row=1, column=14).value = OLD_RATE_HEADER
    shift_images(ws, 13, 1)
    return True


def migrate_master_sizer31(ws) -> bool:
    """20-column layout (N = 'Droplets/s (20-60 um)', O = 'Avg Lamella Thickness') -> 29: drop the 'Sizer 3.0.0 ...'
    columns appended on 2026-10-16, rename N to the 30-60 um band (rows measured before keep their values; their
    Sizer cell stays blank and a note on the header says what they are), and insert the nine new columns after N,
    every image moving right with its column. Cannot fire twice."""
    if not (ws.cell(row=1, column=14).value == OLD_RATE_HEADER
            and ws.cell(row=1, column=15).value == 'Avg Lamella Thickness'):
        return False
    for c in range(ws.max_column, 20, -1):            # right to left, so the indices still hold
        if str(ws.cell(row=1, column=c).value or "").startswith("Sizer 3.0.0"):
            ws.delete_cols(c)
    ws.insert_cols(15, len(NEW_COLUMNS))
    shift_images(ws, 14, len(NEW_COLUMNS))
    ws.cell(row=1, column=14).value = RATE_HEADER
    for i, h in enumerate(NEW_COLUMNS):
        ws.cell(row=1, column=15 + i).value = h
    try:
        from openpyxl.comments import Comment
        ws.cell(row=1, column=14).comment = Comment(RATE_NOTE, "HPATR")
    except Exception:                                 # noqa: BLE001 -- a note is a courtesy, never a failure
        pass
    return True


def migrate_measurement_columns(ws) -> bool:
    """Every migration from the 17-column layout on, in order. Returns True if any changed the sheet. (The older
    ones, before the 17-column layout, live in GUI_Clean's save: no log that old is left.)"""
    changed = False
    for m in (migrate_master_mean_diameter, migrate_master_orifice_length, migrate_master_droplet_rate,
              migrate_master_sizer31):
        changed = m(ws) or changed
    return changed


def stamp_headers(ws) -> None:
    """The canonical names on row 1, each new one styled like the first header (insert_cols gives a column no style)."""
    from copy import copy
    a1 = ws.cell(row=1, column=1)
    for c, name in enumerate(MASTER_HEADERS, start=1):
        cell = ws.cell(row=1, column=c)
        cell.value = name
        if not cell.has_style or cell.fill != a1.fill:
            cell.font, cell.fill, cell.border = copy(a1.font), copy(a1.fill), copy(a1.border)
            cell.alignment = copy(a1.alignment)


# ---- one run's measured numbers ----------------------------------------------------------------------------------

def run_summary_fields(run_dir: Path) -> dict:
    """{'Timestamp': ..., 'Speed (steps/s)': ...} from the run's run_summary.xlsx (GUI_Clean writes both): the
    Timestamp is the row's key in master_log, the speed the syringe flow."""
    from openpyxl import load_workbook
    wb = load_workbook(Path(run_dir) / "run_summary.xlsx", read_only=True)
    try:
        ws = wb["Metadata"] if "Metadata" in wb.sheetnames else wb.worksheets[0]
        return {str(r[0]).strip(): r[1] for r in ws.iter_rows(values_only=True) if r and r[0] is not None}
    finally:
        wb.close()


def band_concentrations(meas_dir: Path) -> dict | None:
    """Droplets per cm^3 in the measurement volume in each of BANDS (sizer 3.x): border weights, crowded and other
    unmeasured droplets added back in proportion, '<30' = everything below d_min (counted, not sized). They add up to
    the concentration in summary.json. None before 3.x."""
    meas_dir = Path(meas_dir)
    summ = json.loads((meas_dir / "summary.json").read_text(encoding="utf-8"))
    c = summ.get("concentration") or {}
    if not c.get("volume_mm3") or not c.get("frames"):
        return None
    k = float(c.get("crowded_factor") or 1.0) / (float(c["frames"]) * float(c["volume_mm3"])) * 1000.0
    out = {name: 0.0 for _lo, _hi, name in BANDS}
    with open(meas_dir / "droplet_sizes.csv", newline="", encoding="utf-8") as fh:
        for r in csv.DictReader(fh):
            if r.get("in_volume") != "1" or not r.get("weight"):
                continue
            w = float(r["weight"])
            if r.get("in_focus") != "1":
                out[BANDS[0][2]] += w                  # in the volume, below d_min: counted, not sized
                continue
            d = float(r["diameter_um"])
            for lo, hi, name in BANDS[1:]:
                if (lo is None or d >= lo) and (hi is None or d < hi):
                    out[name] += w
                    break
            else:
                out[BANDS[0][2]] += w                  # sized below 30 um cannot happen at d_min 30; never lost
    return {name: round(v * k, 3) for name, v in out.items()}


def measured_values(meas_dir: Path, liquid_dir: Path, steps_s=None) -> dict:
    """header -> value for the MEASURED columns of one run, from its own files. '' where a number does not exist for
    its sizer (3.1.0's rate before 3.1.0, concentrations before 3.x); '/' for a no-spray run's sizes, as the app."""
    meas_dir, liquid_dir = Path(meas_dir), Path(liquid_dir)
    summ = json.loads((meas_dir / "summary.json").read_text(encoding="utf-8"))
    cs = json.loads((liquid_dir / "classical_summary.json").read_text(encoding="utf-8"))
    prov = summ.get("provenance") or {}
    version = str(prov.get("sizer_version") or "")
    atom = cs.get("atomised_pct_pooled")
    dribble = atom is not None and float(atom) < NO_SPRAY_ATOM_PCT
    vals = {h: '' for h in MEASURED}
    vals.update({'D32 (um)': '/' if dribble else summ.get('d32_in_focus_um', ''),
                 'Mean Diameter (um)': '/' if dribble else summ.get('droplet_d_mean_um', ''),
                 'Atomised (%)': atom if atom is not None else '', 'Sizer': version})
    if version.startswith("3."):
        conc = (summ.get("concentration") or {}).get("in_volume_per_mm3")
        vals['Concentration (/cm3)'] = round(conc * 1000.0, 3) if conc is not None else ''
        vals['Atomised in volume (%)'] = cs.get('atomised_pct_in_volume_pooled') or ''
        vals.update(band_concentrations(meas_dir) or {})
    if version == "3.1.0" and steps_s not in (None, ''):
        t = rates.from_run_files(meas_dir, liquid_dir, steps_s, band=RATE_BAND_UM)
        if t and t.get("band_rate_per_s") is not None:
            vals[RATE_HEADER] = int(round(t["band_rate_per_s"]))
            vals['In-band flow (mL/min)'] = t["band_flow_ml_min"]
    return {h: ('' if v is None else v) for h, v in vals.items()}


# ---- the log on disk -----------------------------------------------------------------------------------------------

class MasterLogBusy(RuntimeError):
    pass


def master_log_path() -> Path | None:
    """HPATR_MASTER_LOG: 'off' (the test suite) -> None; a path -> that file; unset -> the LaCie's
    Experiments/Logs/master_log.xlsx (None without the LaCie)."""
    env = os.environ.get("HPATR_MASTER_LOG")
    if env:
        return None if env.lower() == "off" else Path(env)
    try:
        sys.path.insert(0, str(HERE.parents[1] / "src"))
        from config_loader import find_lacie_drive
        drive = find_lacie_drive()
    except Exception:                                 # noqa: BLE001
        return None
    return Path(drive) / "Experiments" / "Logs" / "master_log.xlsx" if drive else None


LOCK_FRESH_S = 3 * 24 * 3600     # macOS: an owner file older than this is a leftover (one from 2026-09-09 sits there)


def excel_has_it_open(path: Path) -> bool:
    """Windows: Excel locks the file (opening it for writing fails). macOS: Excel does not lock the workbook but
    writes an owner file '~$<name>' beside it while it is open: open = Excel running and that file fresh."""
    path = Path(path)
    if os.name == "nt":
        try:
            with open(path, "r+b"):
                return False
        except PermissionError:
            return True
        except OSError:
            return False
    owner = path.with_name("~$" + path.name)
    try:
        fresh = owner.exists() and time.time() - owner.stat().st_mtime < LOCK_FRESH_S
        return fresh and subprocess.run(["pgrep", "-x", "Microsoft Excel"], capture_output=True).returncode == 0
    except OSError:
        return False


def update_runs(run_dirs, score_thresh: float = 0.30, path: Path | None = None, log=print) -> dict:
    """Fill the MEASURED columns of the existing master_log rows of `run_dirs` (matched on run_summary's Timestamp)
    from each run's own results; nothing else in a row is touched and no row is added (a run with no row is
    reported). The log is migrated to the current layout first. A backup is kept once a day
    (Logs/Backups/master_log BACKUP <date> before analysis update.xlsx), the save goes to a temporary file in the
    same folder and replaces the log only when complete. Refuses while Excel has it open. Returns what it did."""
    path = Path(path) if path is not None else master_log_path()
    report = {"path": str(path) if path else None, "updated": [], "no_row": [], "skipped": [], "backup": None}
    if path is None or not path.exists():
        report["skipped"] = [f"no master_log ({path or 'off'})"]
        return report
    wanted = {}
    for rd in run_dirs:
        rd = Path(rd)
        try:
            f = run_summary_fields(rd)
            ts = str(f.get("Timestamp") or "").strip()
            an = rd / "shadowgraph" / "analysis"
            vals = measured_values(an / f"droplets_{score_thresh:.2f}", an / f"liquid_{score_thresh:.2f}",
                                   f.get("Speed (steps/s)"))
        except Exception as e:                        # noqa: BLE001 -- one run's gap never stops the others
            report["skipped"].append(f"{rd.name}: {type(e).__name__}: {e}")
            continue
        if ts:
            wanted[ts] = (rd.name, vals)
    if not wanted:
        return report
    from openpyxl import load_workbook
    ro = load_workbook(path, read_only=True)            # cheap look first: does any run have a row at all?
    try:
        stamps = {str(r[0]).strip() for r in ro.active.iter_rows(min_row=2, max_col=1, values_only=True) if r[0]}
    finally:
        ro.close()
    report["no_row"] = sorted(n for ts, (n, _v) in wanted.items() if ts not in stamps)
    if not any(ts in stamps for ts in wanted):
        return report
    if excel_has_it_open(path):
        raise MasterLogBusy("master_log is open in Excel: close it and analyse again (the results are saved in "
                            "each run's folder; only the log was not updated)")
    mtime = path.stat().st_mtime
    wb = load_workbook(path)
    ws = wb.active
    n_img = len(ws._images)
    migrate_measurement_columns(ws)
    stamp_headers(ws)
    col = {ws.cell(row=1, column=c).value: c for c in range(1, ws.max_column + 1)}
    missing = [h for h in MEASURED if h not in col]
    if missing:
        raise RuntimeError(f"master_log has no column {missing}: an older layout this cannot migrate")
    for r in range(2, ws.max_row + 1):
        ts = str(ws.cell(row=r, column=1).value or "").strip()
        if ts in wanted:
            name, vals = wanted[ts]
            for h in MEASURED:
                ws.cell(row=r, column=col[h]).value = vals.get(h, '')
            report["updated"].append(name)
    if len(ws._images) != n_img:
        raise RuntimeError("master_log images changed in memory; nothing saved")
    backups = path.parent / "Backups"
    backups.mkdir(exist_ok=True)
    bk = backups / f"master_log BACKUP {time.strftime('%Y-%m-%d')} before analysis update.xlsx"
    if not bk.exists():
        shutil.copy2(path, bk)
        report["backup"] = str(bk)
    tmp = path.with_name(f".~{path.stem}.saving.xlsx")
    wb.save(tmp)
    if path.stat().st_mtime != mtime:
        tmp.unlink()
        raise MasterLogBusy("master_log changed while it was being updated (GUI_Clean saving a run?): try again")
    os.replace(tmp, path)
    log(f"master_log: updated {len(report['updated'])} row(s)"
        + (f"; no row for {len(report['no_row'])} run(s)" if report["no_row"] else "")
        + (f"; backup {bk.name}" if report["backup"] else ""))
    return report
