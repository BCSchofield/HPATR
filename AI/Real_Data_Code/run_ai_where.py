#!/usr/bin/env python3
"""
run_ai_where.py -- run the AI chain (process_capture: frames, background, inference, measurement) on the
runs whose Notes start with given text, e.g. the falling-bead captures of the size calibration
(2026-10-14; docs/BEAD_CALIBRATION.md):

    python run_ai_where.py "D:\\Calibration\\Imaging Work\\2026\\10\\15" --notes "BEAD fall"
    python run_ai_where.py <folder> --notes "BEAD fall" --dry-run        (list what would run)

The Taguchi app's batch would also do it, but it takes whole day folders: the 43 static bead captures in
the same folder need no AI (bead_calibration.py measures them without it). Restartable: a run that already
holds a complete analysis at this stride is not inferred again (process_capture's reuse). Writes only into
each run's own shadowgraph/ folder, as the app's batch does; the .cine is only read. Needs the AI
environment (torch + detectron2): the home PC's Detectron env.

LOOSE .cine FILES (e.g. falling beads recorded in PCC rather than GUI_Clean, 2026-10-14): a .cine sitting
directly in a given folder, whose FILE NAME starts with the Notes text ("BEAD fall 56.28 p1.cine"), becomes
a run folder next to it, "<folder>/<file name>/", with a run_summary.xlsx whose Notes are the file name
(so bead_calibration.py finds it like any run). process_capture then copies the .cine into that run folder
(its contract: the .cine lives in the run), so each pour is stored twice; the original is never changed.
"""
from __future__ import annotations

import argparse
import sys
import time
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))


def read_notes(run: Path) -> str:
    import openpyxl
    x = run / "run_summary.xlsx"
    if not x.is_file():
        return ""
    wb = openpyxl.load_workbook(x, read_only=True, data_only=True)
    try:
        if "Metadata" not in wb.sheetnames:
            return ""
        for r in wb["Metadata"].iter_rows(min_row=2, values_only=True):
            if r and r[0] is not None and str(r[0]).strip() == "Notes":
                return str(r[1] or "").strip()
        return ""
    finally:
        wb.close()


def loose_cines(paths, notes_prefix: str) -> list[tuple[Path, Path]]:
    """[(cine, the run folder it will get)] for .cine files directly in `paths` whose name starts with
    notes_prefix (case ignored; underscores read as spaces)."""
    want = notes_prefix.strip().lower()
    out = []
    for p in paths:
        p = Path(p)
        if not p.is_dir():
            continue
        for c in sorted(p.glob("*.cine")):
            if c.stem.replace("_", " ").strip().lower().startswith(want):
                out.append((c, p / c.stem))
    return out


def make_run_folder(cine: Path, run: Path) -> None:
    """The run folder for a loose .cine: a run_summary.xlsx whose Notes are the file name (underscores as
    spaces). Nothing else; process_capture fills shadowgraph/."""
    import openpyxl
    run.mkdir(parents=True, exist_ok=True)
    x = run / "run_summary.xlsx"
    if x.exists():
        return
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Metadata"
    ws.append(["Field", "Value"])
    ws.append(["Notes", cine.stem.replace("_", " ").strip()])
    ws.append(["Source", f"loose .cine (e.g. recorded in PCC): {cine}"])
    wb.save(x)


def select_runs(paths, notes_prefix: str) -> list[tuple[Path, Path | None, str]]:
    """[(run folder, its .cine or None, its Notes)] for every run under `paths` whose Notes start with
    notes_prefix (case and surrounding spaces ignored), in folder order."""
    want = notes_prefix.strip().lower()
    out = []
    for p in paths:
        p = Path(p)
        found = [p / "run_summary.xlsx"] if (p / "run_summary.xlsx").is_file() else \
            sorted(q for q in p.rglob("run_summary.xlsx") if not any(s.startswith((".", "_")) for s in q.parts))
        for x in found:
            run = x.parent
            notes = read_notes(run)
            if notes.lower().startswith(want):
                cines = sorted((run / "shadowgraph" / "raw" / "CINE").glob("*.cine"))
                if not cines and (run.parent / f"{run.name}.cine").is_file():
                    cines = [run.parent / f"{run.name}.cine"]          # a loose .cine's own run folder
                out.append((run, cines[0] if cines else None, notes))
    return out


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("paths", nargs="+", type=Path, help="day folders (or run folders) to look in")
    ap.add_argument("--notes", required=True, help='Notes prefix, e.g. "BEAD fall"')
    ap.add_argument("--stride", type=int, default=None, help="keep every Nth frame (default: process_capture's)")
    ap.add_argument("--dry-run", action="store_true", help="list the runs; run nothing")
    a = ap.parse_args()
    loose = loose_cines(a.paths, a.notes)
    if loose:
        print(f"{len(loose)} loose .cine file(s) named '{a.notes} ...': each gets a run folder next to it")
        for cine, run in loose:
            print(f"  {cine.name}  ->  {run.name}/")
        if not a.dry_run:
            for cine, run in loose:
                make_run_folder(cine, run)
    runs = select_runs(a.paths, a.notes)
    known = {r for r, _c, _n in runs}
    for cine, run in loose:                     # (a dry run has no folders yet; list them as they will be)
        if run not in known:
            runs.append((run, cine, cine.stem.replace("_", " ").strip()))
    print(f"{len(runs)} run(s) with Notes starting '{a.notes}':")
    for run, cine, notes in runs:
        print(f"  {run.name}   {notes!r}" + ("" if cine else "   (NO .cine: skipped)"))
    if a.dry_run or not runs:
        return
    from process_capture import DEFAULT_STRIDE, process_capture
    stride = a.stride or DEFAULT_STRIDE
    failed = 0
    for i, (run, cine, notes) in enumerate(runs, 1):
        if cine is None:
            continue
        t0 = time.time()
        print(f"\n=== {i}/{len(runs)}  {run.name}  ({notes}) ===", flush=True)
        try:
            process_capture(cine, run, stride=stride, reuse=True, sharpness_rule=True)
            print(f"=== done in {time.time() - t0:.0f} s ===", flush=True)
        except Exception as e:                       # noqa: BLE001 -- reported; the others carry on
            failed += 1
            print(f"=== FAILED: {type(e).__name__}: {e} ===", flush=True)
    print(f"\n{len(runs) - failed} of {len(runs)} done" + (f", {failed} failed" if failed else ""))
    sys.exit(1 if failed else 0)


if __name__ == "__main__":
    main()
