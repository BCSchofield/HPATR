"""
Lamella analysis Excel generator.

After a batch run, combines lamella_thickness.csv with the experiment's
run_summary.xlsx to produce a combined Excel file with pressure + thickness
on a shared time axis.

Walks up from the TIFF folder to find run_summary.xlsx automatically:
    TIFFs/  →  raw/  →  shadowgraph/  →  <experiment_root>/run_summary.xlsx
"""
from __future__ import annotations
import csv
import os
from typing import Optional


def find_run_summary(tiff_dir: str) -> Optional[str]:
    """Walk up 3 directory levels from tiff_dir to find run_summary.xlsx."""
    path = os.path.abspath(tiff_dir)
    for _ in range(3):
        path = os.path.dirname(path)
        candidate = os.path.join(path, "run_summary.xlsx")
        if os.path.isfile(candidate):
            return candidate
    return None


def read_run_summary(summary_path: str) -> dict:
    """
    Parse run_summary.xlsx.

    Returns dict with:
        fps: float or None
        camera_windows: list of (cam_start, cam_end, run_index) sorted by run_index
            cam_start/cam_end are seconds-since-experiment-start floats
    """
    import openpyxl
    wb = openpyxl.load_workbook(summary_path, data_only=True)

    # Read FPS from Metadata sheet
    fps = None
    if "Metadata" in wb.sheetnames:
        ws = wb["Metadata"]
        for row in ws.iter_rows(values_only=True):
            if row and str(row[0]).strip().upper() == "FPS":
                try:
                    fps = float(row[1])
                except (TypeError, ValueError):
                    pass

    # Read camera windows from Pressure sheet
    camera_windows = []
    if "Pressure" in wb.sheetnames:
        ws = wb["Pressure"]
        headers = [str(c.value).strip() if c.value else "" for c in next(ws.iter_rows(min_row=1, max_row=1))]

        # Find all cam_start_N / cam_end_N column pairs
        run_indices = sorted({
            int(h.split("_")[-1])
            for h in headers
            if h.startswith("cam_start_") and h.split("_")[-1].isdigit()
        })

        for idx in run_indices:
            col_s = f"cam_start_{idx}"
            col_e = f"cam_end_{idx}"
            if col_s not in headers or col_e not in headers:
                continue
            ci_s = headers.index(col_s)
            ci_e = headers.index(col_e)
            cam_start = cam_end = None
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row[ci_s] is not None:
                    try:
                        cam_start = float(row[ci_s])
                    except (TypeError, ValueError):
                        pass
                if row[ci_e] is not None:
                    try:
                        cam_end = float(row[ci_e])
                    except (TypeError, ValueError):
                        pass
            if cam_start is not None:
                camera_windows.append((cam_start, cam_end, idx))

    return {"fps": fps, "camera_windows": camera_windows}


def generate_analysis_excel(
    tiff_dir: str,
    thickness_csv: str,
    fps: float,
    cam_start: float,
    cam_end: Optional[float],
    run_index: int,
    summary_path: str,
) -> str:
    """
    Generate combined Excel with pressure + lamella thickness on shared time axis.

    Returns path to the written Excel file.
    """
    import openpyxl
    import pandas as pd

    # Read thickness CSV
    frames = []
    with open(thickness_csv, newline="") as f:
        for row in csv.DictReader(f):
            idx = int(row["frame_index"])
            t = cam_start + idx / fps
            frames.append({
                "time_s":       round(t, 4),
                "frame_index":  idx,
                "filename":     row["filename"],
                "thickness_px": row["thickness_px"],
                "thickness_mm": row["thickness_mm"],
                "ok":           row["ok"],
            })
    thickness_df = pd.DataFrame(frames)

    # Read pressure data from run_summary
    wb_src = openpyxl.load_workbook(summary_path, data_only=True)
    ws_p = wb_src["Pressure"]
    headers = [c.value for c in next(ws_p.iter_rows(min_row=1, max_row=1))]
    rows = list(ws_p.iter_rows(min_row=2, values_only=True))
    press_df = pd.DataFrame(rows, columns=headers)[["timestamps", "pressures"]].copy()
    press_df.columns = ["time_s", "pressure_bar"]
    press_df = press_df.dropna(subset=["time_s", "pressure_bar"])
    press_df["time_s"] = press_df["time_s"].astype(float).round(4)
    press_df["pressure_bar"] = press_df["pressure_bar"].astype(float)

    # Filter pressure to camera window only
    press_window = press_df[press_df["time_s"] >= cam_start].copy()
    if cam_end is not None:
        press_window = press_window[press_window["time_s"] <= cam_end]

    # Combined sheet: merge on nearest time (pressure is sparse, thickness is dense)
    # Use pandas merge_asof to align pressure to nearest thickness timestamp
    thick_sorted = thickness_df.sort_values("time_s").reset_index(drop=True)
    press_sorted = press_window.sort_values("time_s").reset_index(drop=True)
    combined = pd.merge_asof(
        thick_sorted,
        press_sorted,
        on="time_s",
        direction="nearest",
        tolerance=1.0,
    )

    # Write Excel
    out_path = os.path.join(
        os.path.dirname(thickness_csv),
        f"lamella_analysis_run{run_index}.xlsx",
    )
    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        press_df.to_excel(writer, sheet_name="Pressure (full)", index=False)
        press_window.to_excel(writer, sheet_name="Pressure (camera window)", index=False)
        thickness_df.to_excel(writer, sheet_name="Lamella Thickness", index=False)
        combined.to_excel(writer, sheet_name="Combined", index=False)

    return out_path
