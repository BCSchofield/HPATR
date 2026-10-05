"""run_discovery.py -- find experiment run folders and describe each one.

No Qt in here: everything is plain data in, plain data out, so it is testable
without a display and callable from a worker thread.

A RUN FOLDER is named  <HHMMSS>_<sccm>sccm_<rpm>rpm_<sps>sps_or<mm>_bh<mm>  and
holds  shadowgraph/raw/CINE/<one>.cine  plus (usually) run_summary.xlsx. This
module READS only. It never creates, moves or deletes anything in a run folder.

Surveyed on the LaCie (2026-10-05), which shaped the rules below:
  * 37 folders match the name format: 27 on 2026/10/05 (a clean L9 x 3) and 10
    on 2026/10/01 (already measured, under the LEGACY measurement_/classical_
    folder names).
  * 61 older folders (e.g. 04/08/110302_N6_1.0BAR) use a different scheme. They
    are reported as skipped, with the reason -- never silently dropped.
  * run_summary.xlsx fields vary: 10 runs lack GLR / Fluid / densities, so every
    field is optional. FPS is 800 on some runs and 390 on others.
  * xlsx values are mostly STRINGS even when numeric ("300", "1.2mm"), ranges use
    an en-dash ("0-3086 sccm" with U+2013), and unrecorded values may be text.
"""
from __future__ import annotations

import os
import re
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any, Callable, Iterable

from . import pipeline_spec as spec

# The folder name is written by GUI_Clean._run_id(): HHMMSS_<flow>_<rpm>_<sps>_<or>_<bh>.
# A value that was never entered becomes an explicit token (norpm, nosps, noor,
# nobh) rather than being dropped, and the flow can be "?sccm" or "unknownsccm"
# when no flow reading was available. So the grammar is NOT just digits.
_N = r"\d+(?:\.\d+)?"
RUN_NAME_RE = re.compile(
    rf"^(?P<time>\d{{6}})_(?P<sccm>{_N}sccm|\?sccm|unknownsccm)_(?P<rpm>{_N}rpm|norpm)"
    rf"_(?P<sps>{_N}sps|nosps)_(?P<orifice>or{_N}|noor)_(?P<bh>bh{_N}|nobh)$"
)
NAME_FORMAT = "<HHMMSS>_<n>sccm_<n>rpm_<n>sps_or<x>_bh<n>"

# GUI_Clean._run_id's own docstring: "run_summary.xlsx remains authoritative; this
# [the folder name] is a human-scannable copy ... Do not parse it back as data --
# a renamed folder would silently disagree with the spreadsheet." So the workbook
# wins wherever it has a value; the name fills gaps (and supplies the flow
# SETPOINT, which the workbook records only as an achieved range). Disagreement
# is a warning, never silent.
WORKBOOK_FIELDS = (          # RunInfo attribute, workbook field, label
    ("rpm", "Bubbler RPM", "rpm"),
    ("sps", "Speed (steps/s)", "sps"),
    ("orifice_mm", "Orifice", "orifice"),
    ("bubbler_height", "Bubbler Height (mm)", "bubbler height"),
)


def parse_run_name(name: str) -> dict[str, float | None] | None:
    """Factor values spelled in a run folder name; a no-token or ?sccm gives None.
    Returns None if the name is not a run name at all."""
    m = RUN_NAME_RE.match(name)
    if not m:
        return None

    def num(token: str) -> float | None:
        found = re.search(_N, token)
        return float(found.group()) if found and not token.startswith(("?", "no", "unknown")) \
            else None

    return {"time": m["time"], "sccm": num(m["sccm"]), "rpm": num(m["rpm"]),
            "sps": num(m["sps"]), "orifice_mm": num(m["orifice"]),
            "bubbler_height": num(m["bh"])}


SUMMARY_XLSX = "run_summary.xlsx"
METADATA_SHEET = "Metadata"

# Folder names carry the gas-flow SETPOINT; the xlsx carries the ACHIEVED flow
# range. They differ by ~3% on every real run (3000 vs 3086/3110/3098), which is
# normal. Beyond this the folder name is probably wrong.
FLOW_DRIFT_WARN = 0.10

PRE_SIZER = "pre-2.0.0"      # no sizer_version in summary.json: model-mask-area diameters

# How deep to look when the user picks a parent folder: run <- day <- month <- year.
DEFAULT_SCAN_DEPTH = 3


# ============================================================================
# xlsx values
# ============================================================================

_UNRECORDED = frozenset({"", "not recorded", "not set", "n/a", "na", "none", "unknown",
                         "nan", "null", "-", "--"})
_DASHES = str.maketrans({"–": "-", "—": "-", "−": "-"})
_NUM = r"[-+]?(?:\d+\.?\d*|\.\d+)(?:[eE][-+]?\d+)?"
_UNIT = r"[A-Za-z%°/][A-Za-z0-9%°/^.]*"
_RANGE_RE = re.compile(rf"^\s*({_NUM})\s*-\s*({_NUM})\s*({_UNIT})?\s*$")
_NUMBER_RE = re.compile(rf"^\s*({_NUM})\s*({_UNIT})?\s*$")


@dataclass(frozen=True)
class Value:
    """One xlsx cell, interpreted. Never raises: anything unreadable becomes
    `missing` or plain `text` with no number."""
    raw: Any
    text: str = ""
    number: float | None = None
    range: tuple[float, float] | None = None
    unit: str = ""
    missing: str | None = None

    def pick(self, how: str = "max") -> float | None:
        """A single number: the value itself, or one end/the middle of a range
        ("max" for achieved flow, which is the plateau the run reached)."""
        if self.number is not None:
            return self.number
        if self.range is None:
            return None
        lo, hi = self.range
        return {"max": hi, "min": lo, "mean": (lo + hi) / 2}[how]


def coerce_value(raw: Any) -> Value:
    """Interpret a Metadata cell: "1.2mm" -> 1.2 mm, "0-3086 sccm" (en-dash)
    -> range, "NOT RECORDED" -> missing, 390 -> 390."""
    if raw is None:
        return Value(raw, missing="empty")
    if isinstance(raw, bool):
        return Value(raw, text=str(raw))
    if isinstance(raw, (int, float)):
        return Value(raw, text=str(raw), number=float(raw))
    text = str(raw).strip()
    if text.lower() in _UNRECORDED:
        return Value(raw, text=text, missing=text or "empty")
    t = text.translate(_DASHES)
    m = _RANGE_RE.match(t)
    if m:
        return Value(raw, text=text, range=(float(m.group(1)), float(m.group(2))),
                     unit=m.group(3) or "")
    m = _NUMBER_RE.match(t)
    if m:
        return Value(raw, text=text, number=float(m.group(1)), unit=m.group(2) or "")
    return Value(raw, text=text)


def read_metadata(xlsx: Path) -> dict[str, Value]:
    """The Metadata sheet as {field: Value}. Raises on an unreadable workbook or
    a missing sheet -- the caller turns that into an Issue."""
    import openpyxl

    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    try:
        if METADATA_SHEET not in wb.sheetnames:
            raise KeyError(f"no '{METADATA_SHEET}' sheet (has {wb.sheetnames})")
        out: dict[str, Value] = {}
        for row in wb[METADATA_SHEET].iter_rows(values_only=True):
            if not row or row[0] is None:
                continue
            key = str(row[0]).strip()
            if key and key != "Field":
                out[key] = coerce_value(row[1] if len(row) > 1 else None)
        return out
    finally:
        wb.close()


# ============================================================================
# per-run description
# ============================================================================


@dataclass(frozen=True)
class Issue:
    level: str          # "error" excludes the run from a batch; "warn" is shown only
    message: str


@dataclass
class AnalysisState:
    """What is already on disk for this run at one score threshold."""
    measured: bool = False            # droplets + liquid summaries both present
    droplets_only: bool = False       # droplets summary present but no classical stage
    legacy_names: bool = False        # found under measurement_/classical_
    meas_dir: Path | None = None
    clas_dir: Path | None = None
    sizer_version: str | None = None  # measure_run's provenance; None on old runs
    reusable: bool | None = None      # can_reuse(); None = not evaluated

    @property
    def label(self) -> str:
        if self.measured:
            v = f"v{self.sizer_version}" if self.sizer_version else PRE_SIZER
            return f"measured {v}" + (" \u00b7 legacy folders" if self.legacy_names else "")
        if self.droplets_only:
            return "droplets measured, no classical stage"
        if self.reusable:
            return "frames + inference on disk"
        return "new"


@dataclass
class RunInfo:
    path: Path
    name: str
    name_ok: bool
    time: str | None = None
    date: str | None = None                 # YYYY-MM-DD
    sccm: float | None = None               # SETPOINT, from the folder name
    rpm: float | None = None                # workbook value, else folder name
    sps: float | None = None
    orifice_mm: float | None = None
    bubbler_height: float | None = None
    sources: dict[str, str] = field(default_factory=dict)   # attribute -> "name" | "workbook"
    cine: Path | None = None
    n_cines: int = 0
    cine_bytes: int = 0
    has_xlsx: bool = False
    meta: dict[str, Value] = field(default_factory=dict)
    notes: str = ""
    fps: float | None = None
    flow_achieved_sccm: float | None = None
    analysis: AnalysisState = field(default_factory=AnalysisState)
    issues: list[Issue] = field(default_factory=list)

    @property
    def errors(self) -> list[Issue]:
        return [i for i in self.issues if i.level == "error"]

    @property
    def warnings(self) -> list[Issue]:
        return [i for i in self.issues if i.level == "warn"]

    @property
    def runnable(self) -> bool:
        """Can a batch process this run? Needs a good name and exactly one .cine."""
        return self.name_ok and self.n_cines == 1

    @property
    def analysable(self) -> bool:
        """Can the statistics use it without re-measuring?"""
        return self.name_ok and self.analysis.measured

    @property
    def usable(self) -> bool:
        return self.runnable or self.analysable

    @property
    def condition_key(self) -> tuple | None:
        """Runs with the same key are replicates of one another."""
        if not self.name_ok:
            return None
        return (self.sccm, self.rpm, self.sps, self.orifice_mm, self.bubbler_height)

    @property
    def label(self) -> str:
        if not self.name_ok:
            return self.name
        when = f"{self.date}  " if self.date else ""
        def g(v):
            return "?" if v is None else f"{v:g}"
        return (f"{when}{self.time}   {g(self.sccm)} sccm \u00b7 {g(self.rpm)} rpm \u00b7 "
                f"{g(self.sps)} sps")


def _date_from_path(path: Path) -> str | None:
    """.../YYYY/MM/DD/<run> -> YYYY-MM-DD, or None for a folder not in that tree."""
    d, m, y = path.parent.name, path.parent.parent.name, path.parent.parent.parent.name
    if re.fullmatch(r"\d{2}", d) and re.fullmatch(r"\d{2}", m) and re.fullmatch(r"\d{4}", y):
        return f"{y}-{m}-{d}"
    return None


def _same(a: float | None, b: float | None) -> bool:
    return a is None or b is None or abs(a - b) < 1e-9


def read_analysis_state(run: Path, thr: float, stride: int | None = None,
                        model_dir: Path | None = None, check_reuse: bool = False
                        ) -> AnalysisState:
    """What is already on disk. Reads the CURRENT folder names first, then the
    LEGACY ones (runs measured before 2026-10-05 carry measurement_/classical_;
    re-measuring 20+ GB to rename a folder would be absurd)."""
    import json

    an = Path(run) / "shadowgraph" / "analysis"
    state = AnalysisState()
    for meas_name, clas_name, legacy in (
            (spec.MEAS_DIR, spec.CLAS_DIR, False),
            (spec.LEGACY_MEAS_DIR, spec.LEGACY_CLAS_DIR, True)):
        m, c = an / meas_name.format(thr=thr), an / clas_name.format(thr=thr)
        if (m / "summary.json").is_file() and (c / "classical_summary.json").is_file():
            state.measured, state.legacy_names = True, legacy
            state.meas_dir, state.clas_dir = m, c
            try:
                prov = json.loads((m / "summary.json").read_text(encoding="utf-8")
                                  ).get("provenance", {})
                state.sizer_version = prov.get("sizer_version")
            except (OSError, ValueError):
                pass
            break
    if not state.measured:
        # Measured by the model stage but never classically (e.g. 10/01's
        # ..._nobh run, which predates classical_liquid.py). Not usable for the
        # statistics, which need both, but not "new" either.
        state.droplets_only = any(
            (an / name.format(thr=thr) / "summary.json").is_file()
            for name in (spec.MEAS_DIR, spec.LEGACY_MEAS_DIR))
    if check_reuse and stride is not None:
        try:
            state.reusable = bool(spec.rdc_module("process_capture").can_reuse(
                Path(run), stride, model_dir))
        except Exception:
            state.reusable = False
    return state


def load_run_info(path: Path, *, thr: float | None = None, stride: int | None = None,
                  model_dir: Path | None = None, check_reuse: bool = True) -> RunInfo:
    """Describe one run folder. Never raises on bad data: every problem becomes
    an Issue on the returned RunInfo."""
    path = Path(path)
    parsed = parse_run_name(path.name)
    info = RunInfo(path=path, name=path.name, name_ok=parsed is not None)
    if parsed:
        info.time = parsed["time"]
        for attr in ("sccm", "rpm", "sps", "orifice_mm", "bubbler_height"):
            setattr(info, attr, parsed[attr])
            if parsed[attr] is not None:
                info.sources[attr] = "name"
    else:
        info.issues.append(Issue("error", f"folder name is not {NAME_FORMAT}"))

    if not path.is_dir():
        info.issues.append(Issue("error", "folder does not exist"))
        return info

    # ---- the .cine -------------------------------------------------------
    list_files = spec.rdc_module("_fsutil").list_files
    cines = list_files(path / "shadowgraph" / "raw" / "CINE", "*.cine")
    info.n_cines = len(cines)
    if len(cines) == 1:
        info.cine = cines[0]
        try:
            info.cine_bytes = cines[0].stat().st_size
        except OSError:
            pass

    # ---- already measured? ------------------------------------------------
    s = spec.effective(spec.RunSettings(score_thresh=thr, stride=stride))
    info.analysis = read_analysis_state(path, s.score_thresh, s.stride, model_dir,
                                        check_reuse=check_reuse)
    if len(cines) != 1:
        problem = ("no .cine in shadowgraph/raw/CINE" if not cines
                   else f"{len(cines)} .cine files in shadowgraph/raw/CINE (expected 1)")
        # Fatal only if there is nothing already measured to analyse.
        info.issues.append(Issue("warn" if info.analysis.measured else "error",
                                 problem + (" (but already measured)"
                                            if info.analysis.measured else "")))

    # ---- run_summary.xlsx -------------------------------------------------
    xlsx = path / SUMMARY_XLSX
    info.has_xlsx = xlsx.is_file()
    if not info.has_xlsx:
        info.issues.append(Issue("warn", f"no {SUMMARY_XLSX} (conditions come from the "
                                         f"folder name only)"))
    else:
        try:
            info.meta = read_metadata(xlsx)
        except Exception as exc:
            info.issues.append(Issue("warn", f"{SUMMARY_XLSX} unreadable: "
                                             f"{type(exc).__name__}: {exc}"))
    _apply_metadata(info)
    if info.name_ok:
        for attr, _, what in WORKBOOK_FIELDS:
            if getattr(info, attr) is None:
                info.issues.append(Issue(
                    "warn", f"{what} is not in the folder name or {SUMMARY_XLSX}"))
        if info.sccm is None:
            info.issues.append(Issue("warn", "no gas-flow setpoint in the folder name "
                                             "(the workbook only records the achieved range)"))

    if info.date is None:
        info.date = _date_from_path(path)
    return info


def _meta_number(info: RunInfo, key: str) -> float | None:
    v = info.meta.get(key)
    return v.pick("max") if v is not None else None


def _apply_metadata(info: RunInfo) -> None:
    meta = info.meta
    if not meta:
        return
    notes = meta.get("Notes")
    info.notes = notes.text if notes is not None else ""
    fps = meta.get("FPS")
    info.fps = fps.number if fps is not None else None
    info.flow_achieved_sccm = _meta_number(info, "Flow Range (sccm)")

    ts = meta.get("Timestamp")
    if ts is not None and ts.text:
        try:
            info.date = datetime.fromisoformat(ts.text).strftime("%Y-%m-%d")
        except ValueError:
            info.issues.append(Issue("warn", f"unreadable Timestamp '{ts.text}'"))

    if not info.name_ok:
        return
    for attr, key, what in WORKBOOK_FIELDS:
        xl = _meta_number(info, key)
        if xl is None:
            continue                        # workbook has nothing: keep the name's value
        named = getattr(info, attr)
        if named is not None and not _same(named, xl):
            info.issues.append(Issue(
                "warn", f"folder name says {what} {named:g} but {SUMMARY_XLSX} says "
                        f"{xl:g} -- using the workbook's"))
        setattr(info, attr, xl)
        info.sources[attr] = "workbook"
    if info.flow_achieved_sccm and info.sccm:
        drift = abs(info.flow_achieved_sccm - info.sccm) / info.sccm
        if drift > FLOW_DRIFT_WARN:
            info.issues.append(Issue(
                "warn", f"achieved flow {info.flow_achieved_sccm:g} sccm is "
                        f"{drift:.0%} from the {info.sccm} sccm in the folder name"))


# ============================================================================
# finding run folders
# ============================================================================


@dataclass
class ScanResult:
    runs: list[Path] = field(default_factory=list)
    skipped: list[tuple[Path, str]] = field(default_factory=list)


def _subdirs(path: Path) -> list[Path]:
    try:
        with os.scandir(path) as it:
            return sorted((Path(e.path) for e in it
                           if e.is_dir(follow_symlinks=True) and not e.name.startswith(".")),
                          key=lambda p: p.name)
    except OSError:
        return []


def find_runs(roots: Iterable[Path], max_depth: int = DEFAULT_SCAN_DEPTH) -> ScanResult:
    """Run folders at or under each root. A root that IS a run folder is taken
    as-is; otherwise its subfolders are searched down to `max_depth`.

    Never descends into a run folder (its shadowgraph/ tree holds thousands of
    frames on exFAT). A folder that looks like a run but is named in the older
    scheme is reported in `skipped` with the reason, not silently ignored."""
    result = ScanResult()
    seen: set[Path] = set()

    def add(p: Path) -> None:
        key = p.resolve()
        if key not in seen:
            seen.add(key)
            result.runs.append(p)

    def walk(p: Path, depth: int) -> None:
        if RUN_NAME_RE.match(p.name):
            add(p)
            return
        if (p / "shadowgraph").is_dir():
            result.skipped.append((p, f"name is not {NAME_FORMAT}"))
            return
        if depth < max_depth:
            for child in _subdirs(p):
                walk(child, depth + 1)

    for root in roots:
        root = Path(root)
        if not root.is_dir():
            result.skipped.append((root, "not a folder"))
        elif RUN_NAME_RE.match(root.name):
            add(root)
        else:
            walk(root, 0)
    result.runs.sort(key=lambda p: (_sort_key(p), str(p)))
    return result


def _sort_key(path: Path) -> str:
    """Chronological: the date from the path where there is one, then HHMMSS."""
    m = RUN_NAME_RE.match(path.name)
    return f"{_date_from_path(path) or '0000-00-00'} {m['time'] if m else path.name}"


def load_runs(paths: Iterable[Path], progress: Callable[[int, int, RunInfo], None] | None = None,
              **kwargs) -> list[RunInfo]:
    """load_run_info for each path, in order, optionally reporting progress."""
    paths = list(paths)
    out = []
    for i, p in enumerate(paths, 1):
        info = load_run_info(p, **kwargs)
        out.append(info)
        if progress:
            progress(i, len(paths), info)
    out.sort(key=lambda r: (f"{r.date or '0000-00-00'} {r.time or r.name}", str(r.path)))
    return out


# ============================================================================
# summary
# ============================================================================


@dataclass
class Summary:
    n_runs: int = 0
    n_usable: int = 0
    n_errors: int = 0
    n_warnings: int = 0
    n_measured: int = 0
    n_days: int = 0
    n_unassigned: int = 0                       # no name-derived conditions
    replicates: dict[tuple, int] = field(default_factory=dict)   # condition -> run count

    @property
    def n_conditions(self) -> int:
        return len(self.replicates)

    def headline(self) -> str:
        if not self.n_runs:
            return "0 runs selected"
        counts = sorted(set(self.replicates.values()))

        def plural(n: int, word: str) -> str:
            return f"{n} {word}{'' if n == 1 else 's'}"

        if len(counts) == 1:
            design = f"{plural(self.n_conditions, 'condition')} × {plural(counts[0], 'replicate')}"
        elif counts:
            design = (f"{self.n_conditions} conditions, {counts[0]}–{counts[-1]} "
                      f"replicates (unbalanced)")
        else:
            design = "no conditions"
        parts = [f"{self.n_runs} runs", design, f"{self.n_measured} already measured"]
        if self.n_days > 1:
            parts.append(f"{self.n_days} days")
        if self.n_errors:
            parts.append(f"{self.n_errors} with errors")
        return " · ".join(parts)


def summarise(runs: Iterable[RunInfo]) -> Summary:
    runs = list(runs)
    s = Summary(n_runs=len(runs))
    s.n_usable = sum(r.usable for r in runs)
    s.n_errors = sum(bool(r.errors) for r in runs)
    s.n_warnings = sum(bool(r.warnings) for r in runs)
    s.n_measured = sum(r.analysis.measured for r in runs)
    s.n_days = len({r.date for r in runs if r.date})
    for r in runs:
        key = r.condition_key
        if key is None:
            s.n_unassigned += 1
        else:
            s.replicates[key] = s.replicates.get(key, 0) + 1
    return s


def sizer_versions(runs: Iterable[RunInfo]) -> set[str]:
    """Distinct sizer versions among the already-measured runs. More than one
    means the set is not comparable: taguchi_analysis.py refuses to mix them
    (2.0.0 vs 2.1.0 moved the atomised fraction ~3% relative), so the UI says so
    before the user has spent hours on it. A run whose summary records no
    sizer_version was measured before the sizer existed (diameters from the model
    mask area) -- process_capture's own log calls that "pre-2.0.0"."""
    return {r.analysis.sizer_version or PRE_SIZER
            for r in runs if r.analysis.measured}
