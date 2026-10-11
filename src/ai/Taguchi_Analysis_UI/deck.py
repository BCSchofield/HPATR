"""deck.py -- taguchi_deck.pptx: the whole analysis as slides, every figure on one. No Qt.

    title -> design -> what was found -> the effects (main effects, contributions, every run, S/N,
    GLR) -> size (spread by count and volume, distributions, percentiles) -> steadiness (by
    condition, spectra, time series) -> drift, replicates, interactions -> model checks,
    correlations -> liquid pieces -> data quality -> per-condition summary -> caveats ->
    appendix: the image galleries

STYLE. Built on assets/deck_template.pptx: Ben's hand-made 9x3 deck ("Atomisation Ink" theme:
Cambria headings, navy title slides, grey-blue cards, orange / teal / purple / yellow accents)
with its slides removed, so the generated deck looks like the decks he makes by hand.

Every figure slide has a TITLE that is its takeaway and a card with a DESCRIPTION of what the
figure shows plus the DEFINITIONS of the terms it uses, written from the numbers (the same tables
and stats the report uses, so the deck cannot disagree with the report), and speaker notes on how
to read it. A condition named on a slide carries its factor levels, and where it stands out, the
factor driving it (the factor whose level pushes that condition's value furthest the same way,
with that factor's share of the variation). Tall figures (per-run time series, all responses'
replicates, model checks) are redrawn in slide-sized parts into figures/slides/ (PNG only). A
figure that was not drawn leaves its slide out.

python-pptx is optional: `available()` says why the deck cannot be made, and publish() then
writes everything else and records the reason as a warning.
"""
from __future__ import annotations

import math
import os
import re
from pathlib import Path

import numpy as np

from . import figures_extra, fsops, report, run_metrics, stats, tables

DECK_NAME = "taguchi_deck.pptx"
TEMPLATE = Path(__file__).parent / "assets" / "deck_template.pptx"
SLIDE_FIG_DIR = "slides"                           # figures/slides/: slide-sized parts of tall figures
WIDTH_EMU, HEIGHT_EMU = 9144000, 5143500          # 10 x 5.625 in (16:9), as the hand-made 9x3 deck
SERIES_PER_SLIDE = 3                               # conditions per time-series slide
SIZE_KEYS = ("d32", "d_mean", "d_median", "d90", "dv50")

# the template's theme ("Atomisation Ink")
INK = (0x1B, 0x26, 0x30)
NAVY = (0x22, 0x31, 0x3F)
BODY = (0x3A, 0x46, 0x50)
MUTED = (0x5A, 0x66, 0x70)
CARD = (0xEE, 0xF2, 0xF4)
WHITE = (0xFF, 0xFF, 0xFF)
ORANGE = (0xE0, 0x7A, 0x3F)
TEAL = (0x2A, 0x9D, 0x8F)
PURPLE = (0x7B, 0x6C, 0xB0)
YELLOW = (0xF2, 0xC1, 0x4E)
GREY = (0x9A, 0xA5, 0xAD)

# ---- definitions, shown on the slides that use each term -------------------------------------------

DEFINE = {
    "D32": "D32 (Sauter mean diameter): Σd³/Σd² of the in-focus droplets: the size of a droplet with the spray's "
           "volume-to-surface ratio.",
    "D10": "D10 (mean diameter): the arithmetic mean of the in-focus diameters.",
    "Dv": "Dv10 / Dv50 / Dv90: the diameters below which 10 / 50 / 90% of the liquid volume (Σd³) lies; Dv50 is "
          "the volume median.",
    "span": "Span: (Dv90 − Dv10) / Dv50, how wide the volume distribution is relative to its median.",
    "lattice": "Pixel lattice: diameters come from whole-pixel areas, so only certain values occur; below 50 µm "
               "they are too coarse to trust.",
    "main effect": "Main effect: the mean of every run at one level of a factor.",
    "SD": "SD (standard deviation): the typical spread of values about their mean.",
    "SE": "SE (standard error of a mean): √(MS error / n), how precisely the mean is known, from the ANOVA error.",
    "CI": "CI (95% confidence interval): a range that would contain the true value in 95% of repeated experiments.",
    "bootstrap": "Frame-bootstrap CI: frames resampled within a run, so it measures noise inside the run, not "
                 "between runs.",
    "t CI": "Replicate CI: mean ± t·SD/√n over the replicates (Student's t), i.e. run-to-run variation.",
    "CV": "CV (coefficient of variation): SD / mean, spread relative to the mean.",
    "p": "p: the chance of an effect at least this large if the factor did nothing; below 0.05 is evidence of an "
         "effect.",
    "q": "q (BH): p adjusted for the number of tests run (Benjamini–Hochberg false discovery rate).",
    "contribution": "% of variation: the factor's sum of squares (Type II) as a share of the response's total.",
    "lack of fit": "Lack of fit: what the factors acting separately cannot explain (interactions), tested "
                   "against replicate noise.",
    "pure error": "Error (pure error): the spread between replicates of the same condition, i.e. run-to-run noise.",
    "S/N": "S/N (signal-to-noise, dB): per condition from its replicates, −10·log10(mean y²) when smaller is "
           "better, −10·log10(mean 1/y²) when larger is better. Higher is always better: it rewards a good mean "
           "and a small spread together.",
    "GLR": "GLR: gas-to-liquid mass-flow ratio; it combines the gas and liquid flows, so it is a trend, not a factor.",
    "power law": "Power law: the fit y = a·GLRᵇ; b is the exponent.",
    "r": "r (Pearson): linear correlation, −1 to +1.",
    "count CV": "Droplet-count CV: SD / mean of the droplets per frame (all droplets).",
    "thin": "Frames below half: % of frames with fewer than half the run's mean droplet count.",
    "lag-1": "Lag-1 r: correlation of each frame's droplet count with the next frame's; near 0 = independent "
             "frames, high = bursts lasting several frames.",
    "slow share": "Slow share: % of the droplet-count variance below 2 Hz in the Welch spectrum; white noise "
                  "would put about 10% there.",
    "Welch": "Welch spectrum: the count series is cut into overlapping segments (256 frames, 50% overlap, linear "
             "trend removed), each is windowed and Fourier transformed, and the power is averaged: how the "
             "count's variance is spread over frequency.",
    "r(p,n)": "r(pressure, droplets): correlation between the logged upstream pressure and the droplet count "
              "inside the camera window.",
    "drift": "Residual drift: straight-line slope of the main-effects residuals against capture time.",
    "residual": "Standardised residual: (observed − fitted) / (residual SD·√(1 − leverage)); beyond ±2 happens "
                "about 1 time in 20 by chance.",
    "normal plot": "Normal probability plot: sorted residuals against normal quantiles; a straight line means "
                   "normally distributed errors.",
    "aliased": "Aliased: two effects the design cannot tell apart.",
    "non-additive": "Not additive: share of the spread of the cell means that parallel lines (factors acting "
                    "separately) cannot reproduce.",
    "robust z": "Robust z: (value − median) / (1.4826·MAD); how unusual a run is, not swayed by the outliers "
                "themselves.",
    "moved": "Moved by the sharpness rule: round droplets sizer 2.2.0 reclassified as out of focus.",
    "pieces": "Un-atomised pieces: connected regions of the classical liquid mask (10 µm/px).",
    "atomised": "Atomised fraction: share of the liquid area that is droplets (classical).",
}


def defs(*keys: str) -> str:
    return "  ·  ".join(DEFINE[k] for k in keys if k in DEFINE)


def available() -> str | None:
    """None when the deck can be built, else the reason (and the fix)."""
    try:
        import pptx  # noqa: F401
    except ImportError:
        return ("python-pptx is not installed in this Python environment, so no PowerPoint deck was written "
                "(install it with: python -m pip install python-pptx)")
    return None


# ---- small helpers ----------------------------------------------------------------------------------

def _short(label: str) -> str:
    return label.split(" (")[0]


def _plain(md: str) -> str:
    """Markdown emphasis out, for slide text."""
    return re.sub(r"\*\*(.+?)\*\*|\*(.+?)\*", lambda m: m.group(1) or m.group(2), md).strip()


def _p(p) -> str:
    return report.p_eq(p) if p is not None else "untested"


def _f(x, nd: int = 3) -> str:
    return report.fmt(x, nd)


def _label(results, key: str) -> str:
    return _short(results.responses[key].label)


def _conds(results) -> dict[str, list]:
    out: dict[str, list] = {}
    for r in sorted(results.runs, key=lambda r: (tables.cond_order(r.condition), r.replicate or 0, r.name)):
        out.setdefault(r.condition, []).append(r)
    return out


def _unit(f) -> str:
    m = re.search(r"\(([^)]+)\)", f.label)
    if m:
        return m.group(1)
    return "rpm" if "rpm" in f.label.lower() else ""


def _level(f, v) -> str:
    u = _unit(f)
    return f"{tables.level_text(v)} {u}".strip()


def cond_text(results, c: str) -> str:
    """'T7 (3000 sccm, 900 rpm, 8000 steps/s)'."""
    rs = _conds(results).get(c, [])
    if not rs:
        return c
    return f"{c} (" + ", ".join(_level(f, rs[0].levels[f.key]) for f in results.factors) + ")"


def _cond_levels(results, c: str) -> str:
    rs = _conds(results).get(c, [])
    return ", ".join(_level(f, rs[0].levels[f.key]) for f in results.factors) if rs else ""


def _per_run(rows: list[dict], col: str) -> dict[str, float]:
    return {r["run"]: float(r[col]) for r in rows
            if isinstance(r.get(col), (int, float)) and math.isfinite(r[col])}


def _cond_mean(results, per_run: dict[str, float]) -> dict[str, float]:
    acc: dict[str, list] = {}
    for r in results.runs:
        if r.name in per_run:
            acc.setdefault(r.condition, []).append(per_run[r.name])
    return {c: float(np.mean(v)) for c, v in acc.items()}


def _by_condition_mean(rows: list[dict], col: str) -> dict[str, float]:
    acc: dict[str, list] = {}
    for r in rows:
        v = r.get(col)
        if isinstance(v, (int, float)) and math.isfinite(v):
            acc.setdefault(r["condition"], []).append(v)
    return {c: float(np.mean(v)) for c, v in acc.items()}


def _extremes(d: dict[str, float]):
    if not d:
        return None
    lo = min(d, key=d.get)
    hi = max(d, key=d.get)
    return (lo, d[lo]), (hi, d[hi])


def _analysed(results, key: str) -> bool:
    return key in results.responses and results.responses[key].anova is not None


def _values(results, key: str) -> dict[str, float]:
    return {r.name: float(r.values[key]) for r in results.runs
            if r.values.get(key) is not None and math.isfinite(r.values[key])}


def driver(results, per_run: dict[str, float], c: str) -> str:
    """What pushes condition c's value of a measure away from the average: the factor whose level
    (at c) moves the measure furthest in the same direction, and that factor's share of the
    measure's variation (main-effects ANOVA on the per-run values). '' if nothing to say."""
    runs = [r for r in results.runs if r.name in per_run]
    if len(runs) < 4 or not results.factors:
        return ""
    y = np.array([per_run[r.name] for r in runs])
    grand = float(y.mean())
    cr = [r for r in runs if r.condition == c]
    if not cr:
        return ""
    dev = float(np.mean([per_run[r.name] for r in cr])) - grand
    if dev == 0:
        return ""
    a = stats.anova(list(y), {f.key: [r.levels[f.key] for r in runs] for f in results.factors},
                    [r.condition for r in runs])
    best = None
    for f in results.factors:
        lv = cr[0].levels[f.key]
        eff = float(np.mean([per_run[r.name] for r in runs if r.levels[f.key] == lv])) - grand
        if eff * dev > 0 and (best is None or abs(eff) > abs(best[1])):
            best = (f, eff, lv)
    if best is None:
        return ""
    f, eff, lv = best
    levels = sorted({r.levels[f.key] for r in runs}, key=lambda x: (str(type(x)), x))
    where = ("lowest" if lv == levels[0] else "highest" if lv == levels[-1] else "middle") if len(levels) > 2 else \
            ("low" if lv == levels[0] else "high")
    row = a.rows.get(f.key)
    share = (f", {row.contribution_pct:.0f}% of the variation"
             if row is not None and row.contribution_pct is not None else "")
    return f"driven by its {where} {_short(f.label).lower()} ({_level(f, lv)}{share})"


def _with_driver(results, per_run, c) -> str:
    d = driver(results, per_run, c)
    return f"{cond_text(results, c)}, {d}" if d else cond_text(results, c)


def top_factor(results, key: str):
    """(factor label, % of variation, p) of the factor with the largest share, or None."""
    res = results.responses.get(key)
    if res is None or res.anova is None:
        return None
    best = None
    for f in results.factors:
        row = res.anova.rows.get(f.key)
        if row is None or row.contribution_pct is None:
            continue
        if best is None or row.contribution_pct > best[1]:
            best = (f.label, row.contribution_pct, row.p)
    return best


def _replicate_cvs(results, key: str) -> dict[str, float]:
    out = {}
    for c, rs in _conds(results).items():
        v = np.array([r.values[key] for r in rs if r.values.get(key) is not None], float)
        if v.size >= 2 and v.mean():
            out[c] = float(100 * v.std(ddof=1) / abs(v.mean()))
    return out


def lattice_bound(results, key: str) -> bool:
    """A response that takes few distinct values across the runs (it sits on the pixel lattice, or
    near zero): replicates can be identical, so its CV says nothing about repeatability."""
    res = results.responses.get(key)
    return res is not None and (res.quantisation_warning or res.n_distinct <= len(results.runs) // 2)


# a slide spec: (title, description, definitions, notes)

DESCRIBE = {
    "fig_main_effects": "Mean of every run at each factor level, for the headline responses.",
    "fig_contributions": "Share of each response's total variation by source: factors, lack of fit, error.",
    "fig_per_run": "Every run's value with its frame-bootstrap 95% CI, grouped by condition.",
    "fig_sn": "Taguchi signal-to-noise ratio by factor level (higher is better).",
    "fig_glr": "Each run against its gas-to-liquid ratio, with the power-law fit.",
    "fig_size_spread_count": "Share of each run's in-focus droplets per size bin.",
    "fig_size_spread_volume": "Share of each run's liquid volume per size bin.",
    "fig_size_distribution": "Number and volume distributions per condition, and cumulative volume.",
    "fig_size_percentiles": "Dv10 to Dv90 and Dv50 for every run.",
    "fig_steadiness": "Droplet-count CV, frames below half the mean, lag-1 correlation and slow share, by condition.",
    "fig_spectra": "Welch spectrum of the droplet count per frame, per replicate.",
    "fig_series": "Droplets per frame through each capture, with the logged pressure.",
    "fig_series_liquid": "Liquid area per frame through each capture.",
    "fig_drift": "Responses against capture time, coloured by condition.",
    "fig_replicates": "Replicates per condition with the 95% CI of the condition mean.",
    "fig_interactions": "Condition means for each pair of factors.",
    "fig_diagnostics": "Standardised residuals of the main-effects fit: vs fitted, normal plot, vs time.",
    "fig_correlations": "Correlation of every pair of analysed responses across runs.",
    "fig_pieces": "Un-atomised liquid pieces: size distribution, count per frame, largest piece.",
    "fig_quality": "Sizer and data-quality checks per run; red = far from the other runs.",
    "gallery": "Each run's least and most atomised frame (classical overlays).",
}


def _describe(image: Path) -> str:
    stem = Path(image).stem
    for key in sorted(DESCRIBE, key=len, reverse=True):
        if stem.startswith(key):
            return DESCRIBE[key]
    return ""


# ---- takeaways ----------------------------------------------------------------------------------------

def _and(items: list[str]) -> str:
    return items[0] if len(items) == 1 else ", ".join(items[:-1]) + " and " + items[-1]


def _names(results, keys: list[str]) -> str:
    """Response names for a title: short (no ", in-focus" / ", classical"), first capitalised, joined."""
    return _and([_label(results, k).split(",")[0] if i == 0 else _label(results, k).split(",")[0].lower()
                 for i, k in enumerate(keys)])


def tw_main_effects(results):
    tops = [(k, t) for k in report.HEADLINE if (t := top_factor(results, k)) is not None]
    if not tops:
        title = "Main effects"
    elif len(tops) >= 2 and len({t[0] for _, t in tops}) == 1:
        title = (f"{_short(tops[0][1][0])} drives {'both ' if len(tops) == 2 else ''}"
                 + _names(results, [k for k, _ in tops]))
    else:
        title = "; ".join(f"{_label(results, k)} follows {_short(t[0]).lower()}" for k, t in tops)
    bits = []
    for k in report.HEADLINE:
        res = results.responses.get(k)
        if res is None or not res.effects:
            continue
        spans = []
        for f in results.factors:
            effs = res.effects.get(f.key, [])
            if effs:
                lo = min(effs, key=lambda e: e.mean)
                hi = max(effs, key=lambda e: e.mean)
                spans.append(f"{_short(f.label)} {_f(lo.mean)} ({_level(f, lo.level)}) to {_f(hi.mean)} "
                             f"({_level(f, hi.level)})")
        t = top_factor(results, k)
        bits.append(f"{_label(results, k)}: " + "; ".join(spans)
                    + (f". {_short(t[0])} explains {t[1]:.0f}% ({_p(t[2])})" if t else ""))
    return (title, ". ".join(bits) + ".", defs("main effect", "SE", "bootstrap", "p", "D32", "D10", "atomised"),
            "Each panel: one factor, its levels on the x axis. Dashed line = grand mean; black bars = SE; grey "
            "band = frame-bootstrap 95% CI. p from the ANOVA, top right.")


def tw_contributions(results):
    items = [(k, top_factor(results, k)) for k, r in results.responses.items() if r.anova and r.anova.error]
    items = [(k, t) for k, t in items if t]
    if not items:
        return "Where the variation comes from", "", defs("contribution"), ""
    counts: dict[str, int] = {}
    for _, t in items:
        counts[t[0]] = counts.get(t[0], 0) + 1
    f, n = max(counts.items(), key=lambda kv: kv[1])
    err = sorted(((r.anova.error.contribution_pct or 0, k) for k, r in results.responses.items()
                  if r.anova and r.anova.error), reverse=True)
    desc = (f"Share of each response's total variation by source. {_short(f)} is the largest share for {n} of "
            f"{len(items)} responses. Most left to run-to-run noise: {_label(results, err[0][1])} "
            f"({err[0][0]:.0f}%); least: {_label(results, err[-1][1])} ({err[-1][0]:.0f}%).")
    return (f"{_short(f)} is the largest source for {n} of {len(items)} responses", desc,
            defs("contribution", "lack of fit", "pure error"),
            "A response with a big error share is one the factors barely move.")


def tw_per_run(results):
    out, med = [], []
    for k in report.HEADLINE:
        if not _analysed(results, k):
            continue
        vals = _values(results, k)
        cvs = _replicate_cvs(results, k)
        if not cvs:
            continue
        worst = max(cvs, key=cvs.get)
        lo = min(vals, key=vals.get)
        hi = max(vals, key=vals.get)
        cond = {r.name: r.condition for r in results.runs}
        med.append(float(np.median(list(cvs.values()))))
        out.append(f"{_label(results, k)} from {_f(vals[lo])} ({lo[:6]}, {cond_text(results, cond[lo])}) to "
                   f"{_f(vals[hi])} ({hi[:6]}, {cond_text(results, cond[hi])}); replicate CV median {med[-1]:.1f}%, "
                   f"widest {cond_text(results, worst)} at {cvs[worst]:.1f}%")
    title = f"Replicates agree to within {max(med):.0f}% (median CV)" if med else "Every run"
    return (title, "; ".join(out) + ".", defs("CV", "bootstrap"),
            "One bar per run, grouped by condition (replicates share a colour).")


def tw_sn(results):
    bits = []
    for k in report.HEADLINE:
        res = results.responses.get(k)
        if res is None or not res.sn_defined:
            continue
        best = []
        for f in results.factors:
            eff = res.sn_effects.get(f.key)
            if eff:
                lv = max(eff, key=eff.get)
                best.append(_level(f, lv))
        bits.append((k, best))
    if not bits:
        return "Signal-to-noise ratio", "", defs("S/N"), ""
    same = len(bits) >= 2 and all(b == bits[0][1] for _, b in bits)
    title = (f"S/N favours {', '.join(bits[0][1])} for " +
             (("both " if len(bits) == 2 else "") + _names(results, [k for k, _ in bits])
              if same else _label(results, bits[0][0])))
    desc = " ".join(f"Best levels for {_label(results, k)}: {', '.join(b)}." for k, b in bits)
    return (title, desc, defs("S/N"), "")


def tw_glr(results):
    trends = [(k, results.trends[(k, "glr")]) for k in report.HEADLINE
              if (k, "glr") in results.trends and results.trends[(k, "glr")].r is not None]
    if not trends:
        return "Gas-to-liquid ratio", "", defs("GLR"), ""
    bits = [f"{_label(results, k)}: r = {t.r:+.2f}, {_p(t.p)}"
            + (f", exponent {t.exponent:+.2f}" if t.exponent is not None else "") for k, t in trends]
    k, t = max(trends, key=lambda kt: abs(kt[1].r))
    return (f"{_label(results, k)} {'rises' if t.r > 0 else 'falls'} with GLR (r = {t.r:+.2f})",
            "Each run against its gas-to-liquid ratio. " + "; ".join(bits) + ".",
            defs("GLR", "power law", "r", "p"), "")


def _bins_run_share(bins, measure: str, *, below=None, above=None) -> dict[str, float]:
    keep = [i for i, (lo, hi) in enumerate(zip(bins.edges[:-1], bins.edges[1:]))
            if (below is not None and hi <= below) or (above is not None and lo >= above)]
    out = {}
    for s in bins.runs:
        pct = s.pct_count if measure == "count" else s.pct_volume
        out[s.name] = float(sum(pct[i] for i in keep))
    return out


def _bins_share(bins, measure: str, *, below=None, above=None) -> dict[str, float]:
    rows = bins.table(measure, "condition")
    keep = [lab for lab, lo, hi in zip(bins.labels, bins.edges[:-1], bins.edges[1:])
            if (below is not None and hi <= below) or (above is not None and lo >= above)]
    return {r["condition"]: float(sum(r.get(b) or 0.0 for b in keep)) for r in rows}


def tw_spread_count(results, bins):
    ex = _extremes(_bins_share(bins, "count", below=run_metrics.SMALL_UM)) if bins is not None and bins.runs else None
    if ex is None:
        return "Droplet size spread by count", "", defs("lattice"), ""
    (lo, vlo), (hi, vhi) = ex
    per = _bins_run_share(bins, "count", below=run_metrics.SMALL_UM)
    return (f"{vhi:.0f}% of {hi}'s droplets are below {run_metrics.SMALL_UM:g} µm, against {vlo:.0f}% at {lo}",
            f"Share of each run's in-focus droplets per size bin. Most small droplets: {_with_driver(results, per, hi)}. "
            f"Fewest: {_with_driver(results, per, lo)}.",
            defs("lattice"), "Each bar sums to 100% of that run's in-focus droplets.")


def tw_spread_volume(results, bins):
    ex = _extremes(_bins_share(bins, "volume", above=run_metrics.LARGE_UM)) if bins is not None and bins.runs else None
    if ex is None:
        return "Where the liquid is", "", "", ""
    (lo, vlo), (hi, vhi) = ex
    per = _bins_run_share(bins, "volume", above=run_metrics.LARGE_UM)
    return (f"{vhi:.0f}% of {hi}'s liquid is in droplets above {run_metrics.LARGE_UM:g} µm, against {vlo:.0f}% "
            f"for {lo}",
            f"Share of each run's liquid volume (Σd³) per size bin: a few large droplets carry much of the liquid. "
            f"Most: {_with_driver(results, per, hi)}. Least: {_with_driver(results, per, lo)}.",
            defs("lattice"), "Stacked bars summing to 100% of each run's droplet volume.")


def tw_size_distribution(results):
    rows = [r for r in tables.size_percentiles_by_condition(results) if r.get("Dv50 (um)") is not None]
    if not rows:
        return "Droplet size distributions", "", defs("Dv"), ""
    lo = min(rows, key=lambda r: r["Dv50 (um)"])
    hi = max(rows, key=lambda r: r["Dv50 (um)"])
    big = f"% volume > {run_metrics.LARGE_UM:g} um"
    per = _per_run(tables.size_percentiles(results), "Dv50 (um)")
    return (f"Dv50 runs from {lo['Dv50 (um)']:.1f} µm ({lo['condition']}) to {hi['Dv50 (um)']:.1f} µm "
            f"({hi['condition']})",
            f"Number and volume distributions per condition (replicates pooled) and cumulative volume. Coarsest: "
            f"{_with_driver(results, per, hi['condition'])}, with {hi.get(big, 0):.0f}% of its liquid above "
            f"{run_metrics.LARGE_UM:g} µm. Finest: {_with_driver(results, per, lo['condition'])}, with "
            f"{lo.get(big, 0):.0f}% of its liquid above {run_metrics.LARGE_UM:g} µm.",
            defs("Dv", "lattice"),
            "Left: % of droplets per 5 µm bin. Middle: % of liquid volume per bin. Right: cumulative volume; "
            "where each curve crosses 10 / 50 / 90% is its Dv10 / Dv50 / Dv90.")


def tw_size_percentiles(results):
    rows = [r for r in tables.size_percentiles_by_condition(results) if r.get("span") is not None]
    if not rows:
        return "Volume percentiles per run", "", defs("Dv", "span"), ""
    lo = min(rows, key=lambda r: r["span"])
    hi = max(rows, key=lambda r: r["span"])
    d90 = max(rows, key=lambda r: r["Dv90 (um)"])
    per = _per_run(tables.size_percentiles(results), "span")
    return (f"Size spread (span) is widest at {hi['condition']} ({hi['span']:.2f}), narrowest at "
            f"{lo['condition']} ({lo['span']:.2f})",
            f"Dv10 to Dv90 (bar) and Dv50 (dot) for every run. Widest: {_with_driver(results, per, hi['condition'])}. "
            f"Narrowest: {_with_driver(results, per, lo['condition'])}. Largest Dv90: "
            f"{cond_text(results, d90['condition'])}, {d90['Dv90 (um)']:.0f} µm.",
            defs("Dv", "span"), "Replicates of one condition should sit together.")


def _steady(results, col):
    rows = tables.steadiness(results)
    return _per_run(rows, col), _by_condition_mean(rows, col)


def tw_steadiness(results):
    slow_run, slow = _steady(results, f"slow share < {run_metrics.SLOW_HZ:g} Hz (%)")
    _, lag = _steady(results, "lag-1 r")
    _, cv = _steady(results, "droplet-count CV")
    terms = defs("count CV", "thin", "lag-1", "slow share")
    if not slow:
        bits = [f"{name} {min(v.values()):.2f} to {max(v.values()):.2f}"
                for name, v in (("count CV", cv), ("lag-1 r", lag)) if v]
        return ("Steadiness and pulsing", "Steadiness measures by condition" + (": " + "; ".join(bits) if bits else "")
                + ".", terms, "")
    top = sorted(slow, key=slow.get, reverse=True)[:2]
    calm = min(slow, key=slow.get)
    desc = (f"Most pulsing: {_with_driver(results, slow_run, top[0])}, slow share {slow[top[0]]:.0f}%"
            + (f"; then {cond_text(results, top[1])}, slow share {slow[top[1]]:.0f}%" if len(top) > 1 else "")
            + f". Steadiest: {_with_driver(results, slow_run, calm)}, slow share {slow[calm]:.0f}%.")
    if lag and cv:
        desc += f" Across conditions lag-1 r runs {min(lag.values()):.2f}–{max(lag.values()):.2f}, count CV " \
                f"{min(cv.values()):.2f}–{max(cv.values()):.2f}."
    return ("Pulsing is strongest at " + " and ".join(top) + " (slow share "
            + ", ".join(f"{slow[c]:.0f}%" for c in top) + ")", desc, terms,
            "All droplets (in and out of focus), so unaffected by the sizer version.")


def tw_spectra(results):
    slow_run, slow = _steady(results, f"slow share < {run_metrics.SLOW_HZ:g} Hz (%)")
    _, peak = _steady(results, "slow peak (Hz)")
    if not slow:
        return "Droplet-count spectra", "", defs("Welch"), ""
    top = max(slow, key=slow.get)
    return (f"{top} has the strongest low-frequency pulsing" + (f", peaking near {peak[top]:.1f} Hz" if top in peak else ""),
            f"One line per replicate; shaded below 2 Hz. Strongest: {_with_driver(results, slow_run, top)}. Slow share "
            "by condition: " + ", ".join(f"{c} {slow[c]:.0f}%" for c in sorted(slow, key=slow.get, reverse=True)) + ".",
            defs("Welch", "slow share"),
            "A flat spectrum is noise; power piled up at the left is slow pulsing (the spray coming and going over "
            "seconds).")


def tw_series(results, part: list[str]):
    rp_run, rp = _steady(results, "r pressure vs droplets")
    _, cv = _steady(results, "droplet-count CV")
    rp = {c: v for c, v in rp.items() if c in part}
    if rp:
        c = max(rp, key=rp.get)
        title = f"{', '.join(part)}: droplets per frame track the pressure most at {c} (r = {rp[c]:+.2f})"
        desc = (f"Strongest coupling: {_with_driver(results, rp_run, c)}. " + "; ".join(
            f"{cond_text(results, k)}: r {rp[k]:+.2f}, count CV {cv.get(k, float('nan')):.2f}" for k in part
            if k in rp) + ".")
    else:
        title = "Droplets per frame through each capture: " + ", ".join(part)
        desc = "; ".join(cond_text(results, k) for k in part) + "."
    return (title, desc, defs("r(p,n)", "count CV"),
            "Blue: droplets per frame through the capture, dashed = run mean. Red: logged upstream pressure (barA). "
            "Pressure swings the droplet count follows mean the supply, not the atomiser, sets the pulsing.")


def tw_series_liquid(results, part: list[str]):
    cv_run, cv = _steady(results, "liquid-area CV")
    cv = {c: v for c, v in cv.items() if c in part}
    if not cv:
        return "Liquid area through each capture: " + ", ".join(part), "", defs("CV"), ""
    c = max(cv, key=cv.get)
    return (f"{', '.join(part)}: liquid area is burstiest at {c} (CV {cv[c]:.2f})",
            f"Burstiest: {_with_driver(results, cv_run, c)}. " +
            "; ".join(f"{cond_text(results, k)}: CV {v:.2f}" for k, v in cv.items()) + ".",
            defs("CV"), "Spikes are slugs of un-atomised liquid crossing the field of view.")


def tw_drift(results):
    pairs = stats.drifts(results)
    conf = [(f, r) for f, r in stats.order_confounding(results) if abs(r) >= report.CONFOUNDED_R]
    sig = [d for d, q in pairs if q < report.ALPHA]
    if not pairs:
        title = "Responses over the session"
    elif sig:
        title = "Drift over the session in " + ", ".join(_label(results, d.key) for d in sig)
    else:
        title = "No drift over the session beyond what the factors explain"
    desc = (f"Responses against capture time, coloured by condition. Residual drift tested for {len(pairs)} responses"
            + (f"; smallest p = {min(d.p for d, _ in pairs):.2f}" if pairs else "") + "."
            + (" " + ", ".join(f"{_short(f)} rose with run order (r = {r:+.2f})" for f, r in conf) + "." if conf else ""))
    return (title, desc, defs("drift", "q"),
            "A light check for anything slow (window dirt, gas supply, the silicone, nozzle fouling).")


def tw_replicates(results, keys: list[str], what: str):
    rel = {k: float(np.median(list(c.values()))) for k in keys if _analysed(results, k)
           and (c := _replicate_cvs(results, k))}
    ranked = {k: v for k, v in rel.items() if not lattice_bound(results, k)}
    if not ranked:
        return f"Replicates: {what}", "", defs("t CI", "CV"), ""
    best = min(ranked, key=ranked.get)
    worst = max(ranked, key=ranked.get)
    odd = [k for k in rel if k not in ranked]
    loosest = max(_replicate_cvs(results, worst).items(), key=lambda kv: kv[1], default=None)
    return (f"Replicates, {what}: {_label(results, best)} repeats best, {_label(results, worst)} worst",
            "Median replicate CV: " + ", ".join(f"{_label(results, k)} {v:.1f}%" for k, v in ranked.items()) + "."
            + (f" Loosest {_label(results, worst)}: {cond_text(results, loosest[0])}, CV {loosest[1]:.1f}%." if loosest else "")
            + (" Few distinct values (pixel lattice or near zero), so CV not meaningful: "
               + ", ".join(_label(results, k) for k in odd) + "." if odd else ""),
            defs("t CI", "CV"), "Dots = runs, black bar = condition mean, grey = 95% CI of the mean.")


def _nonadditive_share(results, key: str, fa, fb) -> float | None:
    """% of the spread of the fa x fb cell means that an additive (parallel-lines) fit misses."""
    la = sorted({r.levels[fa.key] for r in results.runs}, key=lambda x: (str(type(x)), x))
    lb = sorted({r.levels[fb.key] for r in results.runs}, key=lambda x: (str(type(x)), x))
    M = np.full((len(la), len(lb)), np.nan)
    for i, a in enumerate(la):
        for j, b in enumerate(lb):
            v = [r.values[key] for r in results.runs if r.levels[fa.key] == a and r.levels[fb.key] == b
                 and r.values.get(key) is not None]
            if v:
                M[i, j] = np.mean(v)
    if np.isnan(M).any() or M.size < 4:
        return None
    fit = M.mean(1, keepdims=True) + M.mean(0, keepdims=True) - M.mean()
    tot = ((M - M.mean()) ** 2).sum()
    return float(100 * ((M - fit) ** 2).sum() / tot) if tot > 0 else None


def tw_interactions(results):
    fs = results.factors
    best = None
    for k in figures_extra.INTERACTION_KEYS:
        if not _analysed(results, k):
            continue
        for i, fa in enumerate(fs):
            for fb in fs[i + 1:]:
                s = _nonadditive_share(results, k, fa, fb)
                if s is not None and (best is None or s > best[0]):
                    best = (s, k, fa, fb)
    fractional = not figures_extra.is_full_factorial(results)
    terms = defs("non-additive", "aliased", "lack of fit") if fractional else defs("non-additive", "lack of fit")
    if best is None:
        return "Condition means for each pair of factors", "", terms, ""
    s, k, fa, fb = best
    others = [_short(f.label).lower() for f in fs if f not in (fa, fb)]
    alias = f" (aliased with {', '.join(others)})" if others and fractional else ""
    return (f"Least parallel: {_label(results, k)}, {_short(fa.label).lower()} × {_short(fb.label).lower()}{alias}",
            f"For {_label(results, k)}, {s:.0f}% of the spread of the {_short(fa.label).lower()} × "
            f"{_short(fb.label).lower()} cell means is not additive."
            + (f" In this fractional design each cell also changes {', '.join(others)}, so most of that is likely "
               f"{', '.join(others)}'s own effect rather than an interaction." if alias else ""),
            terms, "Parallel lines = the two factors act independently.")


def tw_diagnostics(results, keys):
    worst = None
    for k in keys:
        fit = stats.main_effects_fit(results, k) if _analysed(results, k) else None
        if fit is None:
            continue
        for r, z in zip(fit.runs, fit.std_resid):
            if math.isfinite(z) and (worst is None or abs(z) > abs(worst[2])):
                worst = (k, r, z)
    names = " and ".join(_label(results, k) for k in keys if _analysed(results, k))
    if worst is None:
        return f"Model checks: {names}", "", defs("residual", "normal plot"), ""
    k, r, z = worst
    title = (f"Model checks, {names}: no residual beyond ±3" if abs(z) <= 3 else
             f"Model checks: {r.name[:6]} stands out on {_label(results, k)}")
    return (title, f"Largest standardised residual: run {r.name[:6]} of {cond_text(results, r.condition)} on "
            f"{_label(results, k)}, {z:+.2f}.", defs("residual", "normal plot"),
            "Left: residual vs fitted (a funnel = noise grows with the response). Middle: normal probability "
            "(straight = normal errors). Right: residual vs capture time (a slope = drift).")


def tw_correlations(results):
    rows = tables.response_correlations(results)
    labels = [r["response"] for r in rows]
    pairs = []
    for i, a in enumerate(rows):
        for b in labels[i + 1:]:
            v = a.get(b)
            if v is not None:
                pairs.append((abs(v), v, a["response"], b))
    if not pairs:
        return "How the responses move together", "", defs("r"), ""
    pairs.sort(reverse=True)
    strong = sum(1 for p in pairs if p[0] >= 0.8)
    _, v, a, b = pairs[0]
    return (f"{strong} of {len(pairs)} response pairs correlate at |r| ≥ 0.8",
            f"Strongest: {_short(a)} and {_short(b)} (r = {v:+.2f}). Weakest: {_short(pairs[-1][2])} and "
            f"{_short(pairs[-1][3])} (r = {pairs[-1][1]:+.2f}). Across all runs, so it includes the factors' own "
            f"effects: two responses that both follow gas flow correlate whether or not one causes the other.",
            defs("r"), "")


def tw_pieces(results):
    rows = tables.liquid_pieces(results)
    per = _per_run(rows, "pieces per frame")
    pf = _cond_mean(results, per)
    big = _by_condition_mean(rows, "largest area (mm2)")
    ex = _extremes(pf)
    if ex is None:
        return "Un-atomised liquid pieces", "", defs("pieces"), ""
    (lo, vlo), (hi, vhi) = ex
    bb = max(big, key=big.get) if big else None
    return (f"Most un-atomised pieces at {hi} ({vhi:.0f} per frame), fewest at {lo} ({vlo:.0f})",
            f"Most pieces: {_with_driver(results, per, hi)}. Fewest: {_with_driver(results, per, lo)}."
            + (f" Largest pieces at {cond_text(results, bb)}, {big[bb]:.0f} mm² on average." if bb else ""),
            defs("pieces"), "Many small pieces = liquid breaking up; few large ones = slugs passing through.")


def tw_quality(results):
    flags = {n: f for n, f in tables.quality_flags(results).items() if f}
    rows = tables.data_quality(results)
    moved = [r["moved by sharpness rule (% of in-focus)"] for r in rows
             if r.get("moved by sharpness rule (% of in-focus)") is not None]
    versions = sorted({str(r.sizer_version) for r in results.runs})
    desc = (f"Sizer {', '.join(versions)}." + (f" The sharpness rule moved {min(moved):.0f}–{max(moved):.0f}% of "
                                                f"round droplets to out of focus." if moved else ""))
    cond = {r.name: r.condition for r in results.runs}
    if not flags:
        return "No run stands out on the data-quality checks", desc, defs("robust z", "moved"), ""
    desc += " " + "; ".join(f"{n[:6]} of {cond_text(results, cond[n])}: " + ", ".join(
        f"{x['column']} {x['value']:.3g} (typical {x['typical']:.3g})" for x in fl) for n, fl in sorted(flags.items())) + "."
    return (f"{len(flags)} run(s) stand out on a data-quality check", desc, defs("robust z", "moved"),
            "Dark specks = share of the median background darker than 75% of its surroundings: dirt or a splash "
            "on the window.")


def tw_gallery(results, c: str):
    bits = []
    for r in _conds(results).get(c, []):
        ex = figures_extra._extreme_frames(r, results.thr)
        lo = next((p for w, _, p in ex if w == "lowest"), None)
        hi = next((p for w, _, p in ex if w == "highest"), None)
        bits.append(f"{r.name[:6]}: D32 {_f(r.values.get('d32'))} µm, atomised {_f(r.values.get('atom'))}%"
                    + (f" (frames {lo:.1f}% to {hi:.0f}%)" if lo is not None and hi is not None else ""))
    return (f"{cond_text(results, c)}: least and most atomised frames", "; ".join(bits) + ".",
            defs("atomised"),
            "Brown = un-atomised liquid, circles = droplets (green in focus, magenta out of focus). Left: the "
            "frame with the lowest atomised fraction; right: the highest.")


def caveats(results) -> list[tuple[str, str]]:
    out = []
    n = len(results.multiple)
    if n:
        out.append(("Many tests", f"{n} factor tests: about {n * report.ALPHA:.1f} would reach p < {report.ALPHA} by "
                                  f"chance. Trust a result most when its q (BH) is also below {report.ALPHA}."))
    if not figures_extra.is_full_factorial(results):
        out.append(("Fractional design", "Two-factor interactions are aliased with other factors' main effects. "
                                         "Lack of fit says whether interactions matter, not which."))
    lof = [r.label for r in results.responses.values() if r.anova and r.anova.lack_of_fit
           and r.anova.lack_of_fit.p is not None and r.anova.lack_of_fit.p < report.ALPHA]
    if lof:
        out.append(("Interactions are likely", "Significant lack of fit for: " + ", ".join(_short(l) for l in lof) + "."))
    for f, r in stats.order_confounding(results):
        if abs(r) >= report.CONFOUNDED_R:
            out.append(("Run order", f"{_short(f)} rose with run order (r = {r:+.2f}), so anything that changed slowly "
                                     f"would appear as its effect. No drift was detected; randomising run order "
                                     f"removes the question."))
    quant = [r.label for r in results.responses.values() if r.quantisation_warning]
    if quant:
        out.append(("Pixel lattice", "Few distinct values, so p-values not meaningful: "
                                     + ", ".join(_short(q) for q in quant) + "."))
    if any(str(r.sizer_version) == "2.2.0" for r in results.runs):
        out.append(("Sizer 2.2.0", "D32 is the D32 of round, sharp droplets. Part of any effect on it is selection "
                                   "by the sharpness rule (moved share on the data-quality slide)."))
    out.append(("Absence of evidence", "Not significant is not 'no effect': a small design only detects large "
                                       "effects."))
    return out


# ---- slide building ---------------------------------------------------------------------------------

def _rgb(c):
    from pptx.dml.color import RGBColor
    return RGBColor(*c)


class _Deck:
    def __init__(self, footer: str):
        from pptx import Presentation
        from pptx.util import Emu
        if TEMPLATE.is_file():
            self.prs = Presentation(str(TEMPLATE))
            names = {l.name: l for l in self.prs.slide_layouts}
            self.dark = names.get("TITLE_DARK")
            self.content = names.get("CONTENT")
        else:                                           # template missing: plain default slides
            self.prs = Presentation()
            self.prs.slide_width, self.prs.slide_height = Emu(WIDTH_EMU), Emu(HEIGHT_EMU)
            self.dark = self.content = None
        self.blank = self.prs.slide_layouts[6] if len(self.prs.slide_layouts) > 6 else self.prs.slide_layouts[0]
        self.footer = footer
        self.n = 0

    # -- primitives
    def _text(self, slide, x, y, w, h, text, size, *, bold=False, colour=INK, font=None):
        from pptx.util import Inches, Pt
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = box.text_frame
        tf.word_wrap = True
        for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
            setattr(tf, side, Inches(0.02))
        paras = text if isinstance(text, list) else [text]
        for i, item in enumerate(paras):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            runs = item if isinstance(item, tuple) else ((item, size, bold, colour),)
            for t, sz, b, c in runs:
                run = para.add_run()
                run.text = t
                run.font.size, run.font.bold = Pt(sz), b
                run.font.color.rgb = _rgb(c)
                if font:
                    run.font.name = font
            para.space_after = Pt(4)
        return box

    def _shape(self, slide, kind, x, y, w, h, fill, line=None):
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.util import Inches
        shp = slide.shapes.add_shape({"card": MSO_SHAPE.ROUNDED_RECTANGLE, "dot": MSO_SHAPE.OVAL,
                                      "bar": MSO_SHAPE.RECTANGLE}[kind], Inches(x), Inches(y), Inches(w), Inches(h))
        shp.fill.solid()
        shp.fill.fore_color.rgb = _rgb(fill)
        if line is None:
            shp.line.fill.background()
        else:
            shp.line.color.rgb = _rgb(line)
        shp.shadow.inherit = False
        if kind == "card":
            shp.adjustments[0] = 0.04
        return shp

    def _slide(self, title: str, dark: bool = False):
        layout = (self.dark if dark else self.content) or self.blank
        s = self.prs.slides.add_slide(layout)
        self.n += 1
        ph = next((p for p in s.placeholders if p.placeholder_format.type is not None
                   and p.placeholder_format.idx == 100), None)
        if ph is not None:
            ph.text = title
            run = ph.text_frame.paragraphs[0].runs[0]
            if not dark:
                run.font.size = _pt(24 if len(title) <= 55 else 20 if len(title) <= 80 else 17)
                run.font.bold = True
        else:
            self._text(s, 0.5, 0.3, 9.0, 0.75, title, 22, bold=True, colour=INK, font="Cambria")
        if not dark:
            self._text(s, 0.5, 5.24, 8.3, 0.22, self.footer, 7, colour=MUTED)
            self._text(s, 9.0, 5.22, 0.5, 0.25, str(self.n), 9, colour=MUTED)
        return s

    @staticmethod
    def _card_height(detail, terms, w=9.0, how: str = "") -> float:
        per_line = 15.0 * w                             # ~chars per line at 10.5 pt / 8.5 pt
        lines = math.ceil(len(detail) / per_line) + (math.ceil(len(terms) / (per_line * 1.2)) if terms else 0) \
            + (math.ceil(len(how) / (per_line * 1.1)) if how else 0)
        return 0.17 * max(1, lines) + 0.18 + (0.04 if terms else 0) + (0.04 if how else 0)

    def _side_terms(self, slide, x, y, w, h, terms, accent, how: str = "") -> None:
        """Card beside the figure (as the hand-made deck's 'How to read it' card): how to read the figure,
        then the terms it uses."""
        self._shape(slide, "card", x, y, w, h, CARD)
        self._shape(slide, "dot", x + 0.15, y + 0.14, 0.22, 0.22, accent)
        items = [t.strip() for t in terms.split("  ·  ") if t.strip()]
        size = 8.5 if len(terms) + len(how) < 650 else 7.5
        self._text(slide, x + 0.45, y + 0.1, w - 0.55, 0.3,
                   [(("How to read it" if how else "Terms", 11, True, INK),)], 11)
        paras = []
        if how:
            paras.append(((how, size, False, BODY),))
            if items:
                paras.append((("Terms", size + 1, True, INK),))
        for t in items:
            head, _, rest = t.partition(": ")
            paras.append(((head + ": ", size, True, BODY), (rest, size, False, BODY)) if rest else ((t, size, False, BODY),))
        self._text(slide, x + 0.15, y + 0.45, w - 0.27, h - 0.55, paras, size)

    def _card_text(self, slide, x, y, w, detail, terms, accent, how: str = "") -> float:
        """Description card with an accent bar; returns its bottom (inches)."""
        h = self._card_height(detail, terms, w, how)
        self._shape(slide, "card", x, y, w, h, CARD)
        self._shape(slide, "bar", x, y + 0.06, 0.06, h - 0.12, accent)
        paras = [((detail, 10.5, False, BODY),)]
        if how:
            paras.append((("How to read it  ", 9, True, BODY), (how, 9, False, BODY)))
        if terms:
            paras.append((("Terms  ", 8.5, True, MUTED), (terms, 8.5, False, MUTED)))
        self._text(slide, x + 0.16, y + 0.06, w - 0.25, h - 0.1, paras, 10.5)
        return y + h

    # -- slide kinds
    def title(self, heading: str, sub: list[str]) -> None:
        s = self.prs.slides.add_slide(self.dark or self.blank)
        self.n += 1
        phs = {p.placeholder_format.idx: p for p in s.placeholders}
        for i, c in enumerate((ORANGE, TEAL, PURPLE)):
            self._shape(s, "dot", 0.6 + 0.6 * i, 0.75, 0.42, 0.42, c)
        if 100 in phs:
            phs[100].text = heading
        else:
            self._text(s, 0.6, 1.55, 8.8, 1.5, heading, 38, bold=True, colour=WHITE, font="Cambria")
        lines = [x for x in sub if x]
        if 101 in phs:
            tf = phs[101].text_frame
            tf.text = lines[0] if lines else ""
            for line in lines[1:]:
                tf.add_paragraph().text = line
        else:
            self._text(s, 0.6, 3.2, 8.8, 1.2, lines, 16, colour=WHITE)
        self._text(s, 0.6, 5.0, 8.8, 0.3, self.footer, 8, colour=GREY)

    def section(self, heading: str, sub: str = "") -> None:
        self.title(heading, [sub] if sub else [])

    def cards(self, title: str, items: list[tuple[str, str]], accent=ORANGE) -> None:
        """2-column grid of rounded cards, each with an accent dot, a heading and a body."""
        s = self._slide(title)
        n = len(items)
        rows = max(1, math.ceil(n / 2))
        top, gap = 1.2, 0.14
        h = (5.1 - top - gap * (rows - 1)) / rows
        for i, (head, body) in enumerate(items):
            x = 0.5 + (i % 2) * 4.6
            y = top + (i // 2) * (h + gap)
            self._shape(s, "card", x, y, 4.4, h, CARD)
            self._shape(s, "dot", x + 0.2, y + 0.18, 0.32, 0.32, accent)
            longest = max(len(b) for _, b in items)
            fit = h / max(1.0, longest / 52 * 0.2 + 0.35)   # ~52 chars per line at 10.5 pt in 3.6 in
            size = 10.5 if fit >= 1 else 9.5 if fit >= 0.85 else 8.5
            self._text(s, x + 0.65, y + 0.12, 3.6, h - 0.2,
                       [((head, 13, True, INK),), ((body, size, False, BODY),)], 11)

    def figure(self, spec: tuple, image: Path, accent=ORANGE) -> bool:
        from PIL import Image
        from pptx.util import Inches
        if not Path(image).is_file():
            return False
        title, detail, terms, notes, *more = spec
        how = more[0] if more else ""
        detail = detail or _describe(image)
        s = self._slide(title)
        with Image.open(image) as im:
            w, h = im.size
        # how-to-read and terms either sit in the card above (full-width figure) or in a card beside
        # the figure; whichever leaves the figure bigger
        side_w, gap = 2.45, 0.15
        top_a = 1.12 + self._card_height(detail, terms, 9.0, how) + 0.08
        top_b = 1.12 + self._card_height(detail, "") + 0.08
        scale_a = min(9.0 / w, (5.2 - top_a) / h)
        scale_b = min((9.0 - side_w - gap) / w, (5.2 - top_b) / h)
        if (terms or how) and scale_b > scale_a:
            top = self._card_text(s, 0.5, 1.12, 9.0, detail, "", accent) + 0.08
            pw, ph = w * scale_b, h * scale_b
            box_w = 9.0 - side_w - gap
            s.shapes.add_picture(str(image), Inches(0.5 + (box_w - pw) / 2), Inches(top), Inches(pw), Inches(ph))
            self._side_terms(s, 0.5 + box_w + gap, top, side_w, 5.2 - top, terms, accent, how)
        else:
            top = self._card_text(s, 0.5, 1.12, 9.0, detail, terms, accent, how) + 0.08
            pw, ph = w * scale_a, h * scale_a
            s.shapes.add_picture(str(image), Inches(0.5 + (9.0 - pw) / 2), Inches(top), Inches(pw), Inches(ph))
        s.notes_slide.notes_text_frame.text = "\n\n".join(t for t in (detail, how, notes, terms) if t)
        return True

    def table(self, title: str, header: list[str], rows: list[list], detail: str = "", terms: str = "",
              accent=TEAL) -> None:
        from pptx.util import Inches, Pt
        s = self._slide(title)
        top = (self._card_text(s, 0.5, 1.12, 9.0, detail, terms, accent) + 0.1) if (detail or terms) else 1.15
        n_rows = len(rows) + 1
        height = min(5.15 - top, 0.3 * n_rows)
        tbl = s.shapes.add_table(n_rows, len(header), Inches(0.5), Inches(top), Inches(9.0), Inches(height)).table
        size = 11 if n_rows <= 11 and len(header) <= 6 else 8
        for i in range(n_rows):
            for j in range(len(header)):
                cell = tbl.cell(i, j)
                cell.text = str(header[j] if i == 0 else rows[i - 1][j])
                cell.fill.solid()
                cell.fill.fore_color.rgb = _rgb(NAVY if i == 0 else (WHITE if i % 2 else CARD))
                for para in cell.text_frame.paragraphs:
                    for run in para.runs:
                        run.font.size = Pt(size)
                        run.font.bold = i == 0
                        run.font.color.rgb = _rgb(WHITE if i == 0 else INK)

    def save(self, path: Path) -> Path:
        path = Path(path)
        tmp = path.with_name(f".{path.stem}.{os.getpid()}.tmp.pptx")
        try:
            self.prs.save(str(tmp))
            try:
                fsops.replace(tmp, path)
            except PermissionError:
                raise RuntimeError(f"cannot replace {path.name}: close it in PowerPoint and analyse again") from None
        finally:
            tmp.unlink(missing_ok=True)
        return path


def _pt(x):
    from pptx.util import Pt
    return Pt(x)


def _chunks(seq: list, n: int) -> list[list]:
    return [seq[i:i + n] for i in range(0, len(seq), n)]


def _finding_card(line: str) -> tuple[str, str]:
    text = _plain(line)
    if ": " in text:
        head, body = text.split(": ", 1)
        return head, body[:1].upper() + body[1:]
    return "", text


def build(path: Path, results, bins, prov, figures_dir: Path, title: str = "Taguchi analysis") -> Path:
    """Write the deck to `path` from an analysed (not refused) Results."""
    figs = Path(figures_dir)
    parts = figs / SLIDE_FIG_DIR
    d = _Deck(f"{title} · generated {prov.generated} · revision {prov.revision}")

    conds = _conds(results)
    reps = sorted({len(v) for v in conds.values()})
    shape = (f"{len(results.runs)} runs: {len(conds)} conditions × {reps[0]} replicates" if len(reps) == 1
             else f"{len(results.runs)} runs in {len(conds)} conditions ({reps[0]}–{reps[-1]} replicates)")
    dates = sorted({(r.metrics or {}).get("captured", "")[:10] for r in results.runs} - {""})
    d.title(title, [shape + " · factors: " + ", ".join(f.label for f in results.factors),
                    (("Captured " + ", ".join(dates) + " · ") if dates else "")
                    + f"sizer {', '.join(sorted({str(r.sizer_version) for r in results.runs}))} · "
                      f"error term: {results.error_source}"])
    d.table(f"The design: {shape}", ["Condition"] + [f.label for f in results.factors] + ["Runs"],
            [[c] + [tables.level_text(rs[0].levels[f.key]) for f in results.factors] + [len(rs)]
             for c, rs in conds.items()],
            detail=f"Each condition was run {reps[0]} times" + ("" if len(reps) == 1 else " or more")
                   + f"; the error term is {results.error_source}. Runs left out: {len(results.skipped)}.",
            terms=defs("pure error"), accent=NAVY)
    found = [_finding_card(l) for l in report._findings(results, bins) if not l.startswith("  ")]
    if found:
        d.cards("What was found", found, ORANGE)

    # the effects
    d.figure(tw_main_effects(results), figs / "fig_main_effects.png", ORANGE)
    d.figure(tw_contributions(results), figs / "fig_contributions.png", ORANGE)
    d.figure(tw_per_run(results), figs / "fig_per_run.png", ORANGE)
    d.figure(tw_sn(results), figs / "fig_sn.png", ORANGE)
    d.figure(tw_glr(results), figs / "fig_glr.png", ORANGE)

    # droplet size
    d.figure(tw_spread_count(results, bins), figs / "fig_size_spread_count.png", TEAL)
    d.figure(tw_spread_volume(results, bins), figs / "fig_size_spread_volume.png", TEAL)
    d.figure(tw_size_distribution(results), figs / "fig_size_distribution.png", TEAL)
    d.figure(tw_size_percentiles(results), figs / "fig_size_percentiles.png", TEAL)

    # steadiness
    d.figure(tw_steadiness(results), figs / "fig_steadiness.png", PURPLE)
    d.figure(tw_spectra(results), figs / "fig_spectra.png", PURPLE)
    groups = _chunks(list(conds), SERIES_PER_SLIDE)
    for i, part in enumerate(groups, 1):
        files = figures_extra.fig_series(results, parts, only=part, name=f"fig_series_{i}", formats=("png",))
        if files:
            d.figure(tw_series(results, part), files[0], PURPLE)
    for i, part in enumerate(groups, 1):
        files = figures_extra.fig_series_liquid(results, parts, only=part, name=f"fig_series_liquid_{i}",
                                                formats=("png",))
        if files:
            d.figure(tw_series_liquid(results, part), files[0], PURPLE)

    # drift, replicates, interactions
    d.figure(tw_drift(results), figs / "fig_drift.png", GREY)
    analysed = [k for k in results.responses if _analysed(results, k)]
    size_keys = [k for k in analysed if k in SIZE_KEYS]
    other_keys = [k for k in analysed if k not in SIZE_KEYS]
    for i, (keys, what) in enumerate(((size_keys, "droplet size"), (other_keys[:5], "atomisation and steadiness"),
                                      (other_keys[5:], "pulsing and pressure")), 1):
        if keys:
            files = figures_extra.fig_replicates(results, parts, keys=keys, name=f"fig_replicates_{i}",
                                                 formats=("png",))
            if files:
                d.figure(tw_replicates(results, keys, what), files[0], GREY)
    d.figure(tw_interactions(results), figs / "fig_interactions.png", GREY)

    # model checks
    for i, keys in enumerate(_chunks(list(figures_extra.DIAG_KEYS), 2), 1):
        files = figures_extra.fig_diagnostics(results, parts, keys=keys, name=f"fig_diagnostics_{i}",
                                              formats=("png",))
        if files:
            d.figure(tw_diagnostics(results, keys), files[0], GREY)
    d.figure(tw_correlations(results), figs / "fig_correlations.png", GREY)

    # liquid pieces, data quality
    d.figure(tw_pieces(results), figs / "fig_pieces.png", TEAL)
    d.figure(tw_quality(results), figs / "fig_quality.png", YELLOW)

    keys = [k for k in ("d32", "atom", "dv50", "af_flow", "lag1") if _analysed(results, k)]
    if keys:
        cell = {(r["condition"], r["response"]): r for r in tables.condition_summary(results)}
        rows = []
        for c, rs in conds.items():
            line = [c] + [tables.level_text(rs[0].levels[f.key]) for f in results.factors]
            for k in keys:
                r = cell.get((c, results.responses[k].label))
                line.append("–" if r is None or r["mean"] is None else
                            report.fmt(r["mean"], 4) if r["CI high"] is None else
                            f"{report.fmt(r['mean'], 4)} ± {report.fmt(r['CI high'] - r['mean'], 2)}")
            rows.append(line)
        d.table("Per-condition summary", ["Cond."] + [_short(f.label) for f in results.factors]
                + [_label(results, k) for k in keys], rows,
                detail="Mean ± half-width of the 95% CI from the replicates. Every response, with its repeat SD, "
                       "is in the workbook.", terms=defs("t CI", "SD", "Dv", "lag-1"), accent=NAVY)

    d.cards("Caveats before reading too much into this", caveats(results), YELLOW)

    # appendix: images
    gal = [(c, figs / figures_extra.GALLERY_DIR / f"gallery_{c}.png") for c in conds]
    gal = [(c, p) for c, p in gal if p.is_file()]
    if gal:
        d.section("Appendix: images", "Each run's least and most atomised frame, one slide per condition.")
        for c, p in gal:
            d.figure(tw_gallery(results, c), p, GREY)
    return d.save(path)
