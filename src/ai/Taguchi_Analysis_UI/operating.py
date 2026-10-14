"""operating.py -- what each run actually got: the gas delivered and the upstream pressure inside the
camera window (run_summary.xlsx -> Pressure: timestamps, pressures, flows_sccm, cam_start_1, cam_end_1),
and from them its operating REGIME. No Qt.

Why (the CCD day, 2026-10-12): every 0.5 mm run passed its full set flow at 6-8 barA until the silicone
reached the orifice; then two-phase choking pushed the pressure to the supply regulator's ceiling
(~9.2-9.6 barA) and the gas fell to whatever the orifice could pass at that pressure (33-5,554 sccm against
set points of 7,000-11,000). From then on the rig is PRESSURE-LIMITED: the gas set point is no longer an
input, the delivered gas is the atomiser's two-phase capacity, set by feed and rpm. Being pressure-limited
is not bad in itself (C7 / C8, 0.5 mm, feed 3,000, 900 rpm: ~5,500 sccm and 35% atomised, steady and
reproducible), but a fit that believes those runs got their set point is wrong.

    flow-controlled    delivered >= 95% of the set point
    pressure-limited   delivered < 95% and P0 within NEAR_CEILING_BAR of the supply ceiling
    gas short          delivered < 95% and the pressure well below the ceiling (cause unknown: Alicat,
                       supply, a leak); as with pressure-limited, the set point does not describe the run
    unknown            no pressure log or no camera window

`settled`: the camera-window median gas flow within SETTLED_PCT (or SETTLED_SCCM) of the median over the
10 s before the window, and the window's second-half median within the same of its first half. Medians,
because a pressure-limited run pulses (slugs): a straight-line trend across 13 s of pulsing read 10-18% on
runs whose 10 s medians were steady to 2%. A run still falling or recovering when the camera ran (C1, the
first run of the day: 967 sccm before the window, 2,170 in it) is flagged, not dropped.
"""
from __future__ import annotations

import math
from dataclasses import dataclass, field
from pathlib import Path

import numpy as np

CEILING_BAR = 9.5            # supply regulator setting upstream of the Alicat (Ben, 2026-10-13: set ~9.5 for
                             # safety; the rig itself is limited to 10 bar)
DELIVERED_SHARE = 0.95       # delivered / set below this: the set point was not delivered
NEAR_CEILING_BAR = 0.5       # P0 this close to the ceiling (or above): pressure-limited
SETTLED_PCT = 10.0
SETTLED_SCCM = 100.0         # ... or this many sccm, whichever is larger (a dribble's 30-100 sccm is noise)
PRE_WINDOW_S = 10.0
TRACE_POINTS = 400           # the whole-run trace kept for figures is decimated to about this many samples
STALL_S = 0.75               # a step this long in the ~0.22 s log: GUI_Clean's GUI thread stalled (see respace_queued)
BURST_S = 0.05               # readings logged closer than this after a stall were queued during it

FLOW_CONTROLLED, PRESSURE_LIMITED, GAS_SHORT, UNKNOWN = ("flow-controlled", "pressure-limited", "gas short",
                                                        "unknown")
NOT_DELIVERED = (PRESSURE_LIMITED, GAS_SHORT)


def _num(x):
    try:
        v = float(x)
        return v if math.isfinite(v) else None
    except (TypeError, ValueError):
        return None


@dataclass
class Trace:
    t: np.ndarray                 # s since Start, unique and increasing
    p: np.ndarray                 # barA
    q: np.ndarray                 # sccm (0 where not logged)
    cam: tuple | None             # (start, end) of the camera window, s since Start


def respace_queued(t: np.ndarray) -> np.ndarray:
    """Times for readings that were QUEUED while GUI_Clean's GUI thread stalled.

    Before 2026-10-14 the GUI stamped each Alicat reading when its GUI thread handled it, not when it was
    read. On the CCD day that thread stalled for 1-4 s in 30 of 33 runs (the LaCie busy, usually just after
    the previous run's cine finished saving), the reader kept reading every ~0.2 s, and the readings queued
    meanwhile were logged in one burst at a single instant: a gap, then a jump. Unique timestamps alone kept
    one reading of each burst and dropped the rest. Here a burst that follows a stall is spread evenly back
    across the stall, which is when those readings were taken. Logs from the fixed GUI have no such bursts."""
    t = np.asarray(t, float).copy()
    i = 0
    while i < t.size - 1:
        if t[i + 1] - t[i] > STALL_S:
            j = i + 1
            while j + 1 < t.size and t[j + 1] - t[j] < BURST_S:
                j += 1
            n = j - i                                   # readings i+1 .. j were queued during the stall
            if n > 1:
                t[i + 1:j + 1] = t[i] + (t[j] - t[i]) * np.arange(1, n + 1) / n
            i = j
        else:
            i += 1
    return t


def read_trace(xlsx: Path) -> Trace | None:
    """The run's pressure / gas log, or None if the workbook or its Pressure sheet is missing."""
    import openpyxl
    xlsx = Path(xlsx)
    if not xlsx.exists():
        return None
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    try:
        if "Pressure" not in wb.sheetnames:
            return None
        rows = [r for r in wb["Pressure"].iter_rows(min_row=2, values_only=True) if r and r[0] is not None]
    finally:
        wb.close()
    if not rows:
        return None
    t = np.array([_num(r[0]) for r in rows], float)
    p = np.array([_num(r[1]) if len(r) > 1 else None for r in rows], float)
    q = np.array([(_num(r[2]) or 0.0) if len(r) > 2 else 0.0 for r in rows], float)
    t0 = next((_num(r[3]) for r in rows if len(r) > 3 and _num(r[3]) is not None), None)
    t1 = next((_num(r[4]) for r in rows if len(r) > 4 and _num(r[4]) is not None), None)
    ok = np.isfinite(t)
    t, p, q = respace_queued(t[ok]), p[ok], q[ok]
    _, iu = np.unique(t, return_index=True)
    return Trace(t[iu], p[iu], q[iu], (t0, t1) if t0 is not None and t1 is not None else None)


def window_stats(tr: Trace) -> dict:
    """Camera-window statistics of the gas flow and pressure (empty if there is no window)."""
    if tr.cam is None:
        return {}
    t0, t1 = tr.cam
    m = (tr.t >= t0) & (tr.t <= t1)
    if m.sum() < 3:
        return {}
    q, p, t = tr.q[m], tr.p[m], tr.t[m]
    out = {"delivered_sccm": float(np.median(q)), "delivered_p10": float(np.percentile(q, 10)),
           "delivered_p90": float(np.percentile(q, 90)),
           "p0_bar": float(np.nanmedian(p)) if np.isfinite(p).any() else None}
    pre = (tr.t >= t0 - PRE_WINDOW_S) & (tr.t < t0)
    out["pre_sccm"] = float(np.median(tr.q[pre])) if pre.sum() >= 3 else None
    mid = (t0 + t1) / 2
    a, b = q[t < mid], q[t >= mid]
    out["half_change_sccm"] = float(np.median(b) - np.median(a)) if a.size >= 2 and b.size >= 2 else None
    return out


@dataclass
class OperatingPoint:
    set_sccm: float | None
    delivered_sccm: float | None = None
    delivered_p10: float | None = None
    delivered_p90: float | None = None
    p0_bar: float | None = None
    pre_sccm: float | None = None
    half_change_sccm: float | None = None      # second-half minus first-half median in the window
    regime: str = UNKNOWN
    settled: bool | None = None
    ceiling_bar: float = CEILING_BAR
    note: str = ""
    trace: dict = field(default_factory=dict, repr=False)    # whole run, decimated: t, p, q, cam

    @property
    def not_delivered(self) -> bool:
        return self.regime in NOT_DELIVERED

    @property
    def delivered_share(self) -> float | None:
        if self.set_sccm and self.delivered_sccm is not None:
            return self.delivered_sccm / self.set_sccm
        return None

    def text(self) -> str:
        """'11,000 -> 33 sccm at 9.56 barA: pressure-limited' (for tables and tooltips)."""
        if self.delivered_sccm is None:
            return self.note or "no gas log in the camera window"
        s = f"{self.set_sccm:,.0f} → " if self.set_sccm else ""
        s += f"{self.delivered_sccm:,.0f} sccm"
        if self.p0_bar is not None:
            s += f" at {self.p0_bar:.2f} barA"
        s += f": {self.regime}"
        if self.settled is False:
            s += ", not settled"
        return s


_STATS = ("delivered_sccm", "delivered_p10", "delivered_p90", "p0_bar", "pre_sccm", "half_change_sccm")


def classify(stats: dict, set_sccm: float | None, ceiling_bar: float = CEILING_BAR, note: str = "",
             trace: dict | None = None) -> OperatingPoint:
    op = OperatingPoint(set_sccm, ceiling_bar=ceiling_bar, note=note, trace=trace or {})
    for k in _STATS:
        setattr(op, k, stats.get(k))
    if op.delivered_sccm is None:
        op.note = op.note or "no gas log in the camera window"
        return op
    if op.pre_sccm is not None and op.half_change_sccm is not None:
        tol = max(SETTLED_PCT / 100 * op.delivered_sccm, SETTLED_SCCM)
        op.settled = abs(op.delivered_sccm - op.pre_sccm) <= tol and abs(op.half_change_sccm) <= tol
    if not set_sccm:
        op.note = "no gas set point to compare with"
        return op
    if op.delivered_sccm >= DELIVERED_SHARE * set_sccm:
        op.regime = FLOW_CONTROLLED
    elif op.p0_bar is not None and op.p0_bar >= ceiling_bar - NEAR_CEILING_BAR:
        op.regime = PRESSURE_LIMITED
    else:
        op.regime = GAS_SHORT
    return op


def _decimate(tr: Trace) -> dict:
    step = max(1, int(math.ceil(tr.t.size / TRACE_POINTS)))
    clean = lambda a: [None if not math.isfinite(v) else round(float(v), 4) for v in a[::step]]
    return {"t": clean(tr.t), "p": clean(tr.p), "q": clean(tr.q), "cam": list(tr.cam) if tr.cam else None}


def operating_point(xlsx: Path, set_sccm: float | None, ceiling_bar: float = CEILING_BAR) -> OperatingPoint:
    """Read a run's log and classify it. Never raises: a problem is in `note`, the regime 'unknown'."""
    try:
        tr = read_trace(xlsx)
    except Exception as exc:                                  # noqa: BLE001 -- reported, never fatal
        return OperatingPoint(set_sccm, ceiling_bar=ceiling_bar, note=f"pressure log unreadable: {exc}")
    if tr is None:
        return OperatingPoint(set_sccm, ceiling_bar=ceiling_bar, note="no pressure log in run_summary.xlsx")
    note = "" if tr.cam else "no camera window recorded in the pressure log"
    return classify(window_stats(tr), set_sccm, ceiling_bar, note, _decimate(tr))


def reclassify(op: OperatingPoint, set_sccm: float | None, ceiling_bar: float | None = None) -> OperatingPoint:
    """The same log against another set point (a level corrected on the Design page) or ceiling."""
    stats = {k: getattr(op, k) for k in _STATS}
    note = op.note if op.delivered_sccm is None else ""
    return classify(stats, set_sccm, op.ceiling_bar if ceiling_bar is None else ceiling_bar, note, op.trace)
