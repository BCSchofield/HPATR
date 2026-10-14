"""Phase 5: the response-surface outputs on a SIMULATED 33-run CCD day: Ben's run sheet (rsm_design),
captured as measured fake runs with planted surfaces (fakes.ccd_truth: atomised % peaks inside the gas
range and rises with the orifice; no rotation costs 2 points), analysed the way the tab will
(rsm_publish with the orifice restricted to the inserts that exist)."""
from __future__ import annotations

import tempfile
from pathlib import Path

import pytest

from src.ai.Taguchi_Analysis_UI import rsm_design as rb
from src.ai.Taguchi_Analysis_UI import rsm_detect as rsm
from src.ai.Taguchi_Analysis_UI import rsm_figures as rf
from src.ai.Taguchi_Analysis_UI import rsm_publish as rp
from src.ai.Taguchi_Analysis_UI import rsm_report
from src.ai.Taguchi_Analysis_UI import run_discovery as rd
from src.ai.Taguchi_Analysis_UI import tables
from src.ai.Taguchi_Analysis_UI.tests import fakes
from src.ai.Taguchi_Analysis_UI.tests.test_rsm_design import ben_plan

INSERTS = [0.5, 0.8, 1.2, 1.8]


@pytest.fixture(scope="module")
def day():
    with tempfile.TemporaryDirectory() as d:
        root = Path(d)
        fakes.make_measured_ccd_day(root / "src", rb.build_sheet(ben_plan()), seed=3)
        design = rsm.from_runs(rd.load_runs(rd.find_runs([root / "src" / "2026"]).runs, check_reuse=False))
        keys = [f.key for f in design.detect().factors]
        res = rp.publish(design, root / "out", [rp.ModelSpec("full quadratic", keys)],
                         discrete={"orifice": INSERTS}, title="simulated CCD day")
        yield {"res": res, "out": root / "out", "R": next(iter(res.models.values()))}


def test_the_day_publishes_with_every_output(day):
    res = day["res"]
    assert res.ok, res.problems
    assert res.detected.type_key == "ccf" and res.detected.k == 4
    stems = {p.stem for p in res.figures}
    assert {"fig_rsm_contour_atom_rpm100", "fig_rsm_contour_atom_rpm500", "fig_rsm_contour_atom_rpm900",
            "fig_rsm_pairs_atom", "fig_rsm_surface_atom", "fig_rsm_drift", "fig_rsm_validation",
            "fig_rsm_desirability_rpm500", "fig_per_run", "fig_size_spread_volume"} <= stems


def test_the_maps_follow_the_ccd_plan_layout(day):
    L = rf.map_layout(day["R"])
    assert L.pair == ("sccm", "sps") and L.panel == "orifice" and L.figure == "rpm" and L.held == {}


def test_the_planted_atomisation_surface_and_optimum_are_found(day):
    R = day["R"]
    f = R.fits["atom"].full
    c = f.coef()
    assert c["sccm"] > 2 and c["sccm^2"] < -1 and c["sps"] < -1 and c["orifice"] > 0.5
    assert f.r2_pred > 0.9
    # atomised fraction alone: d/dg (3g - 2g^2 + 0.6 g o) = 0 with o near +1 -> g ~ 0.9 -> ~10,800 sccm
    s = rsm_report.stationary_points(R, "atom")
    assert s is not None


def test_the_no_rotation_control_is_caught_and_the_check_run_passes(day):
    ctrl = [c for c in rsm_report.controls(day["R"]) if c["response"].startswith("Atomised fraction")]
    assert len(ctrl) == 1 and ctrl[0]["difference in centre SDs"] < -3          # planted: -2 points
    preds = {p.run.point.kind: p for p in day["R"].fits["atom"].predictions}
    assert preds["check"].inside and not preds["check"].extrapolated
    assert preds["outside"].extrapolated and preds["outside"].inside is False


def test_the_settable_optimum_uses_an_insert_that_exists(day):
    R = day["R"]
    od = R.optimum_discrete
    assert od is not None and round(R.decode("orifice", od.coded["orifice"]), 3) in INSERTS
    nxt = rsm_report.tables(day["res"].models)["next_runs"]
    assert nxt and nxt[0]["Orifice (mm)"] in INSERTS


def test_weak_surfaces_are_kept_out_of_the_best_settings_search(day):
    """Goals named explicitly: the default goals (in-band rate, atomised fraction; 2026-10-14) have no weak
    surface on the fakes, which carry no liquid volumes and so no rate."""
    from src.ai.Taguchi_Analysis_UI import rsm_responses as rr
    R = rr.analyse(day["R"].loaded, discrete={"orifice": INSERTS},
                   goals=[{"key": k, "direction": d} for k, d in (("atom", "max"), ("cv_droplets", "min"),
                                                                  ("d32", "min"))])
    weak = [k for k in ("cv_droplets", "d32") if R.fits[k].r2_pred_used < 0.5]
    assert weak
    assert {g.key for g in R.goals}.isdisjoint(weak)
    assert any("left out of the best-settings search" in w for w in R.warnings)


def test_ccd_conditions_are_grouped_by_setting_in_natural_order(day):
    assert tables.cond_order("C2") < tables.cond_order("C10") < tables.cond_order("T1")
    from pptx import Presentation
    prs = Presentation(str(day["out"] / "rsm_deck.pptx"))
    titles = [next((sh.text_frame.text for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()), "")
              for s in prs.slides]
    series = [t for t in titles if t.startswith("Droplets per frame through each capture")]
    assert series and series[0].endswith("C1, C2, C3")
    assert not any("conflict" in w.lower() for w in day["res"].models["full quadratic"].warnings)


def test_every_surface_slide_says_what_it_shows_and_how_to_read_it(day):
    """Contour maps, 3D surfaces, all-pairs maps and desirability maps: each slide carries 'What it shows'
    and a 'How to read it' guide (Ben, 2026-10-11)."""
    from pptx import Presentation
    prs = Presentation(str(day["out"] / "rsm_deck.pptx"))
    surfaces = 0
    for s in prs.slides:
        pics = [sh for sh in s.shapes if sh.shape_type == 13]
        texts = [sh.text_frame.text for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
        if not pics or not texts:
            continue
        name = pics[0].image.filename if hasattr(pics[0].image, "filename") else ""
        title = texts[0]
        is_surface = (" in 3D" in title or "every pair of factors" in title or title.startswith("Where the goals")
                      or any(p in title for p in (": up with", ": down with", ": no usable surface", ": no clear slope")))
        if not is_surface:
            continue
        body = " ".join(texts)
        assert "What it shows" in body and "How to read it" in body, title
        surfaces += 1
    assert surfaces >= 6 * 3 + 6 + 3        # maps (3 RPM each), 3D + pairs per response, desirability per RPM


def test_the_publish_carries_the_recommendations(day):
    out = day["out"]
    assert {"next_runs.xlsx", "next_runs.csv", "next_runs.md"} <= {p.name for p in out.iterdir()}
    text = (out / "rsm_report.md").read_text(encoding="utf-8")
    assert "## Recommended next runs" in text and "### Confirm the best setting" in text
    from pptx import Presentation
    titles = [next((sh.text_frame.text for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()), "")
              for s in Presentation(str(out / "rsm_deck.pptx")).slides]
    assert any(t.startswith("Next: Confirm the best setting") for t in titles)
