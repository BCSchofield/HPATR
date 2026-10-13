"""Phase 3 of the CCD follow-up (2026-10-13): what each run actually got (operating.py) and what the
response-surface analysis does with runs that did not get their set gas.

The real case: on 2026-10-12 every 0.5 mm run hit the ~9.5 barA supply ceiling once the silicone arrived,
and its gas fell to 33-5,554 sccm against set points of 7,000-11,000. Here the simulated CCD day's ten
0.5 mm runs get a log like that; every other run gets its set point at a few barA."""
from __future__ import annotations

import csv
import tempfile
from pathlib import Path

import numpy as np
import openpyxl
import pytest

from src.ai.Taguchi_Analysis_UI import operating as op
from src.ai.Taguchi_Analysis_UI import rsm_design as rb
from src.ai.Taguchi_Analysis_UI import rsm_detect as rsm
from src.ai.Taguchi_Analysis_UI import rsm_publish as rp
from src.ai.Taguchi_Analysis_UI import rsm_recommend
from src.ai.Taguchi_Analysis_UI import run_discovery as rd
from src.ai.Taguchi_Analysis_UI.tests import fakes
from src.ai.Taguchi_Analysis_UI.tests.test_rsm_design import ben_plan


def add_pressure_log(run: Path, set_sccm: float, delivered: float, p0: float, *, pre: float | None = None,
                     cam=(60.0, 73.0), seed: int = 0) -> None:
    """A Pressure sheet like GUI_Clean's: 4 Hz samples to 85 s, the gas at its set point until 'the silicone
    arrives' at 20 s, then `delivered` at `p0` (with a little noise); cam_start_1 / cam_end_1 on the rows
    nearest the camera window. `pre`: the gas in the 10 s before the window, if different."""
    rng = np.random.default_rng(seed)
    wb = openpyxl.load_workbook(run / "run_summary.xlsx")
    if "Pressure" in wb.sheetnames:                     # the fakes' own short log has no camera window
        del wb["Pressure"]
    ws = wb.create_sheet("Pressure")
    ws.append(["timestamps", "pressures", "flows_sccm", "cam_start_1", "cam_end_1"])
    t = np.arange(0.25, 85, 0.25)
    i0, i1 = int(np.argmin(abs(t - cam[0]))), int(np.argmin(abs(t - cam[1])))
    for i, ti in enumerate(t):
        if ti < 20:
            q, p = set_sccm, 2.0 + p0 * ti / 40
        elif pre is not None and cam[0] - 10 <= ti < cam[0]:
            q, p = pre, p0
        else:
            q, p = delivered, p0
        q = max(0.0, q * (1 + rng.normal(0, 0.01)))
        ws.append([float(ti), round(float(p + rng.normal(0, 0.01)), 3), round(q),
                   float(ti) if i == i0 else None, float(ti) if i == i1 else None])
    wb.save(run / "run_summary.xlsx")


# ---- classification -----------------------------------------------------------------------------------------

def test_the_three_regimes_and_unknown():
    fc = op.classify({"delivered_sccm": 8996, "p0_bar": 1.59, "pre_sccm": 9013, "half_change_sccm": 6}, 9000)
    assert fc.regime == op.FLOW_CONTROLLED and fc.settled and not fc.not_delivered
    pl = op.classify({"delivered_sccm": 5370, "p0_bar": 9.18, "pre_sccm": 5609, "half_change_sccm": -38}, 7000)
    assert pl.regime == op.PRESSURE_LIMITED and pl.not_delivered and pl.settled
    assert pl.text() == "7,000 → 5,370 sccm at 9.18 barA: pressure-limited"
    gs = op.classify({"delivered_sccm": 5000, "p0_bar": 3.0, "pre_sccm": 5000, "half_change_sccm": 0}, 9000)
    assert gs.regime == op.GAS_SHORT and gs.not_delivered
    unk = op.classify({}, 9000)
    assert unk.regime == op.UNKNOWN and not unk.not_delivered and "no gas log" in unk.text()
    # 95% of set is still flow-controlled; just under is not
    assert op.classify({"delivered_sccm": 8550, "p0_bar": 9.4}, 9000).regime == op.FLOW_CONTROLLED
    assert op.classify({"delivered_sccm": 8540, "p0_bar": 9.4}, 9000).regime == op.PRESSURE_LIMITED
    # a lower ceiling setting moves the line
    assert op.classify({"delivered_sccm": 5000, "p0_bar": 7.6}, 9000, ceiling_bar=8.0).regime == op.PRESSURE_LIMITED


def test_settled_uses_medians_and_a_floor():
    # C1 on the day: 967 sccm before the window, 2,170 in it and still climbing
    c1 = op.classify({"delivered_sccm": 2170, "p0_bar": 9.35, "pre_sccm": 967, "half_change_sccm": 1584}, 11000)
    assert c1.settled is False and c1.text().endswith("not settled")
    # a dribble's 34 against 96 sccm is within the 100 sccm floor
    c3 = op.classify({"delivered_sccm": 34, "p0_bar": 9.56, "pre_sccm": 96, "half_change_sccm": -3}, 11000)
    assert c3.settled is True


def test_a_log_is_read_and_classified(tmp_path):
    run = fakes.make_run(tmp_path, "105301_11000sccm_900rpm_5000sps_or0.5_bh0.5", day=("2026", "10", "12"))
    add_pressure_log(run, 11000, 2200, 9.35, pre=1000)
    o = op.operating_point(run / "run_summary.xlsx", 11000)
    assert o.regime == op.PRESSURE_LIMITED and abs(o.delivered_sccm - 2200) < 60 and abs(o.p0_bar - 9.35) < 0.05
    assert o.settled is False and abs(o.pre_sccm - 1000) < 40
    assert o.trace["cam"] == pytest.approx([60.0, 73.0], abs=0.3) and len(o.trace["t"]) <= op.TRACE_POINTS + 2
    assert op.reclassify(o, 2200).regime == op.FLOW_CONTROLLED       # a corrected set point
    missing = op.operating_point(tmp_path / "nowhere.xlsx", 9000)
    assert missing.regime == op.UNKNOWN and "no pressure log" in missing.note


# ---- the simulated CCD day with a pressure-limited 0.5 mm block -----------------------------------------

DELIVERED_05 = {(3000, 900): 5450, (3000, 100): 3870, (4000, 500): 2400, (5000, 900): 1400, (5000, 100): 45}


@pytest.fixture(scope="module")
def day():
    with tempfile.TemporaryDirectory() as d:
        root = Path(d)
        runs = fakes.make_measured_ccd_day(root / "src", rb.build_sheet(ben_plan()), seed=3)
        for i, run in enumerate(runs):
            parts = run.name.split("_")
            sccm, rpm, sps = (float(parts[k].rstrip("sccmrpms")) for k in (1, 2, 3))
            orifice = float(parts[4][2:])
            if orifice == 0.5:
                add_pressure_log(run, sccm, DELIVERED_05[(sps, rpm)], 9.4, seed=i)
            else:
                add_pressure_log(run, sccm, sccm, {0.8: 4.8, 1.2: 3.0, 1.8: 1.6}[orifice], seed=i)
        design = rsm.from_runs(rd.load_runs(rd.find_runs([root / "src" / "2026"]).runs, check_reuse=False))
        keys = [f.key for f in design.detect().factors]
        out = {}
        for mode in ("exclude", "delivered", "set"):
            res = rp.publish(design, root / mode, [rp.ModelSpec("full quadratic", keys)],
                             discrete={"orifice": [0.5, 0.8, 1.2, 1.8]}, title=f"gas regimes ({mode})",
                             goals=[{"key": "atom", "direction": "max", "importance": 3},
                                    {"key": "d_mean", "direction": "min", "importance": 2},
                                    {"key": "af_flow", "direction": "max"}],
                             limited=mode, shared_extras=(mode == "exclude"))
            out[mode] = {"res": res, "R": next(iter(res.models.values())), "dir": root / mode}
        out["report"] = (root / "exclude" / "rsm_report.md").read_text(encoding="utf-8")
        yield out


def test_the_half_mm_block_is_found_pressure_limited_and_left_out(day):
    res, R = day["exclude"]["res"], day["exclude"]["R"]
    assert res.ok, res.problems
    assert len(R.limited) == 10 and all("_or0.5_" in r.point.name for r in R.limited)
    assert {r.op.regime for r in R.limited} == {op.PRESSURE_LIMITED}
    n_fit = sum(1 for r in R.loaded.runs if r.point.in_fit)
    assert n_fit == 21 and R.fits["atom"].full.n == 21               # 31 fitted points minus the ten
    assert all(r.point.excluded.startswith("pressure-limited: ") for r in R.limited)
    assert not any(p.run.point.excluded for p in R.fits["atom"].predictions)   # not "validated" either


def test_what_controlled_the_limited_runs_is_tested(day):
    eff = day["exclude"]["R"].limited_effects
    f = eff["delivered"]
    p = f.p_of
    assert p["sps"] < 0.01 and f.coef()["sps"] < 0           # more silicone, less gas (planted)
    assert p["sccm"] > 0.2                                    # the set point does not matter (planted)


def test_delivered_mode_fits_them_at_the_gas_they_got(day):
    R = day["delivered"]["R"]
    assert len(R.limited) == 10 and R.fits["atom"].full.n == 31
    got = {r.point.name: r.point.values["sccm"] for r in R.loaded.runs if "_or0.5_" in r.point.name}
    assert sorted(set(got.values())) == sorted({round(r.op.delivered_sccm) for r in R.limited})
    assert all("fitted at the gas delivered" in " ".join(r.point.flags) for r in R.loaded.runs
               if "_or0.5_" in r.point.name)


def test_set_mode_is_the_old_behaviour_and_says_so(day):
    R = day["set"]["R"]
    assert R.fits["atom"].full.n == 31 and len(R.limited) == 10
    assert any("fitted at their SET gas" in w for w in R.warnings)


def test_goals_used_and_dropped_are_all_listed(day):
    R = day["exclude"]["R"]
    st = {s.key: s for s in R.goal_status}
    assert set(st) == {"atom", "d_mean", "af_flow"}
    assert {g.key for g in R.goals} == {k for k, s in st.items() if s.used}
    assert all(s.why == "used" for s in st.values() if s.used)
    assert all(s.why != "used" for s in st.values() if not s.used)


def test_the_best_observed_runs_include_pressure_limited_ones(day):
    R = day["exclude"]["R"]
    assert 0 < len(R.best_observed) <= 5
    Ds = [b[0] for b in R.best_observed]
    assert Ds == sorted(Ds, reverse=True)
    assert all(b[1] in R.loaded.runs for b in R.best_observed)


def test_the_report_has_the_gas_section_and_the_goal_table(day):
    text = day["report"]
    assert "## Gas delivered and the operating regime" in text
    assert "### Runs that did not get their set gas (10)" in text
    assert "**What controlled them**" in text
    assert "| Response | Aim | Importance | Used | 0 at | 1 at | Why |" in text
    assert "**Best runs as measured**" in text and "- **Best run as measured**" in text
    assert "- **10 run(s) did not get their set gas** (pressure-limited), kept out of the surfaces" in text
    figs = {p.stem for p in day["exclude"]["res"].figures}
    assert {"fig_rsm_gas_delivered", "fig_gas_traces"} <= figs
    wb = openpyxl.load_workbook(day["exclude"]["dir"] / "rsm_analysis.xlsx", read_only=True)
    assert {"Gas delivered", "Goals used", "Best observed runs", "Pressure-limited effects"} <= set(wb.sheetnames)
    rows = list(csv.DictReader(open(day["exclude"]["dir"] / "rsm_results.csv", newline="", encoding="utf-8")))
    assert {r["in_fit"] for r in rows if "_or0.5_" in r["run"]} == {"0"}


def test_the_gas_flux_puts_both_regimes_on_one_axis(day):
    """Phase 12 (2026-10-13): the covariate section, figure and per-run sheet in the response-surface outputs;
    the pressure-limited runs' flux is from the gas they got in the camera window."""
    text = day["report"]
    assert "## Gas mass flux at the orifice" in text and "](figures/fig_gas_flux.png)" in text
    assert "33 of 33 runs have a flux (33 camera window)" in text
    assert "Orifice length is not recorded for these runs" in text
    sec = text.split("## Gas mass flux at the orifice")[1].split("\n## ")[0]
    assert "| Atomised fraction by area, classical (%) | Gas mass flux at the orifice" in sec
    assert "fig_gas_flux" in {p.stem for p in day["exclude"]["res"].figures}
    wb = openpyxl.load_workbook(day["exclude"]["dir"] / "rsm_analysis.xlsx", read_only=True)
    rows = list(wb["GLR context"].iter_rows(values_only=True))
    h = rows[0]
    got = [dict(zip(h, r)) for r in rows[1:]]
    half = [r for r in got if "_or0.5_" in r["run"]]
    assert len(half) == 10 and {r["gas flux source"] for r in got} == {"camera window"}
    for r in half:
        sps, rpm = (float(r["run"].split("_")[k].rstrip("sccmrpms")) for k in (3, 2))
        G = 1.145 * DELIVERED_05[(sps, rpm)] * 1e-6 / 60 / (np.pi * 0.0005 ** 2 / 4)
        assert r["gas mass flux (kg/m2/s)"] == pytest.approx(G, rel=0.02) and r["regime"] == op.PRESSURE_LIMITED
    trends = list(wb["GLR trends"].iter_rows(values_only=True))
    assert any(r[1] == "Gas mass flux at the orifice, delivered (kg/m2/s)" for r in trends[1:])


def test_the_silicone_film_is_in_the_response_surface_outputs(day):
    """2026-10-13: the orifice split (orifice.py) in the CCD report, workbook (with the diagram and the maths
    beside the table) and deck (the diagram slide just before the gas flux slide, the film slide after)."""
    from pptx import Presentation
    text = day["report"]
    assert "## Silicone film at the exit" in text and "](figures/fig_film_thickness.png)" in text
    assert "](figures/fig_orifice_model.png)" in text and "33 of 33 runs have a film estimate" in text
    wb = openpyxl.load_workbook(day["exclude"]["dir"] / "rsm_analysis.xlsx")
    ws = wb["Silicone film at exit"]
    assert len(ws._images) == 1
    head = [c.value for c in ws[1]]
    rows = [dict(zip(head, [c.value for c in r])) for r in ws.iter_rows(min_row=2, max_row=34)]
    half = [r for r in rows if "_or0.5_" in r["run"]]
    assert len(half) == 10 and all(r["P0 (barA)"] == pytest.approx(9.4, abs=0.05) and r["choked"] == "yes"
                                   for r in half)
    assert all(r["film h (um)"] is not None for r in rows)
    notes = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
    assert any(n.startswith("Film h = (d/2)(1 - sqrt(alpha))") for n in notes)
    prs = Presentation(str(day["exclude"]["res"].deck))
    titles = [next((sh.text_frame.text for sh in sl.shapes if sh.has_text_frame and sh.text_frame.text.strip()), "")
              for sl in prs.slides]
    i = titles.index("How the gas flux and the silicone film are estimated")
    assert "gas mass flux at the orifice" in titles[i + 1] and "silicone film thickness" in titles[i + 2]


def test_a_recommendation_at_the_ceiling_is_flagged(day):
    R = day["exclude"]["R"]
    assert rsm_recommend.pressure_check(R, {"peak_bar": (9.2, 8.5, 9.9)}).startswith("⚠ predicted peak 9.2 barA")
    assert rsm_recommend.pressure_check(R, {"peak_bar": (3.0, 2.5, 3.5)}) == ""
    assert rsm_recommend.pressure_check(R, {}) == ""


# ---- phase 6: orifice on a log scale, and the fitted span ----------------------------------------------------

def test_log_coding_makes_the_geometric_inserts_even():
    f = rsm.FactorLevels("orifice", "Orifice (mm)", [0.8, 1.2, 1.8], 1.2, log=True)
    assert [round(f.real_code(v), 9) for v in (0.8, 1.2, 1.8)] == [-1.0, 0.0, 1.0]
    assert [round(f.design_code(v), 9) for v in (0.8, 1.2, 1.8)] == [-1.0, 0.0, 1.0]
    assert f.decode(0.0) == pytest.approx(1.2) and f.decode(f.real_code(0.5)) == pytest.approx(0.5)
    lin = rsm.FactorLevels("orifice", "Orifice (mm)", [0.8, 1.2, 1.8], 1.2)
    assert lin.real_code(1.2) == pytest.approx(-0.2) and lin.decode(lin.real_code(1.5)) == pytest.approx(1.5)
    assert f.real_code(0) is None                                  # no log of zero


def test_leaving_runs_out_refits_the_span_and_the_best_setting_stays_inside(day):
    R = day["exclude"]["R"]
    o = R.factor("orifice")
    assert o.levels == [1.2, 1.8] and o.log                         # the 0.5 mm block is out; log by default
    assert R.optimum_discrete is not None
    assert round(R.decode("orifice", R.optimum_discrete.coded["orifice"]), 3) in (1.2, 1.8)
    assert any("settable levels outside the fitted range" in w and "0.5" in w and "0.8" in w for w in R.warnings)
    assert any("the fitted runs span 1.2 to 1.8" in w for w in R.warnings)
    # in the other modes every run is fitted, so the design span stands
    assert day["set"]["R"].factor("orifice").levels == [0.5, 1.2, 1.8]


# ---- 2026-10-13, found re-analysing the CCD day: the goal gate and the runs to strengthen the model -------

def test_the_goal_gate_judges_the_simplified_model_that_is_used():
    """With the 0.5 mm block out the real day had 21 runs for 14 terms: the full quadratic's predicted R²
    was -0.28 for atomised fraction while the simplified model, which draws the maps and finds the best
    setting, had 0.89. The gate must judge the model in use."""
    from types import SimpleNamespace as NS

    from src.ai.Taguchi_Analysis_UI import rsm_responses as rr

    def fit(full, red):
        return rr.ResponseFit("k", "k", "max", NS(r2_pred=full) if full is not None else None,
                              NS(r2_pred=red) if red is not None else None, {}, None, None, [], [])
    goals = [NS(key=k) for k in "abcd"]
    fits = {"a": fit(-0.28, 0.89), "b": fit(0.9, 0.3), "c": fit(0.7, None)}
    assert [g.key for g in rr.weak_goals(goals, fits)] == ["b", "d"]     # d has no fit at all
    assert fits["a"].used.r2_pred == 0.89 and fits["c"].used.r2_pred == 0.7


def test_the_goal_table_matches_the_gate(day):
    from src.ai.Taguchi_Analysis_UI import rsm_responses as rr
    R = day["exclude"]["R"]
    for s in R.goal_status:
        fr = R.fits.get(s.key)
        if fr is not None and fr.r2_pred_used is not None:
            assert s.used == (fr.r2_pred_used >= rr.USABLE_R2_PRED), s
            if not s.used:
                assert "(simplified model)" in s.why


def test_runs_to_strengthen_the_model_stay_inside_the_fitted_span(day):
    """0.5 mm on the ln scale fitted to 1.2-1.8 codes far outside, so its prediction variance is huge
    only because it is far away, and every 0.5 mm run was pressure-limited: never offer it."""
    R = day["exclude"]["R"]
    rec = rsm_recommend.augment(R, {"orifice": [0.5, 0.8, 1.2, 1.8]})
    assert rec is not None and rec.rows                       # orifice^2 is not estimable: the model is weak
    assert {r.values["orifice"] for r in rec.rows} <= {1.2, 1.8}
    assert "0.5, 0.8 are outside the fitted span (1.2 to 1.8)" in rec.why
    assert "at 0.5 all 10 runs were pressure-limited" in rec.why
    assert "0.8 ran flow-controlled" in rec.why
    assert "orifice^2 needs a third size" in rec.why
    # with every run in the fit, the whole design span is offered as before
    rec_set = rsm_recommend.augment(day["set"]["R"], {"orifice": [0.5, 0.8, 1.2, 1.8]})
    assert rec_set is None or "outside the fitted span" not in rec_set.why


def test_a_factor_no_steering_surface_uses_is_free_and_held_at_the_centre():
    """The real day: the simplified atomised-fraction surface was gas + orifice only, and the optimiser
    reported rpm 100 / feed 3,000 'on the edge', which was the grid's first value, not an optimum."""
    from types import SimpleNamespace as NS

    from src.ai.Taguchi_Analysis_UI import rsm_responses as rr
    live = {"atom": NS(terms=["1", "sccm", "orifice"]), "peak_bar": NS(terms=["1", "sccm*rpm"])}
    assert rr.free_factors(live, ["sccm", "rpm", "sps", "orifice"]) == ["sps"]
    assert rr.free_factors({"atom": NS(terms=["1"])}, ["sccm"]) == ["sccm"]


def test_free_factors_sit_at_the_centre_of_the_best_setting(day):
    R = day["exclude"]["R"]
    for o in (R.optimum, R.optimum_discrete):
        if o is None:
            continue
        for k in R.free:
            assert k not in o.on_edge
            if k != "orifice":
                assert o.coded[k] == 0.0
    if R.free:
        rep = day["report"]
        assert "in none of the surfaces it optimises" in rep


def test_a_blank_best_is_the_best_the_surface_reaches_not_the_best_run():
    """2026-10-13: with the best RUN as '1 at', every setting predicted beyond it scored a flat 1 and the
    optimiser stopped at the first it tried (feed 3,000 when maximising a rate that rises with feed)."""
    from types import SimpleNamespace as NS

    import numpy as np

    from src.ai.Taguchi_Analysis_UI import rsm_responses as rr
    from src.ai.Taguchi_Analysis_UI import rsm_stats as rs
    fit = NS(predict=lambda x: {"mean": 10 + 4 * np.asarray(x["a"]) + 2 * np.asarray(x["b"])})   # max 16 at (+1, +1)
    goals = [rs.Goal("up", "max", 2.0, 12.0), rs.Goal("down", "min", 5.0, 9.0), rs.Goal("set", "max", 2.0, 12.0),
             rs.Goal("weak", "max", 0.0, 1.0)]
    requested = [{"key": "up", "direction": "max", "low": None, "high": None},
                 {"key": "down", "direction": "min", "low": None, "high": None},
                 {"key": "set", "direction": "max", "low": None, "high": 12.0},          # the user typed a best
                 {"key": "weak", "direction": "max"}]                                    # no usable surface
    out, used = rr.reach_surface(goals, requested, {"up": fit, "down": fit, "set": fit}, ["a", "b"], {})
    g = {x.key: x for x in out}
    assert g["up"].high == pytest.approx(16.0) and g["up"].low == 2.0
    assert g["down"].low == pytest.approx(4.0) and g["down"].high == 9.0           # min end reached at (-1, -1)
    assert g["set"].high == 12.0 and g["weak"].high == 1.0
    assert used == pytest.approx({"up": 16.0, "down": 4.0})


def test_the_simulated_day_optimum_does_not_stop_at_the_first_setting_that_scores_1(day):
    R = day["exclude"]["R"]
    for g in R.goals:
        if g.key in R.surface_best:
            seen = [r.values[g.key] for r in R.loaded.runs if r.point.in_fit and r.values.get(g.key) is not None]
            assert (g.high if g.direction == "max" else g.low) == R.surface_best[g.key]
            assert (g.high >= max(seen)) if g.direction == "max" else (g.low <= min(seen))
    if R.optimum is not None and R.surface_best:
        assert R.optimum.D < 0.999 or all(abs(c) == 1 for c in R.optimum.coded.values())
    assert "best the fitted surface reaches" in day["report"] or not R.surface_best
