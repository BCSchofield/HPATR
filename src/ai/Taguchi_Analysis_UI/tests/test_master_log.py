"""master_log.xlsx at sizer 3.1.0 (plan 2026-10-17, Phase 4; Ben chose the columns): one layout, one set of numbers,
written by GUI_Clean at capture and by the analysis app after an analysis. Synthetic workbooks only: the suite never
opens the real log (conftest: HPATR_MASTER_LOG=off)."""
from __future__ import annotations

import json
import os

import cv2
import numpy as np
import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as XLImage

from src.ai.Taguchi_Analysis_UI import pipeline_spec as spec
from src.ai.Taguchi_Analysis_UI.tests.test_sizer31 import liq31, run31  # noqa: F401  (fixtures)

ML = spec.rdc_module("master_log")
OLD = ('Timestamp', 'Orifice', 'Orifice Length (mm)', 'Bubbler Height (mm)', 'Bubbler RPM', 'Flow Range (sccm)', 'GLR',
       'Pressure Range (barA)', 'Speed (steps/s)', 'Motor Travel (mm)', 'D32 (um)', 'Mean Diameter (um)',
       'Atomised (%)', 'Droplets/s (20-60 um)', 'Avg Lamella Thickness', 'Notes', 'Cone Image', 'Shadowgraph',
       'Pressure Graph', 'Mass Flow Graph')


def old_log(path, tmp_path, rows):
    """A log in the 20-column layout of 2026-10-13, with the six 'Sizer 3.0.0 ...' columns of 2026-10-16 appended
    and an image in the Shadowgraph column of every row."""
    png = tmp_path / "img.png"
    cv2.imwrite(str(png), np.full((20, 30, 3), 128, np.uint8))
    wb = Workbook()
    ws = wb.active
    ws.append(list(OLD) + [f"Sizer 3.0.0 {x}" for x in ("D32 (um)", "Mean", "Atom", "Atom vol", "Conc", "Note")])
    for i, (ts, notes, rate) in enumerate(rows, start=2):
        ws.append([ts, "1.2mm", 1.0, 0.5, 500, "0-9000", 1.0, "1-3", 4000, 36, 50.0, 40.0, 20.0, rate, "", notes,
                   "", "", "", "", 99.9, 1, 2, 3, 4, "old"])
        ws.add_image(XLImage(str(png)), f"R{i}")               # Shadowgraph: column R
    wb.save(path)


def test_the_layout_moves_to_3_1_0_with_every_image_and_value():
    assert len(ML.MASTER_HEADERS) == len(ML.MASTER_COL_WIDTHS) == 29
    i = ML.MASTER_HEADERS.index("Droplets/s (30-60 um)")
    assert ML.MASTER_HEADERS[i + 1:i + 6] == ("<30 um (/cm3)", "30-40 um (/cm3)", "40-60 um (/cm3)",
                                              "60-100 um (/cm3)", "100+ um (/cm3)")
    assert ML.MASTER_HEADERS[i + 6:i + 11] == ("In-band flow (mL/min)", "Concentration (/cm3)",
                                               "Atomised in volume (%)", "Sizer", "Avg Lamella Thickness")


def test_the_migration(tmp_path):
    p = tmp_path / "master_log.xlsx"
    old_log(p, tmp_path, [("2026-10-12 12:57:54", "CCD run", 223433), ("2026-10-05 10:00:00", "Taguchi", 51000)])
    wb = load_workbook(p)
    ws = wb.active
    assert ML.migrate_measurement_columns(ws) and not ML.migrate_measurement_columns(ws)     # once only
    ML.stamp_headers(ws)
    hdr = tuple(ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1))
    assert hdr == ML.MASTER_HEADERS                                  # the appended 3.0.0 columns are gone
    assert sorted(im.anchor._from.col for im in ws._images) == [26, 26]      # R (17) moved 9 right
    col = {h: c for c, h in enumerate(hdr, start=1)}
    assert ws.cell(row=2, column=col["Notes"]).value == "CCD run"     # values move with their columns
    assert ws.cell(row=3, column=col["Droplets/s (30-60 um)"]).value == 51000   # kept; Sizer blank says what it is
    assert ws.cell(row=3, column=col["Sizer"]).value is None
    assert "20-60 um by the 2.2.0" in ws.cell(row=1, column=col["Droplets/s (30-60 um)"]).comment.text


def test_one_runs_numbers_and_the_update(run31, liq31, tmp_path):
    run = run31["run"]
    wb = Workbook()
    wb.active.title = "Metadata"
    for row in (("Field", "Value"), ("Timestamp", "2026-10-17 12:00:00"), ("Speed (steps/s)", "4000")):
        wb.active.append(row)
    wb.save(run / "run_summary.xlsx")
    an = run / "shadowgraph" / "analysis"
    v = ML.measured_values(an / "droplets_0.30", an / "liquid_0.30", 4000)
    assert set(v) == set(ML.MEASURED) and v["Sizer"] == "3.1.0"
    bands = [v[h] for _lo, _hi, h in ML.BANDS]
    assert sum(bands) == pytest.approx(v["Concentration (/cm3)"], rel=1e-3)     # the bands add up to the total
    assert v["Droplets/s (30-60 um)"] > 0 and v["In-band flow (mL/min)"] > 0
    t = spec.rdc_module("rates").from_run_files(an / "droplets_0.30", an / "liquid_0.30", 4000)
    assert v["Droplets/s (30-60 um)"] == round(t["band_rate_per_s"])          # the app's number, by construction
    p = tmp_path / "master_log.xlsx"
    old_log(p, tmp_path, [("2026-10-17 12:00:00", "this run", 1.0), ("2026-10-05 10:00:00", "another", 2.0)])
    rep = ML.update_runs([run, tmp_path / "no_such_run"], path=p, log=lambda s: None)
    assert rep["updated"] == [run.name] and rep["skipped"] and rep["backup"]
    ws = load_workbook(p).active
    hdr = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    row = {h: ws.cell(row=2, column=c).value for c, h in enumerate(hdr, start=1)}
    assert row["Sizer"] == "3.1.0" and row["Droplets/s (30-60 um)"] == v["Droplets/s (30-60 um)"]
    assert row["Notes"] == "this run" and row["Speed (steps/s)"] == 4000          # nothing else in the row changes
    other = {h: ws.cell(row=3, column=c).value for c, h in enumerate(hdr, start=1)}
    assert other["Droplets/s (30-60 um)"] == 2.0 and other["Sizer"] is None      # rows not analysed: untouched
    assert len(ws._images) == 2
    rep2 = ML.update_runs([run], path=p, log=lambda s: None)
    assert rep2["updated"] == [run.name] and rep2["backup"] is None                  # one backup a day
    # a run with no row is reported, never appended
    q = tmp_path / "other_log.xlsx"
    old_log(q, tmp_path, [("2026-10-05 10:00:00", "another", 2.0)])
    rep3 = ML.update_runs([run], path=q, log=lambda s: None)
    assert rep3["no_row"] == [run.name] and not rep3["updated"]
    assert load_workbook(q).active.max_column == 26                                  # not even migrated


def test_the_suite_never_touches_the_real_log():
    assert os.environ.get("HPATR_MASTER_LOG") == "off" and ML.master_log_path() is None
    from src.ai.Taguchi_Analysis_UI import publish
    msgs = []
    rep = publish.update_master_log(["/nowhere"], 0.30, lambda text, level="info": msgs.append(text))
    assert rep is not None and rep["skipped"] and not rep["updated"]


def test_re_measuring_keeps_the_earlier_results(tmp_path):
    pc = spec.rdc_module("process_capture")
    an = tmp_path / "run" / "shadowgraph" / "analysis"

    def measured(version, h):
        for d in ("droplets_0.30", "liquid_0.30"):
            (an / d).mkdir(parents=True, exist_ok=True)
        (an / "droplets_0.30" / "summary.json").write_text(json.dumps(
            {"provenance": {"sizer_version": version, "optics_profile_hash": h}}))
        (an / "liquid_0.30" / "classical_summary.json").write_text("{}")
    measured("3.0.0", "aaa")
    kept = pc.archive_previous(an, 0.30, "3.1.0", "aaa", log=lambda s: None)
    assert kept == an / "_sizer3.0.0" and (kept / "droplets_0.30" / "summary.json").exists()
    assert (kept / "liquid_0.30").exists() and not (an / "droplets_0.30").exists()
    measured("3.1.0", "aaa")
    assert pc.archive_previous(an, 0.30, "3.1.0", "aaa", log=lambda s: None) is None      # the same: resumed
    assert pc.archive_previous(an, 0.30, "3.1.0", "bbb", log=lambda s: None) == an / "_sizer3.1.0"  # new beads
    measured("3.0.0", "aaa")
    assert pc.archive_previous(an, 0.30, "3.1.0", "aaa", log=lambda s: None) == an / "_sizer3.0.0 (2)"  # never over
    assert pc.archive_previous(tmp_path / "empty", 0.30, "3.1.0", "aaa", log=lambda s: None) is None
