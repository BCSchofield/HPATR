"""Phase 4b/5: rsm_publish end to end on measured fake runs (the path the Response surface tab's
Analyse button takes): every output written, the CSV in the CCD plan's column names, the 9x3's
two-model comparison, refusals and isolated failures; nothing written into a run folder."""
from __future__ import annotations

import builtins
import csv
import json
import tempfile
from pathlib import Path

import openpyxl
import pytest

from src.ai.Taguchi_Analysis_UI import pipeline_spec as spec
from src.ai.Taguchi_Analysis_UI import rsm_detect as rsm
from src.ai.Taguchi_Analysis_UI import rsm_figures, rsm_publish as rp
from src.ai.Taguchi_Analysis_UI import run_discovery as rd
from src.ai.Taguchi_Analysis_UI.tests import fakes

MODELS = [rp.ModelSpec("(a) gas x feed + RPM", ["sccm", "sps"], {"rpm": "quadratic"}),
          rp.ModelSpec("(b) gas x feed only", ["sccm", "sps"])]


@pytest.fixture(scope="module", autouse=True)
def small_bootstrap():
    ta = spec.rdc_module("taguchi_analysis")
    old = ta.N_BOOT
    ta.N_BOOT = 20
    yield
    ta.N_BOOT = old


def _design(root: Path) -> rsm.RsmDesign:
    return rsm.from_runs(rd.load_runs(rd.find_runs([root / "2026"]).runs, check_reuse=False))


@pytest.fixture(scope="module")
def world():
    with tempfile.TemporaryDirectory() as d:
        root = Path(d)
        fakes.make_measured_l9x3(root / "src", rpm_effect=0.04, interaction=0.03)
        out = root / "out"
        res = rp.publish(_design(root / "src"), out, MODELS, title="9x3 gas x feed")
        yield {"root": root, "out": out, "res": res}


def test_every_output_is_written(world):
    res, out = world["res"], world["out"]
    assert res.ok, res.problems
    names = {p.name for p in out.iterdir()}
    assert {"rsm_report.md", "rsm_analysis.xlsx", "rsm_results.csv", "rsm_deck.pptx", "rsm_design.json",
            "figures"} <= names
    assert not [n for n in names if n.startswith(".")]
    stems = {p.stem for p in res.figures}
    assert {"fig_rsm_contour_d32", "fig_rsm_surface_d32", "fig_rsm_profiler", "fig_rsm_desirability",
            "fig_rsm_fit_checks", "fig_rsm_model_compare", "fig_size_distribution"} <= stems


def test_the_csv_uses_the_ccd_plans_column_names(world):
    rows = list(csv.DictReader(open(world["out"] / "rsm_results.csv", encoding="utf-8")))
    assert len(rows) == 27
    assert {"gas_sccm", "feed_sps", "rpm", "af", "d32", "mean_d", "count_cv", "peak_bar", "af_flow"} <= set(rows[0])


def test_the_workbook_and_report_carry_both_models(world):
    wb = openpyxl.load_workbook(world["out"] / "rsm_analysis.xlsx", read_only=True)
    assert {"Runs", "Fit summary", "Coefficients", "Model comparison", "Best settings", "Provenance"} <= set(wb.sheetnames)
    models = {r[0] for r in wb["Fit summary"].iter_rows(min_row=2, values_only=True)}
    assert models == {m.name for m in MODELS}
    text = (world["out"] / "rsm_report.md").read_text(encoding="utf-8")
    for section in ("## What was found", "## The design and the models", "## Results by response",
                    "## Best settings", "## Runs kept out of the fit", "## Drift", "## Caveats", "## Provenance"):
        assert section in text, section
    assert "On its own it belongs in the Taguchi tab" in text            # an L9 fitted as a surface on purpose
    assert "Gas flow × Silicone" in text and "r |" in text                  # the RPM vs g*f correlation table


def test_the_deck_has_a_map_and_a_3d_surface_per_response(world):
    from pptx import Presentation
    prs = Presentation(str(world["out"] / "rsm_deck.pptx"))
    titles = [next((sh.text_frame.text for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()), "")
              for s in prs.slides]
    assert any(t.startswith("D32 in-focus: ") for t in titles) and any(t.startswith("D32 in-focus in 3D") for t in titles)
    assert "What was found" in titles and any(t.startswith("Best overall") for t in titles)
    assert [s.slide_layout.name for s in prs.slides][0] == "TITLE_DARK"


def test_nothing_is_written_into_a_run_folder(world):
    def snap():
        return {str(p): (p.stat().st_size, p.stat().st_mtime_ns) for p in (world["root"] / "src").rglob("*")
                if p.is_file()}
    before = snap()
    with tempfile.TemporaryDirectory() as d:
        rp.publish(_design(world["root"] / "src"), Path(d), MODELS[:1], shared_extras=False)
    assert snap() == before


def test_mixed_sizer_versions_refuse_and_write_no_report(world, tmp_path):
    design = _design(world["root"] / "src")
    summ = design.base.rows[0].path / "shadowgraph/analysis/droplets_0.30/summary.json"
    original = summ.read_text()
    try:
        s = json.loads(original)
        s["provenance"]["sizer_version"] = "2.0.0"
        summ.write_text(json.dumps(s))
        res = rp.publish(design, tmp_path, MODELS)
    finally:
        summ.write_text(original)
    assert res.refused and "sizer versions" in res.refused and res.report is None
    assert not (tmp_path / "rsm_report.md").exists()


def test_a_failing_figure_is_reported_and_the_rest_land(world, tmp_path, monkeypatch):
    def boom(*a, **k):
        raise RuntimeError("no ink")
    monkeypatch.setattr(rsm_figures, "fig_surface", boom)
    res = rp.publish(_design(world["root"] / "src"), tmp_path, MODELS, shared_extras=False)
    assert any("surface" in p and "no ink" in p for p in res.problems)
    assert res.report.is_file() and res.workbook.is_file() and res.deck.is_file()
    assert (tmp_path / "figures" / "fig_rsm_contour_d32.png").is_file()


def test_without_python_pptx_there_is_no_deck_and_everything_else_lands(world, tmp_path, monkeypatch):
    real = builtins.__import__

    def no_pptx(name, *a, **k):
        if name == "pptx" or name.startswith("pptx."):
            raise ImportError("No module named 'pptx'")
        return real(name, *a, **k)
    monkeypatch.setattr(builtins, "__import__", no_pptx)
    res = rp.publish(_design(world["root"] / "src"), tmp_path, MODELS[:1], shared_extras=False)
    assert res.deck is None and res.report.is_file() and not (tmp_path / "rsm_deck.pptx").exists()
