"""taguchi_deck.pptx: built from a published synthetic 9 x 3 (gas flow strong, RPM none), its
takeaways checked against what was built in, and the missing-python-pptx path."""
from __future__ import annotations

import builtins
import tempfile
from pathlib import Path

import pytest

from src.ai.Taguchi_Analysis_UI import deck, publish
from src.ai.Taguchi_Analysis_UI import design as dz
from src.ai.Taguchi_Analysis_UI import pipeline_spec as spec
from src.ai.Taguchi_Analysis_UI import run_discovery as rd
from src.ai.Taguchi_Analysis_UI.tests import fakes

pptx = pytest.importorskip("pptx")


@pytest.fixture(scope="module", autouse=True)
def small_bootstrap():
    ta = spec.rdc_module("taguchi_analysis")
    old = ta.N_BOOT
    ta.N_BOOT = 40
    yield
    ta.N_BOOT = old


@pytest.fixture(scope="module")
def world():
    with tempfile.TemporaryDirectory() as d:
        root = Path(d)
        fakes.make_measured_l9x3(root / "src")
        design = dz.detect(rd.load_runs(rd.find_runs([root / "src" / "2026"]).runs, check_reuse=False))
        res = publish.publish(design, root / "out")
        assert res.ok, res.problems
        yield {"root": root, "design": design, "res": res}


def _slides(path: Path):
    from pptx import Presentation
    prs = Presentation(str(path))
    out = []
    for s in prs.slides:
        texts = [sh.text_frame.text for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
        pics = sum(1 for sh in s.shapes if sh.shape_type == 13)
        tables = sum(1 for sh in s.shapes if sh.has_table)
        out.append({"title": texts[0] if texts else "", "texts": texts, "pics": pics, "tables": tables})
    return prs, out


def test_publish_writes_the_deck(world):
    res = world["res"]
    assert res.deck == world["root"] / "out" / deck.DECK_NAME and res.deck.is_file()
    assert not [p for p in res.deck.parent.iterdir() if p.name.startswith(".")]


def test_the_deck_has_its_slides_in_order_and_one_figure_per_figure_slide(world):
    prs, slides = _slides(world["res"].deck)
    assert (prs.slide_width, prs.slide_height) == (deck.WIDTH_EMU, deck.HEIGHT_EMU)
    assert slides[0]["title"] == "Taguchi analysis"
    assert slides[1]["tables"] == 1 and slides[1]["title"].startswith("The design: 27 runs: 9 conditions")
    assert slides[2]["title"] == "What was found"
    assert slides[-1]["title"].startswith("Caveats")
    figure_slides = [s for s in slides if s["pics"]]
    assert figure_slides and all(s["pics"] == 1 for s in figure_slides)
    for slide in prs.slides:                          # nothing placed off the slide
        for sh in slide.shapes:
            assert sh.left >= 0 and sh.top >= 0
            assert sh.left + sh.width <= prs.slide_width and sh.top + sh.height <= prs.slide_height


def test_takeaways_match_what_was_built_in(world):
    """The synthetic campaign's gas flow drives droplet size and atomisation; RPM does nothing."""
    _, slides = _slides(world["res"].deck)
    titles = [s["title"] for s in slides]
    assert "Gas flow drives D32 in-focus, mean diameter and atomised fraction by area" in titles, titles
    assert not any("RPM drives" in t for t in titles)
    caveats = " ".join(slides[-1]["texts"])
    n = len(world["res"].results.multiple)
    assert f"{n} factor tests" in caveats and "Fractional design" in caveats
    assert "rose with run order" in caveats                  # the fakes keep the real run order
    assert "No drift over the session beyond what the factors explain" in titles


def test_every_figure_is_in_the_deck_with_a_description(world):
    """Each figure the publish drew is on a slide (byte for byte), except the tall ones, which
    appear as their slide-sized parts; every picture slide carries a description."""
    import hashlib
    from pptx import Presentation
    figs = world["root"] / "out" / "figures"
    prs = Presentation(str(world["res"].deck))
    embedded = {hashlib.sha256(sh.image.blob).hexdigest() for s in prs.slides for sh in s.shapes
                if sh.shape_type == 13}
    split = {"fig_series", "fig_series_liquid", "fig_replicates", "fig_diagnostics"}
    for png in sorted(figs.glob("fig_*.png")):
        if png.stem in split:
            continue
        assert hashlib.sha256(png.read_bytes()).hexdigest() in embedded, png.name
    for png in sorted((figs / deck.SLIDE_FIG_DIR).glob("*.png")):
        assert hashlib.sha256(png.read_bytes()).hexdigest() in embedded, png.name
    assert {p.stem.rsplit("_", 1)[0] for p in (figs / deck.SLIDE_FIG_DIR).glob("*.png")} >= \
        {"fig_series", "fig_replicates", "fig_diagnostics"}
    for s in prs.slides:
        if any(sh.shape_type == 13 for sh in s.shapes):
            texts = [sh.text_frame.text for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
            assert len(texts) >= 3, texts                        # title, description, footer
            assert s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip()


def test_the_gas_flux_figure_is_drawn_and_has_its_slide(world):
    """Phase 12 (2026-10-13): the gas mass flux at the orifice, beside GLR, with a description."""
    figs = world["root"] / "out" / "figures"
    assert (figs / "fig_gas_flux.png").is_file() and deck._describe(figs / "fig_gas_flux.png") == \
        deck.DESCRIBE["fig_gas_flux"]
    _, slides = _slides(world["res"].deck)
    titles = [s["title"] for s in slides]
    i = next(i for i, t in enumerate(titles) if "with the gas mass flux at the orifice" in t)
    assert slides[i]["pics"] == 1 and titles[i - 1] == "How the gas flux and the silicone film are estimated"
    assert "with GLR" in titles[i - 2] and slides[i - 1]["pics"] == 1
    text = " ".join(slides[i]["texts"])
    assert "coloured by orifice (1.2 mm)" in text and "Gas mass flux G:" in text


@pytest.fixture(scope="module")
def logged():
    """The same campaign with GUI_Clean pressure logs (a camera window), so every run has P0 and a film."""
    from src.ai.Taguchi_Analysis_UI.tests.test_gas_regimes import add_pressure_log
    with tempfile.TemporaryDirectory() as d:
        root = Path(d)
        runs = fakes.make_measured_l9x3(root / "src")
        for i, run in enumerate(runs):
            sccm = float(run.name.split("_")[1][:-4])
            add_pressure_log(run, sccm, sccm, 1.4 + sccm / 4000, seed=i)
        design = dz.detect(rd.load_runs(rd.find_runs([root / "src" / "2026"]).runs, check_reuse=False))
        res = publish.publish(design, root / "out")
        assert res.ok, res.problems
        yield {"root": root, "res": res}


def test_the_film_figure_and_the_model_diagram_are_on_slides_in_order(logged):
    """2026-10-13: the diagram just before the gas flux slide, the film slide after it, each with a
    description; the workbook's film sheet carries the diagram and the maths beside its table."""
    import hashlib
    import openpyxl
    from pptx import Presentation
    figs = logged["root"] / "out" / "figures"
    for name in ("fig_orifice_model", "fig_gas_flux", "fig_film_thickness"):
        assert (figs / f"{name}.png").is_file() and deck._describe(figs / f"{name}.png") == deck.DESCRIBE[name]
    _, slides = _slides(logged["res"].deck)
    titles = [s["title"] for s in slides]
    i = titles.index("How the gas flux and the silicone film are estimated")
    assert "with the gas mass flux at the orifice" in titles[i + 1]
    assert "with the silicone film thickness at the exit" in titles[i + 2]
    assert all(slides[j]["pics"] == 1 for j in (i, i + 1, i + 2))
    assert "Film thickness h:" in " ".join(slides[i + 2]["texts"])
    prs = Presentation(str(logged["res"].deck))
    embedded = {hashlib.sha256(sh.image.blob).hexdigest() for s in prs.slides for sh in s.shapes
                if sh.shape_type == 13}
    for name in ("fig_orifice_model", "fig_film_thickness"):
        assert hashlib.sha256((figs / f"{name}.png").read_bytes()).hexdigest() in embedded
    wb = openpyxl.load_workbook(logged["root"] / "out" / "taguchi_analysis.xlsx")
    ws = wb["Silicone film at exit"]
    head = [c.value for c in ws[1]]
    assert {"Cd used", "Cd source", "P0 (barA)", "gas-only capacity (kg/m2/s)", "gas share alpha",
            "u_L (m/s)", "film h (um)", "atomised fraction (%)", "regime"} <= set(head)
    assert len(ws._images) == 1
    text = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
    assert any(t.startswith("How the gas flux and the silicone film are estimated") for t in text)
    assert any("ideal (not measured)" in t for t in text)
    rows = [dict(zip(head, [c.value for c in r])) for r in ws.iter_rows(min_row=2, max_row=28)]
    assert len(rows) == 27 and all(r["film h (um)"] and 0 < r["gas share alpha"] < 1 for r in rows)
    report = (logged["root"] / "out" / "taguchi_report.md").read_text(encoding="utf-8")
    assert "## Silicone film at the exit" in report and "](figures/fig_orifice_model.png)" in report
    assert "27 of 27 runs have a film estimate. Cd: ideal (not measured)" in report


def test_a_figure_that_was_not_drawn_drops_its_slide(world, tmp_path):
    figs = tmp_path / "figures"
    figs.mkdir()
    res = world["res"]
    path = deck.build(tmp_path / "d.pptx", res.results, res.bins, res.prov, figs)
    _, slides = _slides(path)
    # only the slide-sized parts the deck draws itself remain (time series, replicates, model checks)
    drawn = sorted(p.name for p in (figs / deck.SLIDE_FIG_DIR).glob("*.png"))
    assert sum(s["pics"] for s in slides) == len(drawn) and all(not n.startswith("gallery") for n in drawn)
    assert [s["title"] for s in slides][:3] == ["Taguchi analysis", slides[1]["title"], "What was found"]


def test_without_python_pptx_everything_else_is_written_and_the_report_says_why(world, monkeypatch):
    real_import = builtins.__import__

    def no_pptx(name, *a, **k):
        if name == "pptx" or name.startswith("pptx."):
            raise ImportError("No module named 'pptx'")
        return real_import(name, *a, **k)

    monkeypatch.setattr(builtins, "__import__", no_pptx)
    assert "python-pptx is not installed" in deck.available()
    with tempfile.TemporaryDirectory() as d:
        res = publish.publish(world["design"], Path(d))
        assert res.ok and res.deck is None and res.report.is_file() and res.workbook.is_file()
        assert not (Path(d) / deck.DECK_NAME).exists()
        assert "python-pptx is not installed" in res.report.read_text(encoding="utf-8")


def test_the_deck_uses_the_template_and_defines_its_terms(world):
    """Built on Ben's hand-made deck's theme (navy title slide, CONTENT layout); figure slides carry a
    'Terms' block; conditions named in descriptions carry their factor levels with units."""
    import re
    from pptx import Presentation
    assert deck.TEMPLATE.is_file()
    prs = Presentation(str(world["res"].deck))
    layouts = [s.slide_layout.name for s in prs.slides]
    assert layouts[0] == "TITLE_DARK" and layouts.count("CONTENT") == len(layouts) - layouts.count("TITLE_DARK")
    texts = [" ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame) for s in prs.slides]
    assert sum("Terms" in t for t in texts) >= 10
    called = [m for t in texts for m in re.finditer(r"\bT\d\b(?! has)", t)]
    with_levels = [t for t in texts if re.search(r"\bT\d \(\d+ sccm, \d+ rpm, \d+ steps/s\)", t)]
    assert called and with_levels


def test_every_definition_names_its_term():
    assert all(": " in v[:45] for v in deck.DEFINE.values())
