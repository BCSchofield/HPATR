"""Phase 12 of the CCD follow-up (2026-10-13): the gas mass flux at the orifice as a covariate (beside GLR),
the orifice LENGTH from run_summary.xlsx, the warning when the length differs between runs (a factor the
model does not know about), and the orifice split (orifice.py: gas share, silicone speed, film thickness)."""
from __future__ import annotations

import math
import tempfile
from pathlib import Path

import openpyxl
import pytest

from src.ai.Taguchi_Analysis_UI import covariates as cv
from src.ai.Taguchi_Analysis_UI import design as dz
from src.ai.Taguchi_Analysis_UI import operating
from src.ai.Taguchi_Analysis_UI import orifice
from src.ai.Taguchi_Analysis_UI import rsm_design as rb
from src.ai.Taguchi_Analysis_UI import rsm_detect as rsm
from src.ai.Taguchi_Analysis_UI import rsm_responses as rr
from src.ai.Taguchi_Analysis_UI import run_discovery as rd
from src.ai.Taguchi_Analysis_UI import stats
from src.ai.Taguchi_Analysis_UI.tests import fakes
from src.ai.Taguchi_Analysis_UI.tests.test_gas_regimes import add_pressure_log
from src.ai.Taguchi_Analysis_UI.tests.test_rsm_design import ben_plan

RUN = "115412_7000sccm_900rpm_3000sps_or0.5_bh0.5"         # C7 of the CCD day


@pytest.fixture
def tmp():
    with tempfile.TemporaryDirectory() as d:
        yield Path(d)


def set_length(run: Path, value) -> None:
    """Add GUI_Clean's new Metadata row (from 2026-10-13) to a fake run's workbook."""
    wb = openpyxl.load_workbook(run / "run_summary.xlsx")
    wb["Metadata"].append((rd.ORIFICE_LENGTH, value))
    wb.save(run / "run_summary.xlsx")


def ctx_with(length):
    return cv.RunContext(orifice_mm=1.2, orifice_length_mm=length)


# ---- the flux ------------------------------------------------------------------------------------------------

def test_flux_by_hand():
    # 9,000 sccm through 1.2 mm: 9000e-6/60 m3/s over pi 0.0012^2 / 4 m2 = 132.63 m/s; x 1.2 kg/m3 = 159.15
    G, v = cv.gas_flux(9000, 1.2, 1.2)
    assert v == pytest.approx(9000e-6 / 60 / (math.pi * 0.0012 ** 2 / 4)) == pytest.approx(132.63, abs=0.01)
    assert G == pytest.approx(159.15, abs=0.01) and G / 1.2 == pytest.approx(v)
    # C7 of the CCD day: 5,370 sccm delivered through 0.5 mm at the recorded 1.145 kg/m3
    G, v = cv.gas_flux(5370, 0.5, 1.145)
    assert v == pytest.approx(455.8, abs=0.1) and G == pytest.approx(521.9, abs=0.1)
    # no density: no flux (not guessed), but the velocity equivalent needs none; no orifice or flow: nothing
    assert cv.gas_flux(9000, 1.2, None) == (None, pytest.approx(132.63, abs=0.01))
    assert cv.gas_flux(9000, None, 1.2) == (None, None) and cv.gas_flux(None, 1.2, 1.2) == (None, None)


def test_source_order_camera_window_then_achieved_then_set_point(tmp):
    run = fakes.make_run(tmp, RUN, overrides={"Flow Range (sccm)": "0–7104 sccm"})
    c = cv.run_context(rd.load_run_info(run, check_reuse=False))
    assert c.orifice_mm == 0.5 and c.rho_gas == 1.145
    assert c.gas_flux_source == cv.ACHIEVED and c.flux_sccm == 7104
    assert c.gas_flux == pytest.approx(cv.gas_flux(7104, 0.5, 1.145)[0])
    # the fakes' own short log has no camera window: the flux stays on the achieved flow
    op = operating.operating_point(run / "run_summary.xlsx", 7000)
    assert op.delivered_sccm is None and cv.with_operating(c, op).gas_flux_source == cv.ACHIEVED
    # a log with a camera window: what it got there wins (pressure-limited, as C7 was)
    add_pressure_log(run, 7000, 5370, 9.2)
    op = operating.operating_point(run / "run_summary.xlsx", 7000)
    assert op.regime == operating.PRESSURE_LIMITED
    cv.with_operating(c, op)
    assert c.gas_flux_source == cv.CAMERA and c.flux_sccm == pytest.approx(5370, rel=0.02)
    assert c.gas_flux == pytest.approx(521.9, rel=0.02) and c.exit_velocity_std == pytest.approx(c.gas_flux / 1.145)
    # no flow range in the workbook: the set point from the folder name
    run2 = fakes.make_run(tmp / "b", RUN, omit=("Flow Range (sccm)",))
    c2 = cv.run_context(rd.load_run_info(run2, check_reuse=False))
    assert c2.gas_flux_source == cv.SET_POINT and c2.flux_sccm == 7000
    # no gas density: no flux, said why; the velocity equivalent is still there
    run3 = fakes.make_run(tmp / "c", RUN, omit=("Gas density (kg/m3)",))
    c3 = cv.run_context(rd.load_run_info(run3, check_reuse=False))
    assert c3.gas_flux is None and "no gas density" in c3.flux_note and c3.exit_velocity_std is not None


def test_the_analysis_uses_the_camera_window_where_there_is_one(tmp):
    runs = fakes.make_measured_l9x3(tmp)
    add_pressure_log(runs[0], 3000, 2990, 1.6)
    res = stats.analyse(dz.detect(rd.load_runs(rd.find_runs([tmp / "2026"]).runs, check_reuse=False)),
                        do_bootstrap=False)
    src = {r.name: r.context.gas_flux_source for r in res.runs}
    assert src.pop(runs[0].name) == cv.CAMERA and set(src.values()) == {cv.ACHIEVED}
    t = res.trends[("atom", "gas_flux")]
    assert t.n == 27 and t.r > 0.5 and t.exponent > 0           # the fakes' gas effect, through one orifice
    assert ("atom", "l_over_d") in res.trends and res.trends[("atom", "l_over_d")].r is None   # not recorded
    assert res.warnings == []                                    # no length recorded anywhere: nothing to warn


# ---- the orifice length --------------------------------------------------------------------------------------

def test_orifice_length_is_read_from_the_workbook_with_l_over_d(tmp):
    run = fakes.make_run(tmp, "090432_3000sccm_300rpm_4000sps_or1.2_bh1")
    info = rd.load_run_info(run, check_reuse=False)
    assert info.orifice_length_mm is None                        # every run captured before 2026-10-13
    c = cv.run_context(info)
    assert c.orifice_length_mm is None and c.l_over_d is None
    set_length(run, 3)
    info = rd.load_run_info(run, check_reuse=False)
    assert info.orifice_length_mm == 3.0 and not info.warnings   # not a factor: no "missing" warning either
    assert cv.run_context(info).l_over_d == pytest.approx(2.5)
    run2 = fakes.make_run(tmp / "b", "090432_3000sccm_300rpm_4000sps_or1.2_bh1")
    set_length(run2, "NOT RECORDED")
    assert rd.load_run_info(run2, check_reuse=False).orifice_length_mm is None


def test_the_length_warning_only_when_lengths_differ_or_some_are_unrecorded():
    w = cv.orifice_length_warning([ctx_with(3.0)] * 20 + [ctx_with(5.0)] * 13)
    assert w.startswith("ORIFICE LENGTH DIFFERS BETWEEN RUNS (3.0 mm in 20, 5.0 mm in 13)")
    assert "not a factor in this model, so its effect is mixed into the others" in w
    w = cv.orifice_length_warning([ctx_with(3.0)] * 4 + [ctx_with(5.0)] * 3 + [ctx_with(None)] * 2)
    assert "(3.0 mm in 4, 5.0 mm in 3; not recorded in 2)" in w
    w = cv.orifice_length_warning([ctx_with(3.0)] * 5 + [ctx_with(None)] * 3)
    assert w.startswith("orifice length is recorded for only 5 of 8 runs (3.0 mm); the other 3 do not record it")
    assert cv.orifice_length_warning([ctx_with(3.0), ctx_with(3.02), ctx_with(2.98)]) is None   # one insert
    assert cv.orifice_length_warning([ctx_with(None)] * 5) is None
    assert cv.orifice_length_warning([]) is None


def test_the_taguchi_analysis_warns_when_the_length_changed(tmp):
    runs = fakes.make_measured_l9x3(tmp)
    for i, run in enumerate(runs):
        set_length(run, 5.0 if i < 10 else 3.0)
    res = stats.analyse(dz.detect(rd.load_runs(rd.find_runs([tmp / "2026"]).runs, check_reuse=False)),
                        do_bootstrap=False)
    assert [w for w in res.warnings if "ORIFICE LENGTH" in w] == [
        "ORIFICE LENGTH DIFFERS BETWEEN RUNS (3.0 mm in 17, 5.0 mm in 10): it is not a factor in this model, so "
        "its effect is mixed into the others."]
    assert {round(r.context.l_over_d, 3) for r in res.runs} == {2.5, 4.167}
    assert res.trends[("atom", "l_over_d")].n == 27


def test_the_response_surface_load_warns_when_the_length_changed_and_not_when_equal(tmp):
    runs = fakes.make_measured_ccd_day(tmp / "src", rb.build_sheet(ben_plan()), seed=3, n_frames=12)
    for run in runs:
        set_length(run, 5.0 if "_or1.8_" in run.name else 3.0)

    def load():
        design = rsm.from_runs(rd.load_runs(rd.find_runs([tmp / "src" / "2026"]).runs, check_reuse=False))
        return rr.load(design.detect(), design.base.runs)

    n18 = sum("_or1.8_" in r.name for r in runs)
    w = [x for x in load().warnings if "ORIFICE LENGTH" in x]
    assert len(w) == 1 and f"(3.0 mm in {len(runs) - n18}, 5.0 mm in {n18})" in w[0]
    for run in runs:
        if "_or1.8_" in run.name:
            wb = openpyxl.load_workbook(run / "run_summary.xlsx")
            ws = wb["Metadata"]
            ws.cell(row=ws.max_row, column=2).value = 3.0
            wb.save(run / "run_summary.xlsx")
    assert not [x for x in load().warnings if "orifice length" in x.lower()]


# ---- the orifice split: gas share, silicone speed, film (orifice.py) ----------------------------------------

def test_the_split_by_hand_for_c7():
    # C7: 5,370 sccm N2 (1.145 kg/m3) through 0.5 mm, P0 9.18 barA, feed 3,000 steps/s = 8.544 mL/min
    G, _ = cv.gas_flux(5370, 0.5, 1.145)
    g, R, T = 1.4, 296.8, 293.15
    G_star = 9.18e5 * math.sqrt(g / (R * T)) * (2 / (g + 1)) ** ((g + 1) / (2 * (g - 1)))
    ideal, choked = orifice.gas_only_flux(9.18, g, R)
    assert choked and ideal == pytest.approx(G_star) == pytest.approx(2131, abs=1)
    q = 3000 / 6800 * math.pi * 20.27 ** 2 / 4 * 1e-9          # m3/s, the syringe
    sp = orifice.split(G, ideal, 0.5, q)
    alpha = G / G_star
    assert sp.gas_share == pytest.approx(alpha) == pytest.approx(0.245, abs=0.001)
    assert sp.film_um == pytest.approx(250 * (1 - math.sqrt(alpha))) == pytest.approx(126, abs=1)
    assert sp.u_liquid_m_s == pytest.approx(q / ((1 - alpha) * math.pi * 0.0005 ** 2 / 4)) == \
        pytest.approx(0.96, abs=0.01)
    assert sp.note == ""


def test_choked_and_subsonic_branches_meet_at_the_critical_ratio():
    g, R = 1.4, 287.05
    pc = orifice.PA_BAR / orifice.critical_ratio(g)            # 1.918 barA
    assert orifice.critical_ratio(g) == pytest.approx(0.5283, abs=1e-4)
    above, ch1 = orifice.gas_only_flux(pc * 1.0001, g, R)
    below, ch2 = orifice.gas_only_flux(pc * 0.9999, g, R)
    assert ch1 and not ch2 and above == pytest.approx(below, rel=1e-3)
    sub, choked = orifice.gas_only_flux(1.5, g, R)               # r = 0.675: subsonic
    r = orifice.PA_BAR / 1.5
    by_hand = 1.5e5 * math.sqrt(2 * g / ((g - 1) * R * orifice.T_K) * (r ** (2 / g) - r ** ((g + 1) / g)))
    assert not choked and sub == pytest.approx(by_hand) and sub < orifice.gas_only_flux(3.0, g, R)[0]
    assert orifice.gas_only_flux(1.0, g, R) == (0.0, False)     # no pressure drop, no capacity
    s = orifice.split(100.0, 0.0, 1.2, None)
    assert s.film_um is None and "at or below atmospheric" in s.note
    over = orifice.split(800.0, 700.0, 1.2, 1e-7)               # more gas than gas alone could pass
    assert over.gas_share == 1.0 and over.film_um == 0 and over.u_liquid_m_s is None
    assert "exceeds the ideal gas-only capacity" in over.note


def test_gas_properties_from_the_gas_tag():
    assert orifice.gas_properties("Nitrogen Gas")[:3] == ("nitrogen", 1.4, 296.8)
    assert orifice.gas_properties("N2")[0] == "nitrogen"
    assert orifice.gas_properties("CO2")[:3] == ("co2", 1.29, 188.9)
    assert orifice.gas_properties("Air")[:4] == ("air", 1.4, 287.05, "")
    name, *_rest, note = orifice.gas_properties("Argon")
    assert name == "air" and "'Argon' not known" in note


def test_cd_lookup():
    assert [e["diameter_mm"] for e in orifice.load_cd()] == [0.5, 0.8, 1.2, 1.8]   # the shipped file
    assert all(e["cd"] is None for e in orifice.load_cd())
    assert orifice.lookup_cd(1.2, None) == (1.0, orifice.IDEAL)                   # cd null: ideal
    table = [{"diameter_mm": 1.2, "length_mm": 3.0, "cd": 0.82, "measured": "2026-10-15"},
             {"diameter_mm": 1.2, "length_mm": 5.0, "cd": 0.78, "measured": "2026-10-16"},
             {"diameter_mm": 0.8, "length_mm": None, "cd": 0.85, "measured": "2026-10-15"}]
    assert orifice.lookup_cd(1.2, 5.02, table) == (0.78, "measured 2026-10-16")    # diameter and length
    assert orifice.lookup_cd(1.2, 3.0, table) == (0.82, "measured 2026-10-15")
    assert orifice.lookup_cd(1.2, 4.0, table) == (1.0, orifice.IDEAL)              # no insert of that length
    assert orifice.lookup_cd(0.8, 6.0, table) == (0.85, "measured 2026-10-15")     # the entry has no length
    assert orifice.lookup_cd(0.805, None, table) == (0.85, "measured 2026-10-15")  # within 0.01 mm
    assert orifice.lookup_cd(1.8, None, table) == (1.0, orifice.IDEAL)             # not in the table
    assert orifice.lookup_cd(1.2, None, []) == (1.0, orifice.IDEAL)


def test_the_camera_window_gives_the_film(tmp):
    run = fakes.make_run(tmp, RUN, overrides={"Flow Range (sccm)": "0–7104 sccm"})
    c = cv.run_context(rd.load_run_info(run, check_reuse=False))
    cv.with_operating(c, operating.operating_point(run / "run_summary.xlsx", 7000))
    assert c.film_um is None and "no camera window" in c.film_note             # the fakes' short log
    add_pressure_log(run, 7000, 5370, 9.18)
    cv.with_operating(c, operating.operating_point(run / "run_summary.xlsx", 7000))
    assert c.gas_model == "nitrogen" and c.cd == 1.0 and c.cd_source == orifice.IDEAL and c.choked
    assert c.p0_bar == pytest.approx(9.18, abs=0.02) and c.gas_share == pytest.approx(0.245, abs=0.005)
    assert c.film_um == pytest.approx(126, abs=2) and c.u_liquid_m_s == pytest.approx(0.96, abs=0.02)


def test_no_spray_runs_are_plotted_not_fitted():
    from types import SimpleNamespace as NS
    runs = []
    for i, (x, y) in enumerate([(50, 10), (70, 14), (150, 19), (160, 20), (300, 21), (520, 35), (3.3, 0.2),
                                (5.3, 0.45)]):
        rec = {"no_spray": "atomised 0.2% < 2%"} if x < 10 else {}
        runs.append(stats.RunRecord(f"r{i}", Path(f"r{i}"), f"C{i}", 1, {}, {"atom": y}, {}, None, 1, rec,
                                    NS(gas_flux=x, film_um=500 - x, glr=x / 100)))
    t = cv.trends(runs, [("atom", "Atomised", "max", None)])[("atom", "gas_flux")]
    assert t.n == 6 and len(t.points) == 6 and len(t.excluded) == 2
    assert "2 no-spray run(s) plotted, not fitted" in t.note
    with_all = cv.trend("atom", "gas_flux", [p[0] for p in t.points + t.excluded],
                        [p[1] for p in t.points + t.excluded], ["c"] * 8)
    assert with_all.exponent > 2 * t.exponent                     # the two dribbles set the exponent
    assert len(cv.trends(runs, [("atom", "Atomised", "max", None)])[("atom", "glr")].excluded) == 2
