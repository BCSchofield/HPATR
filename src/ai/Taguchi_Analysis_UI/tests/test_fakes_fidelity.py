"""The fakes must look like the real thing, or every test built on them proves
nothing. These compare the fixtures against the real 2026/10/05 data whenever
the LaCie is mounted (skipped otherwise)."""
from __future__ import annotations

import re
import tempfile
from pathlib import Path

import openpyxl
import pytest

from src.ai.Taguchi_Analysis_UI.tests import fakes

REAL_DAY = Path("/Volumes/LaCie/Experiments/2026/10/05")
needs_lacie = pytest.mark.skipif(not REAL_DAY.is_dir(), reason="LaCie not mounted")
FIRST = fakes.REAL_10_05[0][0]


def _tree(root: Path) -> set[str]:
    """Relative paths, dirs marked with a trailing slash. The cine's own file
    name is masked: its time legitimately differs run to run."""
    out = set()
    for p in sorted(root.rglob("*")):
        rel = p.relative_to(root).as_posix() + ("/" if p.is_dir() else "")
        out.add(re.sub(r"recording_\d{6}\.cine$", "recording_<hhmmss>.cine", rel))
    return out


@needs_lacie
def test_fake_run_tree_matches_real_run():
    with tempfile.TemporaryDirectory() as tmp:
        fake = fakes.make_run(Path(tmp), FIRST)
        assert _tree(fake) == _tree(REAL_DAY / FIRST)


@needs_lacie
def test_fake_metadata_matches_real_field_order_types_and_values():
    real = openpyxl.load_workbook(REAL_DAY / FIRST / "run_summary.xlsx", data_only=True)
    with tempfile.TemporaryDirectory() as tmp:
        fake_dir = fakes.make_run(Path(tmp), FIRST,
                                  notes="Taguchi ReRun 1 - Repeat 1\nWITH NOZZLE ADAPTER")
        fake = openpyxl.load_workbook(fake_dir / "run_summary.xlsx", data_only=True)
    assert fake.sheetnames == real.sheetnames == ["Metadata", "Pressure"]
    real_rows = list(real["Metadata"].iter_rows(values_only=True))
    fake_rows = list(fake["Metadata"].iter_rows(values_only=True))
    assert [r[0] for r in fake_rows] == [r[0] for r in real_rows]            # field order
    assert [type(r[1]).__name__ for r in fake_rows] == [type(r[1]).__name__ for r in real_rows]
    # these must be reproduced exactly for run 1
    keep = {"Timestamp", "Orifice", "Bubbler Height (mm)", "Bubbler RPM", "Speed (steps/s)",
            "FPS", "Notes", "Gas", "Fluid"}
    assert {r[0]: r[1] for r in fake_rows if r[0] in keep} == \
           {r[0]: r[1] for r in real_rows if r[0] in keep}
    assert [c.value for c in next(fake["Pressure"].iter_rows())] == list(fakes.PRESSURE_COLUMNS) \
           == [c.value for c in next(real["Pressure"].iter_rows())]


@needs_lacie
def test_fakes_real_name_table_matches_the_real_day_exactly():
    real_names = sorted(p.name for p in REAL_DAY.iterdir() if p.is_dir())
    assert [n for n, _, _ in fakes.REAL_10_05] == real_names


@needs_lacie
def test_fakes_real_notes_table_matches_real_notes():
    notes = {}
    for name, _, _ in fakes.REAL_10_05:
        ws = openpyxl.load_workbook(REAL_DAY / name / "run_summary.xlsx", read_only=True,
                                    data_only=True)["Metadata"]
        notes[name] = next(r[1] for r in ws.iter_rows(values_only=True) if r[0] == "Notes")
    for name, trial, rep in fakes.REAL_10_05:
        assert notes[name].startswith(f"Taguchi ReRun {trial} - Repeat {rep}"), name


def test_fake_l9x3_is_a_balanced_9x3_in_the_real_non_condition_order():
    names = [n for n, _, _ in fakes.REAL_10_05]
    assert len(names) == 27
    conds = {re.sub(r"^\d{6}_", "", n) for n in names}
    assert len(conds) == 9
    # folder order != condition order (the trap replicate grouping must survive)
    assert names.index("111625_6000sccm_600rpm_8000sps_or1.2_bh1") < \
           names.index("112253_3000sccm_600rpm_6000sps_or1.2_bh1")
    from collections import Counter
    assert set(Counter(re.sub(r"^\d{6}_", "", n) for n in names).values()) == {3}
