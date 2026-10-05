"""Phase 9 deliverables: figures, report, workbook, flat CSVs, odd-frame pack, publish().

Everything runs on a synthetic MEASURED 9 x 3 campaign (known effects: gas flow strong,
silicone slight, RPM none), so each statement the report makes can be checked against what
was built in. Nothing here touches the real runs."""
from __future__ import annotations

import csv
import json
import os
import re
import shutil
import tempfile
from pathlib import Path

import numpy as np
import pytest

from src.ai.Taguchi_Analysis_UI import design as dz
from src.ai.Taguchi_Analysis_UI import figures, odd, provenance, publish, report, size_bins, stats, tables
from src.ai.Taguchi_Analysis_UI import pipeline_spec as spec
from src.ai.Taguchi_Analysis_UI import run_discovery as rd
from src.ai.Taguchi_Analysis_UI import workbook
from src.ai.Taguchi_Analysis_UI.tests import fakes

BOTH = {"flat_csvs": True, "odd_pack": True}


def load_design(root: Path) -> dz.Design:
    return dz.detect(rd.load_runs(rd.find_runs([root / "2026"]).runs, check_reuse=False))


@pytest.fixture(scope="module", autouse=True)
def small_bootstrap():
    ta = spec.rdc_module("taguchi_analysis")
    old = ta.N_BOOT
    ta.N_BOOT = 60
    yield
    ta.N_BOOT = old


@pytest.fixture(scope="module")
def world():
    """One synthetic campaign, analysed and published once with every option on."""
    with tempfile.TemporaryDirectory() as d:
        root = Path(d)
        fakes.make_measured_l9x3(root / "src")
        design = load_design(root / "src")
        out = root / "out"
        calls = []
        res = publish.publish(design, out, options=BOTH, progress=lambda i, n, name: calls.append((i, n)))
        yield {"root": root / "src", "out": out, "design": design, "res": res, "calls": calls}


# ---- publish() ---------------------------------------------------------------------------------------------

def test_publish_writes_every_mandatory_output_and_the_optional_ones_asked_for(world):
    res, out = world["res"], world["out"]
    assert res.ok, res.problems
    names = {p.name for p in out.iterdir()}
    assert {"taguchi_report.md", "taguchi_analysis.xlsx", "figures", "taguchi_design.json",
            "csv", "odd"} <= names
    assert not [n for n in names if n.startswith(".")]            # no temp files left behind


def test_progress_is_reported_once_per_run(world):
    assert world["calls"][0] == (1, 27) and world["calls"][-1] == (27, 27) and len(world["calls"]) == 27


def test_optional_outputs_are_not_written_unless_ticked(world):
    with tempfile.TemporaryDirectory() as d:
        res = publish.publish(world["design"], Path(d))
        assert res.ok and res.csv_dir is None and res.odd is None
        assert not (Path(d) / "csv").exists() and not (Path(d) / "odd").exists()


def test_nothing_is_ever_written_into_a_run_folder(world):
    def snapshot():
        return {str(p): (p.stat().st_size, p.stat().st_mtime_ns)
                for p in world["root"].rglob("*") if p.is_file() or p.is_symlink()}
    before = snapshot()
    with tempfile.TemporaryDirectory() as d:
        publish.publish(world["design"], Path(d), options=BOTH)
    assert snapshot() == before


def test_a_refused_analysis_leaves_the_existing_report_untouched(world):
    with tempfile.TemporaryDirectory() as d:
        out = Path(d)
        good = publish.publish(world["design"], out)
        assert good.ok
        before = {p.name: p.read_bytes() for p in out.iterdir() if p.is_file()}
        # make two runs disagree about the sizer version
        design = load_design(world["root"])
        summ = design.rows[0].path / "shadowgraph/analysis/droplets_0.30/summary.json"
        original = summ.read_text()
        try:
            s = json.loads(original); s["provenance"]["sizer_version"] = "2.0.0"
            summ.write_text(json.dumps(s))
            res = publish.publish(design, out)
        finally:
            summ.write_text(original)
        assert res.refused and "different sizer versions" in res.refused and not res.ok
        assert res.report is None and res.workbook is None and not res.figures
        after = {p.name: p.read_bytes() for p in out.iterdir() if p.is_file()}
        assert after == before


def test_one_failing_step_does_not_lose_the_others_and_is_reported(world, monkeypatch):
    def boom(*a, **k):
        raise RuntimeError("matplotlib exploded")
    monkeypatch.setattr(figures, "make_all", boom)
    messages = []
    with tempfile.TemporaryDirectory() as d:
        res = publish.publish(world["design"], Path(d), log=lambda t, l="info": messages.append((l, t)))
        assert not res.ok and any("figures failed" in p and "matplotlib exploded" in p for p in res.problems)
        assert res.report and res.report.is_file() and res.workbook and res.workbook.is_file()
        assert ("error", res.problems[0]) in messages
        assert "figures failed" in res.report.read_text(encoding="utf-8")       # and the report says so


def test_publishing_again_replaces_the_previous_outputs(world):
    with tempfile.TemporaryDirectory() as d:
        out = Path(d)
        publish.publish(world["design"], out)
        (out / "taguchi_report.md").write_text("STALE")
        publish.publish(world["design"], out)
        assert "STALE" not in (out / "taguchi_report.md").read_text(encoding="utf-8")
        assert not [p for p in out.iterdir() if p.name.startswith(".")]


# ---- figures --------------------------------------------------------------------------------------------------

EXPECTED_FIGS = ("fig_main_effects", "fig_contributions", "fig_per_run", "fig_sn", "fig_glr",
                 "fig_size_spread_count", "fig_size_spread_volume")


def test_every_figure_is_written_as_png_and_svg(world):
    fig_dir = world["out"] / "figures"
    for name in EXPECTED_FIGS:
        png, svg = fig_dir / f"{name}.png", fig_dir / f"{name}.svg"
        assert png.is_file() and svg.is_file(), name
        assert png.read_bytes()[:8] == b"\x89PNG\r\n\x1a\n" and b"<svg" in svg.read_bytes()[:400]
    assert len(world["res"].figures) == 2 * len(EXPECTED_FIGS)


def test_pngs_are_300_dpi_not_screen_resolution(world):
    import struct
    data = (world["out"] / "figures" / "fig_main_effects.png").read_bytes()
    width = struct.unpack(">I", data[16:20])[0]
    assert width > 2500                                   # a ~10 inch figure at 300 dpi


def test_figures_use_the_object_api_not_pyplot():
    """pyplot keeps global state and a GUI backend; the analysis runs on a thread beside Qt."""
    import ast
    tree = ast.parse(Path(figures.__file__).read_text(encoding="utf-8"))
    imported = []
    for node in ast.walk(tree):
        if isinstance(node, ast.Import):
            imported += [a.name for a in node.names]
        elif isinstance(node, ast.ImportFrom):
            imported += [node.module or ""] + [f"{node.module}.{a.name}" for a in node.names]
    assert imported and not [m for m in imported if "pyplot" in m]


@pytest.fixture
def captured(monkeypatch):
    got = {}
    monkeypatch.setattr(figures, "_save", lambda fig, out_dir, name: got.__setitem__(name, fig) or [])
    return got


def test_size_spread_stacks_sum_to_100_with_runs_on_x_grouped_by_condition(world, captured):
    res = world["res"]
    figures.fig_size_spread_count(res.bins, Path("x"))
    figures.fig_size_spread_volume(res.bins, Path("x"))
    for name in ("fig_size_spread_count", "fig_size_spread_volume"):
        ax = captured[name].axes[0]
        bars = [p for p in ax.patches]
        assert len(bars) == 27 * len(res.bins.labels)
        tops = {}
        for p in bars:
            tops[round(p.get_x() + p.get_width() / 2)] = tops.get(round(p.get_x() + p.get_width() / 2), 0) + p.get_height()
        assert len(tops) == 27 and all(abs(v - 100) < 1e-6 for v in tops.values())
        ticks = [t.get_text() for t in ax.get_xticklabels()]
        cond = res.bins.run_condition
        names = {r.name[:6]: r.condition for r in res.results.runs}
        seq = [names[t] for t in ticks]
        assert seq == sorted(seq)                          # replicates of a condition sit together
        legend = [t.get_text() for t in ax.get_legend().get_texts()]
        assert legend == [f"{lab} µm" for lab in reversed(res.bins.labels)]


def test_size_bin_colours_are_one_ordered_ramp_not_a_categorical_palette():
    cols = figures._ramp(9)
    lum = [0.2126 * r + 0.7152 * g + 0.0722 * b for r, g, b, _ in cols]
    assert all(a > b for a, b in zip(lum, lum[1:])) or all(a < b for a, b in zip(lum, lum[1:]))


def test_main_effects_has_one_panel_per_response_and_factor(world, captured):
    figures.fig_main_effects(world["res"].results, Path("x"))
    assert len(captured["fig_main_effects"].axes) == 2 * 3


def test_contributions_bars_partition_100_percent(world, captured):
    figures.fig_contributions(world["res"].results, Path("x"))
    ax = captured["fig_contributions"].axes[0]
    by_row = {}
    for p in ax.patches:
        by_row[round(p.get_y())] = by_row.get(round(p.get_y()), 0) + p.get_width()
    assert by_row and all(abs(v - 100) < 1e-6 for v in by_row.values())


def test_figures_with_nothing_to_draw_write_nothing(world):
    r = world["res"].results
    empty = stats.Results(r.factors, r.runs, {}, [], [])
    with tempfile.TemporaryDirectory() as d:
        assert figures.make_all(empty, None, Path(d)) == []
        assert not list(Path(d).iterdir())
        assert figures.fig_glr(empty, Path(d)) == [] and figures.fig_sn(empty, Path(d)) == []


# ---- report ---------------------------------------------------------------------------------------------------------

@pytest.fixture(scope="module")
def text(world):
    return (world["out"] / "taguchi_report.md").read_text(encoding="utf-8")


def test_report_headlines_the_effects_that_were_built_in(world, text):
    found = text.split("## What was found")[1].split("## Design")[0]
    d32 = next(l for l in found.splitlines() if "D32" in l)
    atom = next(l for l in found.splitlines() if "Atomised" in l)
    assert "Gas flow (sccm)" in d32 and "Gas flow (sccm)" in atom
    assert "Bubbler RPM" not in d32 and "Bubbler RPM" not in atom       # built in: no effect
    assert "GLR" in found and "not a factor" in found


def test_report_never_writes_p_equals_less_than(text):
    assert not re.search(r"p = <", text) and "p < 0.001" in text


def test_report_numbers_are_the_statistics_numbers(world, text):
    a = world["res"].results.responses["d32"].anova
    section = text.split("### D32 in-focus (um)")[1].split("### Atomised")[0]
    row = next(l for l in section.splitlines() if l.startswith("| Silicone"))
    assert report.fmt_p(a.rows["sps"].p) in row and report.fmt(a.rows["sps"].ms) in row
    assert f"| error (pure error) | {report.fmt(a.error.ss)} | 18 |" in section


def test_a_modest_effect_that_fails_the_multiple_test_correction_says_so(world, text):
    found = text.split("## What was found")[1].split("## Design")[0]
    assert "does not survive the multiple-test correction" in found


def test_report_has_every_section_and_embeds_the_figures_it_has(text):
    for heading in ("## What was found", "## Design", "## Results by response", "## Droplet size spread",
                    "## Gas-to-liquid ratio (GLR)", "## Odd runs and frames", "## How to read the statistics",
                    "## Provenance"):
        assert heading in text, heading
    for fig in EXPECTED_FIGS:
        assert f"](figures/{fig}.png)" in text
    assert len(re.findall(r"^### ", text, re.M)) == len(stats.responses_table())


def test_report_explains_the_multiple_test_count_honestly(world, text):
    n = len(world["res"].results.multiple)
    assert f"{n} factor tests were run" in text and f"{n * 0.05:.1f} would come out" in text


def test_report_carries_the_glr_caveat_and_the_size_caveat(text):
    assert "cannot be separated from them" in text and "indicative" in text


def test_report_lists_what_was_left_out_and_why():
    with tempfile.TemporaryDirectory() as d:
        root = Path(d)
        fakes.make_measured_l9x3(root)
        extra = fakes.make_run(root / "x", "130000_3000sccm_300rpm_4000sps_or1.2_bh1", notes="Taguchi ReRun 1 - Repeat 4")
        design = load_design(root)
        design.update_runs(list(design.runs.values()) + [rd.load_run_info(extra, check_reuse=False)])
        design.rows[0].included = False
        res = stats.analyse(design, do_bootstrap=False)
        out = report.build(res, size_bins.from_results(res), provenance.collect(res, None, Path(d)))
        assert "## Runs left out" in out and "not measured yet" in out and "left out" in out


def test_report_for_an_unreplicated_design_names_its_weaker_error_term():
    with tempfile.TemporaryDirectory() as d:
        root = Path(d)
        fakes.make_measured_l9x3(root)
        design = load_design(root)
        a = dz.assign(design)
        for row in design.rows:
            if a.of(row).replicate != 1:
                row.included = False
        res = stats.analyse(design, do_bootstrap=False)
        assert res.error_source == "residual"
        out = report.build(res, None, provenance.collect(res, None, Path(d)))
        assert "no replicates" in out and "mixes noise with any interactions" in out
        assert "lack of fit" not in out.split("## Results by response")[1].split("## How to read")[0].lower().replace(
            "**lack of fit** compares", "")


def test_report_for_an_interaction_warns_that_interactions_are_likely():
    with tempfile.TemporaryDirectory() as d:
        root = Path(d)
        fakes.make_measured_l9x3(root, interaction=0.25, replicate_sd=0.01, per_frame=200)
        res = stats.analyse(load_design(root), do_bootstrap=False)
        out = report.build(res, None, provenance.collect(res, None, Path(d)))
        assert "interactions are likely" in out


def test_report_without_any_glr_says_so_and_invents_nothing(world):
    res = world["res"].results
    saved = [(r, r.context) for r in res.runs]
    try:
        for r in res.runs:
            r.context = type(r.context)(note="no GLR in run_summary.xlsx and no densities to compute one")
        res.trends = {}
        out = report.build(res, None, world["res"].prov)
        sec = out.split("## Gas-to-liquid ratio (GLR)")[1].split("## Odd runs")[0]
        assert "No run has a GLR" in sec and "Nothing was guessed" in sec and "| r |" not in sec
        assert "**GLR**" not in out.split("## Design")[0]
    finally:
        for r, c in saved:
            r.context = c


def test_report_for_a_refused_analysis_says_why_and_has_no_results():
    r = stats.Results([], [], {}, [("a", "left out")], [], refused="nothing to analyse: x")
    prov = provenance.Provenance("now", "abc", False, [], 0.3, 0, 0, [], [])
    out = report.build(r, None, prov)
    assert "## Not analysed" in out and "nothing to analyse: x" in out and "## Results" not in out


def test_report_flags_a_quantised_response_instead_of_trusting_its_p_values(world):
    res = world["res"].results
    d32 = res.responses["d32"]
    d32.quantisation_warning, d32.n_distinct = True, 3
    try:
        out = report.build(res, None, world["res"].prov)
        assert "pixel lattice" in out.split("## Design")[0]
        assert "p-values below are not meaningful" in out
    finally:
        d32.quantisation_warning, d32.n_distinct = False, 27


def test_markdown_tables_escape_pipes_and_newlines():
    t = report.md_table(["a|b", "c"], [["x|y", "line1\nline2"]])
    assert "a\\|b" in t and "x\\|y" in t and "line1 line2" in t


def test_provenance_states_what_is_known_and_marks_dirty_pipelines(world, text):
    prov = world["res"].prov
    rows = dict(prov.rows())
    assert rows["Sizer version(s)"] == "2.1.0" and rows["Score threshold"] == "0.3"
    assert rows["Bootstrap"] == "60 replicates, seed 0" and "∞" in rows["Size bin edges (µm)"]
    assert "| Sizer version(s) | 2.1.0 |" in text
    dirty = provenance.Provenance("now", "abc1234", True, ["2.1.0"], 0.3, 0, 0, [0, 25.0], [])
    assert "uncommitted edits to the pipeline" in dict(dirty.rows())["Repository revision"]
    assert dict(dirty.rows())["Bootstrap"] == "not run"


def test_provenance_reads_the_batch_settings_when_a_batch_ran_here(world):
    with tempfile.TemporaryDirectory() as d:
        jd = Path(d) / "_job"
        jd.mkdir()
        (jd / "job.json").write_text(json.dumps({"settings": {"stride": 10, "device": "cpu", "model_dir": "/m/dennis"},
                                                  "options": {"images": "all"}}))
        p = provenance.collect(world["res"].results, world["res"].bins, Path(d))
        r = dict(p.rows())
        assert r["Model directory"] == "/m/dennis" and r["Inference device"] == "cpu"
        assert r["Frame stride (batch)"] == "10" and r["Per-frame images"] == "every frame"


# ---- workbook ---------------------------------------------------------------------------------------------------------

def test_workbook_has_the_planned_sheets_all_within_excels_name_limit(world):
    import openpyxl
    wb = openpyxl.load_workbook(world["out"] / "taguchi_analysis.xlsx")
    names = wb.sheetnames
    for want in ("Per-run responses", "Design matrix", "Main effects", "ANOVA", "S-N ratios",
                 "Size bins (count)", "Size bins (volume)", "GLR context", "GLR trends", "Provenance"):
        assert want in names, want
    assert all(len(n) <= 31 for n in names)
    assert "Timings" not in names                        # no batch ran in this folder


def test_workbook_values_are_the_statistics_values(world):
    import openpyxl
    res = world["res"].results
    wb = openpyxl.load_workbook(world["out"] / "taguchi_analysis.xlsx")
    ws = wb["Per-run responses"]
    head = [c.value for c in ws[1]]
    assert ws.max_row == 28
    d32col = head.index("D32 in-focus (um)")
    got = {row[0].value: row[d32col].value for row in ws.iter_rows(min_row=2)}
    for r in res.runs:
        assert got[r.name] == pytest.approx(r.values["d32"], abs=1e-6)
    an = wb["ANOVA"]
    h = [c.value for c in an[1]]
    rows = [dict(zip(h, [c.value for c in row])) for row in an.iter_rows(min_row=2)]
    gas = next(r for r in rows if r["response"] == "D32 in-focus (um)" and r["source"] == "Gas flow (sccm)")
    assert gas["p"] == pytest.approx(res.responses["d32"].anova.rows["sccm"].p, abs=1e-7)
    assert gas["df"] == 2 and gas["q (BH)"] is not None


def test_workbook_size_bins_rows_sum_to_100(world):
    import openpyxl
    wb = openpyxl.load_workbook(world["out"] / "taguchi_analysis.xlsx")
    for sheet in ("Size bins (count)", "Size bins (volume)"):
        ws = wb[sheet]
        head = [c.value for c in ws[1]]
        first = head.index("<25")
        for row in ws.iter_rows(min_row=2, values_only=True):
            assert abs(sum(row[first:]) - 100) < 0.01


def test_workbook_includes_the_batch_timings_when_a_batch_ran(world):
    import openpyxl
    with tempfile.TemporaryDirectory() as d:
        out = Path(d)
        (out / "_job").mkdir()
        (out / "_job" / "timings.csv").write_text("run,status,inference_s\nrunA,ok,12.5\nrunB,ok,13.5\n")
        res = publish.publish(world["design"], out)
        assert res.ok
        ws = openpyxl.load_workbook(out / "taguchi_analysis.xlsx")["Timings"]
        assert [c.value for c in ws[1]] == ["run", "status", "inference_s"] and ws.max_row == 3


def test_workbook_text_that_looks_like_a_formula_stays_text():
    import openpyxl
    r = stats.Results([], [], {}, [("run1", "=1+1")], [])
    with tempfile.TemporaryDirectory() as d:
        prov = provenance.Provenance("now", "x", False, [], 0.3, 0, 0, [], [])
        tables_rows = {"skipped_runs": tables.skipped(r)}
        orig = tables.all_tables
        tables.all_tables = lambda results, bins: {**orig(results, bins), **tables_rows}
        try:
            workbook.write(Path(d) / "w.xlsx", r, None, prov)
        finally:
            tables.all_tables = orig
        cell = openpyxl.load_workbook(Path(d) / "w.xlsx")["Skipped runs"]["B2"]
        assert cell.value == "=1+1" and cell.data_type == "s"


def test_a_workbook_open_in_excel_gives_an_instruction_not_a_traceback(world, monkeypatch):
    def locked(src, dst):
        raise PermissionError("in use")
    from src.ai.Taguchi_Analysis_UI import fsops
    monkeypatch.setattr(fsops, "ATTEMPTS", 2)                # a real lock outlasts every retry
    monkeypatch.setattr(workbook.os, "replace", locked)
    with tempfile.TemporaryDirectory() as d:
        with pytest.raises(workbook.WorkbookError, match="close it in Excel"):
            workbook.write(Path(d) / "w.xlsx", world["res"].results, world["res"].bins, world["res"].prov)
        assert not list(Path(d).iterdir())               # and no temp file is left


# ---- flat CSVs ----------------------------------------------------------------------------------------------------------

def test_flat_csvs_match_the_tables_and_the_json_carries_provenance(world):
    out = world["out"] / "csv"
    with open(out / "per_run.csv", newline="", encoding="utf-8") as f:
        assert len(list(csv.DictReader(f))) == 27
    with open(out / "anova.csv", newline="", encoding="utf-8") as f:
        rows = list(csv.DictReader(f))
    assert len(rows) == len(tables.anova_rows(world["res"].results))
    j = json.loads((out / "results.json").read_text(encoding="utf-8"))
    assert j["provenance"]["Sizer version(s)"] == "2.1.0" and "anova" in j["tables"]
    assert len(j["multiple_comparisons"]) == len(world["res"].results.multiple)


# ---- odd-frame pack ----------------------------------------------------------------------------------------------------------

def odd_world():
    d = tempfile.TemporaryDirectory()
    root = Path(d.name)
    fakes.make_measured_l9x3(root / "src")
    res = stats.analyse(load_design(root / "src"), do_bootstrap=False)
    return d, root, res


def read_flags(result) -> list[dict]:
    with open(result.csv_path, newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def first_flagged(res, result):
    """(run record, frame) of the first frame flag_odd raised."""
    row = next(r for r in read_flags(result) if r["frame"])
    return next(r for r in res.runs if r.name == row["run"]), row["frame"]


def test_image_root_redirects_only_images_to_the_folder_that_exists():
    with tempfile.TemporaryDirectory() as d:
        cl = Path(d)
        root = odd._ImageRoot(cl)
        assert root / "images" == cl / "images"                       # nothing exists: unchanged
        (cl / "extreme_images").mkdir()
        assert root / "images" == cl / "extreme_images"
        (cl / "images").mkdir()
        assert root / "images" == cl / "images"                       # every-frame wins
        assert root / "classical_summary.json" == cl / "classical_summary.json"


def test_odd_pack_flags_and_explains_missing_images():
    d, root, res = odd_world()
    with d:
        out = root / "out"
        result = odd.make_pack(res, out)
        assert result.csv_path and result.n_flags > 0 and "d32_outlier" in result.by_category
        rows = list(csv.DictReader(open(result.csv_path, newline="", encoding="utf-8")))
        assert {r["images_copied"] for r in rows if r["category"] != "replicate_outlier"} == {odd.NO_IMAGE_EXTREMES}
        assert "none found" not in result.csv_path.read_text(encoding="utf-8")
        assert "extreme frames only" in result.note


def test_odd_pack_copies_the_image_the_stock_flag_odd_would_have_missed():
    d, root, res = odd_world()
    with d:
        run, frame = first_flagged(res, odd.make_pack(res, root / "out"))
        images = run.rec["cl_dir"] / "extreme_images"           # where the default mode keeps them
        images.mkdir()
        (images / f"{frame}.png").write_bytes(b"png")
        again = odd.make_pack(res, root / "out")
        hit = next(r for r in read_flags(again) if r["run"] == run.name and r["frame"] == frame)
        assert hit["images_copied"] == "classical"
        copied = list((root / "out" / "odd").rglob(f"{frame}_classical.png"))
        assert len(copied) == 1 and copied[0].read_bytes() == b"png"


def test_odd_pack_distinguishes_every_frame_runs_with_a_missing_image():
    d, root, res = odd_world()
    with d:
        run, _ = first_flagged(res, odd.make_pack(res, root / "out"))
        (run.rec["cl_dir"] / "images").mkdir()                   # an every-frame run, this frame missing
        rows = read_flags(odd.make_pack(res, root / "out"))
        mine = [r for r in rows if r["run"] == run.name and r["frame"]]
        others = [r for r in rows if r["run"] != run.name and r["frame"]]
        assert mine and {r["images_copied"] for r in mine} == {odd.NO_IMAGE}
        assert not others or {r["images_copied"] for r in others} == {odd.NO_IMAGE_EXTREMES}


def test_odd_pack_adds_replicate_outliers_which_flag_odd_cannot_know():
    d, root, res = odd_world()
    with d:
        victim = res.runs[3]
        res.responses["d32"].outliers.append({"run": victim.name, "condition": victim.condition, "value": 140.0,
                                              "condition_mean": 100.0, "z": 4.2})
        result = odd.make_pack(res, root / "out")
        rows = list(csv.DictReader(open(result.csv_path, newline="", encoding="utf-8")))
        rep = [r for r in rows if r["category"] == "replicate_outlier"]
        assert len(rep) == 1 and rep[0]["run"] == victim.name and "+4.2 SD" in rep[0]["why"]


def test_odd_pack_with_nothing_to_report_writes_nothing(monkeypatch):
    d, root, res = odd_world()
    with d:
        monkeypatch.setattr(spec.rdc_module("taguchi_analysis"), "flag_odd", lambda *a, **k: 0)
        result = odd.make_pack(res, root / "out")
        assert result.csv_path is None and "nothing out of the ordinary" in result.note
        assert not (root / "out" / "odd").exists()


def test_a_previous_odd_pack_is_replaced_not_mixed_in():
    d, root, res = odd_world()
    with d:
        out = root / "out"
        stale = out / "odd" / "090432" / "d32_outlier"
        stale.mkdir(parents=True)
        (stale / "frame_old_classical.png").write_bytes(b"old")
        odd.make_pack(res, out)
        assert not list((out / "odd").rglob("frame_old_classical.png"))


def test_a_folder_named_odd_holding_other_files_is_not_wiped():
    d, root, res = odd_world()
    with d:
        out = root / "out"
        (out / "odd").mkdir(parents=True)
        (out / "odd" / "my_notes.txt").write_text("keep me")
        odd.make_pack(res, out)
        assert (out / "odd" / "my_notes.txt").read_text() == "keep me"


def test_odd_pack_uses_the_batchs_timings_to_flag_slow_stages():
    d, root, res = odd_world()
    with d:
        out = root / "out"
        (out / "_job").mkdir(parents=True)
        lines = ["run,status,inference_s"] + [f"{r.name},ok,{100 if i else 900}" for i, r in enumerate(res.runs[:6])]
        (out / "_job" / "timings.csv").write_text("\n".join(lines) + "\n")
        result = odd.make_pack(res, out)
        assert result.by_category.get("slow_stage") == 1


# ---- found on the real 10/01 data (and the synthetic one) ----------------------------------------------------------------

def test_per_run_chart_gives_every_condition_of_an_l9_its_own_colour(world, captured):
    figures.fig_per_run(world["res"].results, Path("x"))
    ax = captured["fig_per_run"].axes[0]
    colours = {tuple(p.get_facecolor()) for p in ax.patches}
    assert len(colours) == 9                              # an L9 has 9 conditions; a 7-colour cycle repeated two


def test_a_campaign_with_no_glr_at_all_says_so_instead_of_counting_zero_runs(world, monkeypatch):
    from src.ai.Taguchi_Analysis_UI import covariates
    monkeypatch.setattr(covariates, "run_context",
                        lambda info: covariates.RunContext(note="no GLR in run_summary.xlsx and no densities"))
    res = stats.analyse(world["design"], do_bootstrap=False)
    assert res.warnings == ["no run has a GLR (no GLR in run_summary.xlsx and no densities), "
                            "so there is no GLR analysis."]
    assert not any("other 0" in w for w in res.warnings)


def test_a_refused_analysis_still_logs_the_runs_it_skipped_and_why(world):
    fresh = fakes.make_run(world["root"].parent / "fresh", "130000_3000sccm_300rpm_4000sps_or1.2_bh1")
    design = dz.detect([rd.load_run_info(fresh, check_reuse=False)])
    lines = []
    with tempfile.TemporaryDirectory() as d:
        res = publish.publish(design, Path(d), log=lambda t, l="info": lines.append(t))
    assert res.refused and any("skipping 130000_" in l and "not measured yet" in l for l in lines)


PUBLISHED = Path("/Volumes/LaCie/Experiments/Taguchi/First Taguchi Trial (RPM, SCCM, Silicone Flow)"
                 "/results_2026-10-01/taguchi_results.json")
REAL_01 = Path("/Volumes/LaCie/Experiments/2026/10/01")


@pytest.mark.skipif(not (PUBLISHED.is_file() and REAL_01.is_dir()), reason="LaCie not mounted")
def test_the_real_l9_published_end_to_end_matches_the_published_analysis():
    """Real 10/01 folders (legacy names, no GLR, two side runs) -> every deliverable, read-only."""
    import openpyxl
    pub = json.loads(PUBLISHED.read_text())
    design = dz.detect(rd.load_runs(rd.find_runs([REAL_01]).runs, check_reuse=False))
    for r in dz.outside_design(design):
        r.included = False
    snapshot = lambda: {str(p): (p.stat().st_size, p.stat().st_mtime_ns) for p in REAL_01.rglob("*") if p.is_file()}
    before = snapshot()
    with tempfile.TemporaryDirectory() as d:
        out = Path(d)
        res = publish.publish(design, out, options=BOTH)
        assert res.ok, res.problems
        assert len(res.results.runs) == 9 and res.results.error_source == "residual"
        # the workbook's ANOVA is the published ANOVA
        ws = openpyxl.load_workbook(out / "taguchi_analysis.xlsx")["ANOVA"]
        h = [c.value for c in ws[1]]
        rows = [dict(zip(h, [c.value for c in row])) for row in ws.iter_rows(min_row=2)]
        for key, label in (("d32", "D32 in-focus (um)"), ("atom", "Atomised fraction, classical (%)")):
            for f, flabel in (("sccm", "Gas flow (sccm)"), ("rpm", "Bubbler RPM"), ("sps", "Silicone (steps/s)")):
                row = next(r for r in rows if r["response"] == label and r["source"] == flabel)
                assert row["p"] == pytest.approx(pub["results"][key]["anova"][f]["p"], abs=1e-7)
                assert row["F"] == pytest.approx(pub["results"][key]["anova"][f]["F"], abs=1e-4)
        # the report quotes the same p the workbook holds
        text = (out / "taguchi_report.md").read_text(encoding="utf-8")
        p_gas = pub["results"]["atom"]["anova"]["sccm"]["p"]
        found = text.split("## What was found")[1].split("## Design")[0]
        assert report.p_eq(p_gas) in found
        # no GLR was recorded: said, not guessed
        assert "No run has a GLR" in text and not (out / "figures" / "fig_glr.png").exists()
        # the odd pack found real images (this campaign kept them) and every row says what it copied
        rows = list(csv.DictReader(open(out / "odd" / "odd_flags.csv", newline="", encoding="utf-8")))
        assert res.odd.n_images > 0 and rows and all(r["images_copied"] != "none found" for r in rows)
    assert snapshot() == before                           # the real campaign was only read


# ---- tightened after reviewing which wrong versions would still have passed --------------------------------------------

def test_the_count_chart_shows_counts_and_the_volume_chart_shows_volume(world, captured):
    bins = world["res"].bins
    figures.fig_size_spread_count(bins, Path("x"))
    figures.fig_size_spread_volume(bins, Path("x"))
    first = sorted(bins.runs, key=lambda s: (bins.run_condition[s.name], s.name))[0]
    assert first.pct_count[1] != pytest.approx(first.pct_volume[1])        # the two really differ here
    for name, want in (("fig_size_spread_count", first.pct_count), ("fig_size_spread_volume", first.pct_volume)):
        ax = captured[name].axes[0]
        stack = sorted((p for p in ax.patches if round(p.get_x() + p.get_width() / 2) == 0),
                       key=lambda p: p.get_y())                           # the first run's bars, bottom to top
        assert [p.get_height() for p in stack] == pytest.approx(list(want))


def test_only_effects_that_fail_the_correction_carry_the_warning(world, text):
    found = text.split("## What was found")[1].split("## Design")[0]
    d32 = next(l for l in found.splitlines() if "D32" in l)
    atom = next(l for l in found.splitlines() if "Atomised" in l)
    q = {(rk, fk): qq for rk, fk, _p, qq in world["res"].results.multiple}
    assert q[("d32", "sps")] > 0.05 and "does not survive" in d32
    assert q[("atom", "sps")] < 0.05 and "does not survive" not in atom
    assert d32.count("does not survive") == 1                              # gas flow (q tiny) is not tagged


def test_p_value_wording_thresholds():
    assert report.p_eq(0.0005) == "p < 0.001" and report.p_eq(0.001) == "p = 0.001"
    assert report.p_eq(0.0276) == "p = 0.028" and report.p_eq(None) == "p = –"
    assert report.fmt_p(0.0005) == "<0.001" and report.fmt_p(0.5) == "0.500" and report.fmt_p(None) == "–"


def test_workbook_q_values_are_the_benjamini_hochberg_values(world):
    import openpyxl
    ws = openpyxl.load_workbook(world["out"] / "taguchi_analysis.xlsx")["ANOVA"]
    h = [c.value for c in ws[1]]
    rows = [dict(zip(h, [c.value for c in row])) for row in ws.iter_rows(min_row=2)]
    q = {(rk, fk): qq for rk, fk, _p, qq in world["res"].results.multiple}
    label = {f.key: f.label for f in world["res"].results.factors}
    names = {k: r.label for k, r in world["res"].results.responses.items()}
    for (rk, fk), want in q.items():
        row = next(r for r in rows if r["response"] == names[rk] and r["source"] == label[fk])
        assert row["q (BH)"] == pytest.approx(want, abs=1e-7) and row["q (BH)"] >= row["p"]
