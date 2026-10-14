"""More extreme frames (CCD follow-up phase 10, 2026-10-13): beside the D32 / atomised extremes, each run's
frame with the most in-focus droplets, the most droplets in total, and the largest / smallest mean in-focus
diameter, drawn into droplets_<thr>/extreme_frames_more/ with an index.json.

The pick (measure_run.pick_more_extremes) is shared by measure_run and the back-fill
(render_more_extremes.py), which re-measures only the chosen frames from the saved predictions and refuses a
run whose re-measure disagrees with per_frame.csv. A synthetic run is measured by the real measure_run.py, so
the back-fill is checked against what a new measurement writes, byte for byte."""
from __future__ import annotations

import csv
import hashlib
import json
import shutil
import subprocess
import sys
from pathlib import Path
from types import SimpleNamespace

import cv2
import numpy as np
import pytest
from pycocotools import mask as mask_util

from src.ai.Taguchi_Analysis_UI import deck, report
from src.ai.Taguchi_Analysis_UI import figures_extra as fx
from src.ai.Taguchi_Analysis_UI import pipeline_spec as spec

mr = spec.rdc_module("measure_run")
rme = spec.rdc_module("render_more_extremes")

KEYS = ["most_in_focus", "most_total", "mean_d_highest", "mean_d_lowest"]


def row(frame, n_in, n_out=0, d32=50.0, atom=10.0):
    return {"frame": frame, "droplets_in_focus": n_in, "droplets_out_of_focus": n_out,
            "d32_in_focus_um": d32, "atomised_pct": atom}


# ---- the pick ---------------------------------------------------------------------------------------------

def test_keys_labels_and_units_are_the_agreed_ones():
    frames, note = mr.pick_more_extremes([row("a", 20, 5), row("b", 30, 1)], {"a": 30.0, "b": 40.0})
    assert note == ""
    assert [(e["key"], e["label"], e["unit"]) for e in frames] == [
        ("most_in_focus", "Most in-focus droplets", "in-focus droplets"),
        ("most_total", "Most droplets in total", "droplets"),
        ("mean_d_highest", "Largest mean diameter", "um mean (in focus)"),
        ("mean_d_lowest", "Smallest mean diameter", "um mean (in focus)")]
    assert all(list(e) == list(mr.MORE_INDEX_FIELDS) for e in frames)
    e = frames[1]
    assert (e["frame"], e["value"], e["droplets_in_focus"], e["droplets_total"]) == ("b", 31, 30, 31)


def test_a_sparse_frame_cannot_be_the_mean_extreme():
    rows = [row("one_speck", 1), row("nine", 9), row("ten", 10), row("busy", 50)]
    mean_d = {"one_speck": 5.0, "nine": 8.0, "ten": 30.0, "busy": 40.0}
    frames, note = mr.pick_more_extremes(rows, mean_d)
    by = {e["key"]: e for e in frames}
    assert by["mean_d_lowest"]["frame"] == "ten" and by["mean_d_lowest"]["value"] == 30.0
    assert by["mean_d_highest"]["frame"] == "busy" and note == ""
    assert mr.MIN_DROPLETS_FOR_MEAN == 10


def test_the_floor_is_half_the_runs_median_count():
    """Ben, 2026-10-13: with a flat 10 the largest-mean frame on 130918 was 57.9 um from 17 droplets (the run's
    frames have 54 in a median frame); at half the median it is the 34-droplet frame."""
    rows = [row(f"f{i}", 54) for i in range(5)] + [row("thin", 17), row("fuller", 34)]
    mean_d = {**{f"f{i}": 39.0 for i in range(5)}, "thin": 57.9, "fuller": 55.1}
    assert mr.mean_floor(rows) == 27
    by = {e["key"]: e for e in mr.pick_more_extremes(rows, mean_d)[0]}
    assert by["mean_d_highest"]["frame"] == "fuller"
    by10 = {e["key"]: e for e in mr.pick_more_extremes(rows, mean_d, floor=10)[0]}
    assert by10["mean_d_highest"]["frame"] == "thin"                     # the old rule
    assert mr.mean_floor([row("a", 8), row("b", 12)]) == 10              # never below 10
    assert mr.mean_floor([]) == 10


def test_ties_go_to_the_first_frame():
    rows = [row("f0", 20, 5), row("f1", 20, 5), row("f2", 10, 15)]
    frames, _ = mr.pick_more_extremes(rows, {"f0": 30.0, "f1": 30.0, "f2": 30.0})
    assert [e["frame"] for e in frames] == ["f0", "f0", "f0", "f0"]


@pytest.mark.parametrize("n_ok", [0, 1])
def test_fewer_than_two_qualifying_frames_leave_out_the_mean_entries_and_say_why(n_ok):
    rows = [row("f0", 12 if n_ok else 3), row("f1", 3)]
    frames, note = mr.pick_more_extremes(rows, {"f0": 30.0, "f1": 10.0})
    assert [e["key"] for e in frames] == ["most_in_focus", "most_total"]
    assert f"{n_ok} frame(s)" in note and str(mr.MIN_DROPLETS_FOR_MEAN) in note


def test_no_rows_no_frames():
    assert mr.pick_more_extremes([], {}) == ([], "no frames measured")


def test_frame_mean_reads_back_from_the_csv_identically():
    areas = [13.0, 29.0, 113.0, 7.0]
    ds = [mr.csv_um(a) for a in areas]
    back = [float(str(d)) for d in ds]
    assert mr.frame_mean_d(ds) == mr.frame_mean_d(back[::-1])
    assert np.isnan(mr.frame_mean_d([]))


# ---- a synthetic run, measured by the real measure_run.py -------------------------------------------------

H, W, BG = 120, 240, 40000.0
# (in-focus count, radius px, out-of-focus count): f1 and f4 tie on in-focus droplets (f1 must win), f2 has the
# most in total and the largest droplets, f3 has the smallest but only 3 in focus (guarded out), so f0 is the
# smallest mean.
DESIGN = [(12, 2, 2), (15, 3, 0), (11, 6, 10), (3, 1, 1), (15, 4, 0)]
EXPECT = {"most_in_focus": 1, "most_total": 2, "mean_d_highest": 2, "mean_d_lowest": 0}


def _rle(m):
    r = mask_util.encode(np.asfortranarray(m.astype(np.uint8)))
    r["counts"] = r["counts"].decode("ascii")
    return r


def _det(img_id, m, cat):
    r = _rle(m)
    x, y, w, h = (float(v) for v in mask_util.toBbox(r))
    return {"image_id": img_id, "category_id": cat, "score": 0.9, "segmentation": r,
            "area": float(mask_util.area(r)), "bbox": [x, y, w, h]}


def _frame(img_id, n_in, r_in, n_out):
    img = np.full((H, W), BG, np.float32)
    dets = []
    slots = [(x, y) for y in range(12, 96, 22) for x in range(12, W - 8, 22)]
    for k in range(n_in + n_out):
        m = np.zeros((H, W), np.uint8)
        cv2.circle(m, slots[k], r_in if k < n_in else 3, 1, -1)
        img[m > 0] = BG * (0.2 if k < n_in else 0.85)
        dets.append(_det(img_id, m, mr.DROPLET))
    fil = np.zeros((H, W), np.uint8)
    fil[104:112, 20:220] = 1                              # one filament, so the atomised % is a real ratio
    img[fil > 0] = BG * 0.3
    dets.append(_det(img_id, fil, mr.FILAMENT))
    return img, dets


def _build_run(root: Path) -> Path:
    run = root / "120000_9000sccm_500rpm_4000sps_or1.2_bh0.5"
    raw = run / "shadowgraph" / "raw"
    for d in (raw / "frames" / "16bit", raw / "frames" / "8bit", run / "shadowgraph" / "analysis"):
        d.mkdir(parents=True, exist_ok=True)
    cv2.imwrite(str(raw / "background_median.tiff"), np.full((H, W), BG, np.uint16))
    images, preds = [], []
    for i, (n_in, r_in, n_out) in enumerate(DESIGN):
        stem = f"frame_{i:04d}_n{i * 10 - 1}"
        img, dets = _frame(i + 1, n_in, r_in, n_out)
        cv2.imwrite(str(raw / "frames" / "16bit" / f"{stem}.tiff"), img.astype(np.uint16))
        cv2.imwrite(str(raw / "frames" / "8bit" / f"{stem}.png"), (img / BG * 255).astype(np.uint8))
        images.append({"id": i + 1, "file_name": f"{stem}.png", "width": W, "height": H})
        preds += dets
    (raw / "instances.json").write_text(json.dumps({"images": images}))
    (run / "shadowgraph" / "analysis" / "predictions.json").write_text(json.dumps(preds))
    return run


def _measure(run: Path, out: Path, images: str | None):
    raw = run / "shadowgraph" / "raw"
    cmd = [sys.executable, "-W", "ignore", str(Path(mr.__file__)), "--pred",
           str(run / "shadowgraph" / "analysis" / "predictions.json"), "--root", str(run),
           "--val-dir", "shadowgraph/raw", "--background", str(raw / "background_median.tiff"),
           "--sixteen-bit-dir", str(raw / "frames" / "16bit"), "--score-thresh", "0.30",
           "--out-dir", str(out), "--n-boot", "20"] + (["--images", images] if images else [])
    p = subprocess.run(cmd, capture_output=True, text=True)
    assert p.returncode == 0, p.stdout[-2000:] + p.stderr[-2000:]
    return p.stdout


@pytest.fixture(scope="module")
def measured(tmp_path_factory):
    root = tmp_path_factory.mktemp("more_extremes")
    run = _build_run(root)
    meas = run / "shadowgraph" / "analysis" / "droplets_0.30"
    out = _measure(run, meas, "extremes")
    return {"run": run, "meas": meas, "stdout": out}


@pytest.fixture
def copy(measured, tmp_path):
    """A private copy of the measured run, so a test can tamper with it."""
    dst = tmp_path / measured["run"].name
    shutil.copytree(measured["run"], dst)
    return dst, dst / "shadowgraph" / "analysis" / "droplets_0.30"


def _tree(root: Path, skip: str) -> dict[str, str]:
    return {str(p.relative_to(root)): hashlib.sha256(p.read_bytes()).hexdigest()
            for p in sorted(root.rglob("*")) if p.is_file() and skip not in p.parts}


def test_measure_run_picks_draws_and_indexes_the_four_frames(measured):
    meas = measured["meas"]
    idx = json.loads((meas / mr.MORE_EXTREMES_DIR / "index.json").read_text())
    assert list(idx) == ["schema", "min_droplets_for_mean", "note", "frames"]
    assert (idx["schema"], idx["min_droplets_for_mean"], idx["note"]) == (1, 10, "")
    stems = [f"frame_{i:04d}_n{i * 10 - 1}" for i in range(len(DESIGN))]
    assert [e["key"] for e in idx["frames"]] == KEYS
    assert {e["key"]: e["frame"] for e in idx["frames"]} == {k: stems[i] for k, i in EXPECT.items()}
    for e in idx["frames"]:
        assert e["image"] == f"{e['frame']}.png" and (meas / mr.MORE_EXTREMES_DIR / e["image"]).is_file()
    assert idx["frames"][0]["value"] == 15 and idx["frames"][1]["value"] == 21
    summary = json.loads((meas / "summary.json").read_text())
    assert summary["more_extreme_frames"] == idx["frames"]
    with open(meas / "per_frame.csv", newline="") as f:
        rows = list(csv.reader(f))
    assert rows[0][-1] == "mean_d_in_focus_um" and rows[0][-2] == "blob_area_px"
    with open(meas / mr.MORE_EXTREMES_DIR / "index.csv", newline="") as f:
        assert [r["key"] for r in csv.DictReader(f)] == KEYS
    assert "Most in-focus droplets" in measured["stdout"] and "Smallest mean diameter" in measured["stdout"]
    assert not (meas / f"{mr.MORE_EXTREMES_DIR}.partial").exists()


def test_without_images_the_index_is_still_written(measured, tmp_path):
    out = tmp_path / "droplets_0.30"
    _measure(measured["run"], out, None)
    idx = json.loads((out / mr.MORE_EXTREMES_DIR / "index.json").read_text())
    assert [e["key"] for e in idx["frames"]] == KEYS and all(e["image"] is None for e in idx["frames"])
    assert not list((out / mr.MORE_EXTREMES_DIR).glob("*.png"))


def test_the_backfill_reproduces_measure_run_and_writes_nothing_else(measured, copy):
    run, meas = copy
    shutil.rmtree(meas / mr.MORE_EXTREMES_DIR)
    before = _tree(run, mr.MORE_EXTREMES_DIR)
    status, msg = rme.backfill_run(run, 0.30)
    assert status == "written", msg
    assert "gate passed" in msg
    assert _tree(run, mr.MORE_EXTREMES_DIR) == before
    made = measured["meas"] / mr.MORE_EXTREMES_DIR
    for f in made.iterdir():
        assert (meas / mr.MORE_EXTREMES_DIR / f.name).read_bytes() == f.read_bytes(), f.name
    assert sorted(p.name for p in (meas / mr.MORE_EXTREMES_DIR).iterdir()) == sorted(p.name for p in made.iterdir())


def test_the_backfill_refuses_a_run_whose_remeasure_disagrees(copy):
    run, meas = copy
    shutil.rmtree(meas / mr.MORE_EXTREMES_DIR)
    p = meas / "per_frame.csv"
    with open(p, newline="") as f:
        rows = list(csv.DictReader(f))
    rows[1]["d32_in_focus_um"] = f"{float(rows[1]['d32_in_focus_um']) + 0.5:.2f}"    # the most-in-focus frame
    with open(p, "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]))
        w.writeheader()
        w.writerows(rows)
    tampered = p.read_bytes()
    status, msg = rme.backfill_run(run, 0.30)
    assert status == "failed" and "gate failed" in msg and rows[1]["frame"] in msg and "D32" in msg
    assert not (meas / mr.MORE_EXTREMES_DIR).exists()
    assert not (meas / f"{mr.MORE_EXTREMES_DIR}.partial").exists()
    assert p.read_bytes() == tampered


def test_an_existing_folder_is_kept_unless_forced(copy):
    run, meas = copy
    idx = meas / mr.MORE_EXTREMES_DIR / "index.json"
    was = idx.read_bytes()
    assert rme.backfill_run(run, 0.30)[0] == "exists"
    assert rme.backfill_run(run, 0.30, dry=True)[0] == "exists"
    status, msg = rme.backfill_run(run, 0.30, force=True)
    assert status == "written", msg
    assert idx.read_bytes() == was


def test_dry_run_writes_nothing(copy):
    run, meas = copy
    shutil.rmtree(meas / mr.MORE_EXTREMES_DIR)
    status, msg = rme.backfill_run(run, 0.30, dry=True)
    assert status == "dry" and "most_in_focus frame_0001_n9" in msg
    assert not (meas / mr.MORE_EXTREMES_DIR).exists()


# ---- the app's gallery ------------------------------------------------------------------------------------

def _fake_results(root: Path, with_index: tuple[bool, ...]):
    from PIL import Image
    runs = []
    for i, has in enumerate(with_index):
        run = root / f"12{i}000_run"
        meas = run / "shadowgraph" / "analysis" / "droplets_0.30"
        meas.mkdir(parents=True)
        if has:
            rows = [row("frame_0000_n-1", 20, 5), row("frame_0001_n9", 12, 30), row("frame_0002_n19", 15, 0)]
            frames, note = mr.pick_more_extremes(rows, {"frame_0000_n-1": 40.0, "frame_0001_n9": 55.5,
                                                        "frame_0002_n19": 30.25})
            for e in frames:
                e["image"] = f"{e['frame']}.png"
            mr.write_more_extremes(meas, frames, note,
                                   lambda stem, path: Image.new("RGB", (64, 36), (180, 180, 180)).save(path))
        runs.append(SimpleNamespace(name=run.name, path=run, rec={}, condition="T1", replicate=i + 1,
                                    levels={"gas": 9000}, values={}))
    return SimpleNamespace(thr=0.30, runs=runs, factors=[SimpleNamespace(key="gas", label="Gas flow (sccm)")])


def test_the_gallery_is_drawn_from_runs_with_an_index_and_skips_the_rest(tmp_path):
    from PIL import Image
    res = _fake_results(tmp_path, (True, False, True))
    files = fx.galleries(res, tmp_path / "figs")
    assert [p.name for p in files] == ["gallery_more_T1.png"]
    with Image.open(files[0]) as im:
        w = fx.GALLERY_WIDTH
        assert im.width == 4 * w + 5 * 8
        assert im.height == 34 + 2 * (round(36 * w / 64) + 22 + 8) + 8          # two runs, one row each
    ex = fx._more_frames(res.runs[0], 0.30)
    assert set(ex) == set(KEYS) and ex["mean_d_lowest"][1]["value"] == 30.25
    title, detail, *_ = deck.tw_gallery_more(res, "T1")
    assert "smallest mean 30.25 um" in detail and "most in total 42" in detail and "120000" in detail


def test_no_index_no_gallery(tmp_path):
    res = _fake_results(tmp_path, (False, False))
    assert fx.galleries(res, tmp_path / "figs") == []
    assert not (tmp_path / "figs").exists()


def test_the_report_lists_the_two_kinds_of_sheet_separately():
    text = report._gallery_section(None, ["figs/galleries/gallery_T1.png", "figs/galleries/gallery_more_T1.png"])
    assert "- [T1](figures/galleries/gallery_T1.png)" in text
    assert "- [T1](figures/galleries/gallery_more_T1.png)" in text
    assert text.index("lowest- and highest-atomised") < text.index("most in-focus droplets")
    only_more = report._gallery_section(None, ["figs/galleries/gallery_more_T2.png"])
    assert "highest-atomised" not in only_more and "gallery_more_T2.png" in only_more


def test_each_run_gets_a_2x2_slide_that_fills_the_page(tmp_path):
    """Ben, 2026-10-13: the busiest-frame slides as a 2x2 grid taking up most of the slide, one per run."""
    from PIL import Image
    from pptx import Presentation
    from pptx.util import Emu
    res = _fake_results(tmp_path, (True, False, True))
    for r in res.runs:                                    # frame-shaped images (16:9) for a realistic layout
        d = r.path / "shadowgraph" / "analysis" / "droplets_0.30" / mr.MORE_EXTREMES_DIR
        for png in d.glob("*.png") if d.is_dir() else []:
            Image.new("RGB", (2048, 1152), (200, 200, 200)).save(png)
    res.factors = []
    title, cells, notes = deck.tw_more_run(res, res.runs[0])
    assert [c.split(":")[0] for _, c in cells] == ["Most in-focus droplets", "Most droplets in total",
                                                    "Smallest mean diameter", "Largest mean diameter"]
    assert "30.25 µm from 15 in-focus droplets" in cells[2][1] and "half the run's median" in notes
    assert deck.tw_more_run(res, res.runs[1]) is None                        # no index: no slide
    dk = deck._Deck("footer")
    assert deck.more_grid_slides(dk, res) == 2
    path = dk.save(tmp_path / "g.pptx")
    prs = Presentation(str(path))
    W, H = Emu(prs.slide_width).inches, Emu(prs.slide_height).inches
    for s in prs.slides:
        pics = [sh for sh in s.shapes if sh.shape_type == 13]
        assert len(pics) == 4
        xs = sorted({round(Emu(p.left).inches, 2) for p in pics}); ys = sorted({round(Emu(p.top).inches, 2) for p in pics})
        assert len(xs) == 2 and len(ys) == 2                                  # a 2x2 grid
        assert all(Emu(p.width).inches > 3.2 for p in pics)
        right = max(Emu(p.left + p.width).inches for p in pics); low = max(Emu(p.top + p.height).inches for p in pics)
        assert right <= W and low <= 5.2 and (right - min(xs)) > 0.65 * W      # most of the page, inside it
        texts = [sh.text_frame.text for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
        assert len(texts) >= 3 and s.notes_slide.notes_text_frame.text.strip()


def test_grid_frames_are_shrunk_for_the_deck(tmp_path):
    """The full-size frame PNGs are ~4 MB: embedded as they were, the CCD deck reached 457 MB."""
    from PIL import Image
    big = tmp_path / "f.png"
    Image.effect_noise((2048, 1152), 60).convert("RGB").save(big)
    buf = deck._slide_jpeg(big)
    with Image.open(buf) as im:
        assert im.format == "JPEG" and im.size == (deck.GRID_PX, round(1152 * deck.GRID_PX / 2048))
    assert len(buf.getvalue()) < big.stat().st_size / 3


def test_the_rate_slide_states_the_band_and_the_four_steps():
    res = SimpleNamespace(opts=SimpleNamespace(band_um=(20.0, 60.0)), runs=[SimpleNamespace(values={"rate": 1.0})])
    cards = deck.rate_cards(res)
    assert deck.has_rate(res) and [h for h, _ in cards][:4] == [
        "1. Silicone in", "2. Share made into droplets", "3. Volume of a typical droplet", "4. In the band (20-60 um)"]
    assert "20 um up to (not including) 60 um" in cards[3][1] and all(len(b) < 300 for _, b in cards)
    assert not deck.has_rate(SimpleNamespace(runs=[SimpleNamespace(values={"rate": None})]))


def test_every_goal_gets_maps():
    from src.ai.Taguchi_Analysis_UI import rsm_figures as rf
    R = SimpleNamespace(goal_status=[SimpleNamespace(key="rate_band"), SimpleNamespace(key="atom")],
                        goals=[], fits={k: 1 for k in ("rate_band", "atom", "d32", "lag1", "peak_bar")})
    assert rf.mapped_responses(R) == ["rate_band", "atom", "d32", "lag1", "peak_bar"]


# ---- drawn the classical way (2026-10-14) -----------------------------------------------------------------
dmc = spec.rdc_module("draw_more_classical")


def _classical(run: Path, out: Path):
    cmd = [sys.executable, "-W", "ignore", str(Path(dmc.__file__).with_name("classical_liquid.py")), "--root",
           str(run), "--out-dir", str(out), "--score-thresh", "0.30", "--images-mode", "extremes"]
    p = subprocess.run(cmd, capture_output=True, text=True)
    assert p.returncode == 0, p.stdout[-2000:] + p.stderr[-2000:]


def test_the_classical_pass_draws_the_busiest_frames_and_the_deck_uses_them(copy):
    """classical_liquid.py draws measure_run's busiest / mean-diameter frames too (un-atomised liquid filled),
    lists them in its summary, and the gallery, the 2x2 slides and the GUI read them in preference."""
    run, meas = copy
    liq = run / "shadowgraph" / "analysis" / "liquid_0.30"
    _classical(run, liq)
    idx = json.loads((meas / mr.MORE_EXTREMES_DIR / "index.json").read_text())
    want = sorted({f["frame"] for f in idx["frames"]})
    assert json.loads((liq / "classical_summary.json").read_text())["more_extreme_frames"] == want
    assert all((liq / "extreme_images" / f"{f}.png").is_file() for f in want)
    r = SimpleNamespace(name=run.name, path=run, rec={}, condition="T1", replicate=1, levels={"gas": 9000}, values={})
    ex = fx._more_frames(r, 0.30)
    assert set(ex) == set(KEYS) and fx.more_drawn_classical(ex)
    assert all(p.parent == liq / "extreme_images" for p, _e in ex.values())
    res = SimpleNamespace(thr=0.30, runs=[r], factors=[SimpleNamespace(key="gas", label="Gas flow (sccm)")])
    _title, _cells, notes = deck.tw_more_run(res, r)
    assert deck.CLASSICAL_DRAWN in notes and "model-only" not in notes
    (liq / "extreme_images" / f"{ex['most_total'][1]['frame']}.png").unlink()      # one frame missing: mixed
    ex2 = fx._more_frames(r, 0.30)
    assert ex2["most_total"][1]["drawn"] == "measure_run" and deck._drawn_text([ex2]) == deck.MIXED_DRAWN


def test_the_classical_backfill_draws_only_and_refuses_a_run_whose_numbers_move(copy, tmp_path):
    run, meas = copy
    liq = run / "shadowgraph" / "analysis" / "liquid_0.30"
    _classical(run, liq)
    want = json.loads((liq / "classical_summary.json").read_text())["more_extreme_frames"]
    made = {f: (liq / "extreme_images" / f"{f}.png").read_bytes() for f in want}
    s = json.loads((liq / "classical_summary.json").read_text())
    s.pop("more_extreme_frames")                                  # as a run measured before 2026-10-14
    (liq / "classical_summary.json").write_text(json.dumps(s, indent=2))
    for f in want:
        (liq / "extreme_images" / f"{f}.png").unlink()
    before = _tree(run, "extreme_images")
    assert s["window"]["objects"] == []                                # no window drops: the fast path
    for full in (False, True):                                        # both paths draw the same images
        msg = dmc.one(run, 0.30, False, False, tmp_path, full=full)
        assert "numbers unchanged" in msg and ("full pass" if full else "fast") in msg, msg
        assert {f: (liq / "extreme_images" / f"{f}.png").read_bytes() for f in want} == made
        assert _tree(run, "extreme_images") == before                 # nothing else in the run written
        assert "already drawn" in dmc.one(run, 0.30, False, False, tmp_path)
        for f in want:
            (liq / "extreme_images" / f"{f}.png").unlink()
    rows = (liq / "classical_per_frame.csv").read_text()               # a frame's row moved: both paths refuse
    first = want[0]
    line = next(ln for ln in rows.splitlines() if ln.startswith(first + ","))
    cells = line.split(",")
    cells[4] = f"{float(cells[4]) + 0.5:g}"                            # atomised_pct
    (liq / "classical_per_frame.csv").write_text(rows.replace(line, ",".join(cells)))
    for full in (False, True):
        msg = dmc.one(run, 0.30, False, False, tmp_path, full=full)
        assert "FAILED the gate" in msg and f"{first} atomised_pct" in msg, msg
    (liq / "classical_per_frame.csv").write_text(rows)
    s["atomised_pct_pooled"] += 0.5                                    # the pooled summary: the full pass checks it
    (liq / "classical_summary.json").write_text(json.dumps(s, indent=2))
    msg = dmc.one(run, 0.30, False, False, tmp_path, full=True)
    assert "FAILED the gate" in msg and "atomised_pct_pooled" in msg
    assert not any((liq / "extreme_images" / f"{f}.png").exists() for f in want)
