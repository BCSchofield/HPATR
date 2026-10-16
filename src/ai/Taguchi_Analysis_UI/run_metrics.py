"""run_metrics.py -- everything the Taguchi and response-surface analyses need from ONE run,
beyond the nine responses taguchi_analysis.load_run already gives. No Qt.

    size        in-focus droplet sizes: D10, D32, Dv10 / Dv50 / Dv90, span, % volume
                below 50 um and above 125 um, count and volume histograms
    steadiness  per-frame droplet count (ALL droplets, in and out of focus, so it does not
                depend on the sizer version) and liquid area: CVs, % frames below half the
                mean count, lag correlations, Welch spectrum and the "slow share" (< 2 Hz)
    pressure    the logged upstream pressure and gas flow inside the camera window
                (run_summary.xlsx), and how droplets per frame follow the pressure
    pieces      un-atomised liquid pieces (classical_components.csv): median, largest, per frame
    quality     sizer version, % moved by the 2.2.0 rule, in/out-of-focus counts,
                detections per frame, background brightness and dark specks, static detections
    throughput  atomised liquid flow = liquid mL/min x atomised fraction
    when        capture time (date folders + HHMMSS), for drift checks

Every part is independent: a missing input leaves that part None and says why in
`missing`, never stopping the rest. Results are plain JSON-able dicts and are cached per run
(the cache is keyed on the inputs' sizes and modification times).

The five values the statistics analyse are pulled out by `response_values()`; see
EXTRA_RESPONSES. Everything else is descriptive data for figures and tables.
"""
from __future__ import annotations

import csv
import hashlib
import json
import math
import re
from datetime import datetime
from pathlib import Path

import numpy as np

from . import fsops

SCHEMA = 2                        # 2: grouped-data Dv percentiles (2026-10-11)
SLOW_HZ = 2.0                     # "slow pulsing": variance below this frequency
SMALL_UM, LARGE_UM = 50.0, 125.0  # % volume below / above (the CCD plan's size-spread response)
FRAME_PX = 2048 * 1152            # used only when the frame size is not recorded
STATIC_FRAC = 0.20                # a detection spot present in >= this share of frames is static
STATIC_GRID_PX = 6

# (key, label, direction, quantised) -- the extra responses the statistics analyse, in the
# same shape as taguchi_analysis.RESPONSES. Kept few on purpose: every extra response is
# another chance of a false "significant" (the CCD plan's warning). The two volume-based ones
# (liquid stage 2.3.0, 2026-10-13) are SECONDARY: they rank runs like the area fraction but are
# ~3x noisier and good to about +/-40% in scale. Runs measured before 2.3.0 have no volume
# columns, so these are blank for them. The in-band rate (below) rests on the volume fraction and
# is the default goal from 2026-10-14 (rsm_responses.DEFAULT_GOALS).
EXTRA_RESPONSES = [
    ("dv50", "Dv50, volume median, in-focus (um)", "smaller", False),
    ("af_flow", "Atomised liquid flow index, area-based (mL/min)", "larger", False),
    ("thin_pct", "Frames below half the mean droplet count (%)", "smaller", False),
    ("lag1", "Droplet-count lag-1 correlation (pulsing)", "smaller", False),
    ("peak_bar", "Peak pressure (barA)", "smaller", False),
    ("atom_vol", "Atomised fraction by volume, estimated (%)", "larger", False),
    ("af_flow_vol", "Atomised liquid flow index, volume-based (mL/min)", "larger", False),
    # how MANY droplets of the size you want, not just their size (Ben, 2026-10-13): mean diameter
    # cannot tell 5 droplets at 40 um from 100,000. A CONCENTRATION (droplets in view per frame), not
    # a production rate: it understates fast, high-gas sprays; the 2.2.0 sharpness rule removes more
    # droplets at high gas; below 50 um sizes sit near the pixel lattice. Never a default goal.
    ("band_per_frame", "In-focus droplets per frame in the size band", "larger", False),
    # a production RATE by mass balance (2026-10-13; docs/notes/2026-10-13_droplet_rate_investigation.md):
    # droplets made per second = syringe flow x atomised fraction BY VOLUME / mean in-focus droplet volume,
    # pi/6 <d^3>. Unlike the per-frame counts it does not fall when the spray speeds up: on the CCD day the
    # in-focus count was flat along the 1.2 mm gas axis while this doubled. Rests on the volume estimate
    # (+/-40% absolute, ranking good) and on the in-focus sizes standing for every droplet; <d^3> leans on
    # the largest, worst-sampled droplets. Liquid stage 2.3.0 only. `rate_band` is the default goal
    # (importance 3, with atomised fraction 1; Ben, 2026-10-14: the spray is collected and size-filtered,
    # so throughput in the band is the objective).
    ("rate", "Droplets produced per second, mass balance (estimated)", "larger", False),
    ("rate_band", "Droplets produced per second in the size band, mass balance (estimated)", "larger", False),
]
DEFAULT_BAND_UM = (20.0, 60.0)     # the size Ben collects (2026-10-14; was 30-50). Both edges fall between
                                   # pixel-lattice diameters (3/4 and 28/29 px of area at 10 um/px), so the
                                   # band starts at 4 px (22.6 um): the 1-3 px droplets (11-20 um) are left out


def band_counts(rec: dict | None, band=DEFAULT_BAND_UM):
    """In-focus droplets per frame with lo <= d < hi (half-open, as size_bins), in the run's frame
    order (frames with none count as 0), or None without the per-frame diameters."""
    if not rec or rec.get("dia") is None or not rec.get("frames"):
        return None
    lo, hi = band
    out = []
    for f in rec["frames"]:
        d = np.asarray(rec["dia"].get(f, []), float)
        out.append(int(((d >= lo) & (d < hi)).sum()))
    return np.array(out, float)


def _rate_parts(rec: dict | None, band=DEFAULT_BAND_UM):
    """Per frame, in the run's frame order: in-focus droplets, those in the band (lo <= d < hi) and their
    sum of d^3 (um^3), or None without the per-frame diameters."""
    if not rec or rec.get("dia") is None or not rec.get("frames"):
        return None
    lo, hi = band
    n, nb, s3 = [], [], []
    for f in rec["frames"]:
        d = np.asarray(rec["dia"].get(f, []), float)
        n.append(d.size)
        nb.append(int(((d >= lo) & (d < hi)).sum()))
        s3.append(float((d ** 3).sum()))
    return np.array(n, float), np.array(nb, float), np.array(s3)


def mass_balance_rate(ml_min, atom_vol_frac, count, s3_um3):
    """Droplets produced per second: (mL/min -> mm^3/s) x atomised volume fraction / (pi/6 <d^3>), with
    <d^3> = s3_um3 / count. Works on arrays (bootstrap replicates) as on scalars; 0 droplets -> nan."""
    q = ml_min * 1000.0 / 60.0
    with np.errstate(all="ignore"):
        vbar_mm3 = np.pi / 6.0 * np.asarray(s3_um3, float) / np.asarray(count, float) * 1e-9
        return q * np.asarray(atom_vol_frac, float) / np.where(vbar_mm3 > 0, vbar_mm3, np.nan)


def _vol_per_frame(rec: dict | None):
    """(droplet volume, total liquid volume) per frame in the run's frame order, or None when the
    liquid stage did not write volumes (before 2.3.0)."""
    if not rec or not rec.get("cpf"):
        return None
    try:
        d = np.array([float(rec["cpf"][f]["droplet_vol_mm3"]) for f in rec["frames"]])
        t = np.array([float(rec["cpf"][f]["total_liquid_vol_mm3"]) for f in rec["frames"]])
    except (KeyError, TypeError, ValueError):
        return None
    return d, t


# ---- small helpers ------------------------------------------------------------------------------

def _read_csv(path: Path) -> list[dict]:
    with open(path, newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def _num(x):
    try:
        v = float(x)
        return v if math.isfinite(v) else None
    except (TypeError, ValueError):
        return None


def _jsonable(x):
    if isinstance(x, dict):
        return {str(k): _jsonable(v) for k, v in x.items()}
    if isinstance(x, (list, tuple)):
        return [_jsonable(v) for v in x]
    if isinstance(x, np.ndarray):
        return _jsonable(x.tolist())
    if isinstance(x, (np.floating, float)):
        return None if not math.isfinite(float(x)) else float(x)
    if isinstance(x, (np.integer,)):
        return int(x)
    return x


def volume_percentiles(d, ps=(10, 50, 90)) -> list[float | None]:
    """Dv_p: the diameter below which p% of the droplet VOLUME (sum d^3) lies.

    Grouped-data interpolation: each distinct diameter is a cell reaching halfway to its
    neighbours, and the cumulative volume is interpolated linearly across the cells. The
    in-focus diameters sit on the pixel lattice (~1 um apart near 60 um, thousands of droplets
    per value), so interpolating between the sorted droplets themselves always lands ON a
    lattice value: on the 10/05 9x3 that gave 8 distinct Dv50s in 27 runs and identical
    replicates (found 2026-10-11). Across cells the estimate moves continuously with the
    distribution."""
    d = np.asarray(d, dtype=float)
    d = d[np.isfinite(d)]
    if d.size == 0:
        return [None for _ in ps]
    u, inv = np.unique(d, return_inverse=True)
    w = np.bincount(inv, weights=d ** 3)
    if u.size == 1:
        return [float(u[0]) for _ in ps]
    edges = np.concatenate([[u[0] - (u[1] - u[0]) / 2], (u[:-1] + u[1:]) / 2, [u[-1] + (u[-1] - u[-2]) / 2]])
    cum = np.concatenate([[0.0], np.cumsum(w)]) / w.sum()
    return [float(np.interp(p / 100.0, cum, edges)) for p in ps]


def capture_time(run_dir: Path) -> str | None:
    """ISO time of capture from <...>/YYYY/MM/DD/<HHMMSS>_... ; None if the path doesn't say."""
    run_dir = Path(run_dir)
    m = re.match(r"(\d{2})(\d{2})(\d{2})_", run_dir.name)
    if not m:
        return None
    y, mo, d = (p.name for p in (run_dir.parents[2], run_dir.parents[1], run_dir.parent))
    try:
        hh, mm, ss = (int(g) for g in m.groups())
        return datetime(int(y), int(mo), int(d), hh, mm, ss).isoformat()
    except ValueError:
        return None


# ---- the parts ----------------------------------------------------------------------------------

def size_part(dia_by_frame: dict, edges: list[float]) -> dict:
    d = np.concatenate([np.asarray(v, float) for v in dia_by_frame.values()]) \
        if dia_by_frame else np.empty(0)
    d = d[np.isfinite(d)]
    if d.size == 0:
        return {"n": 0}
    v = d ** 3
    dv10, dv50, dv90 = volume_percentiles(d)
    nb = len(edges) - 1
    idx = np.clip(np.searchsorted(edges, d, side="right") - 1, 0, nb - 1)
    counts = np.bincount(idx, minlength=nb)
    vol = np.bincount(idx, weights=v, minlength=nb)
    return {
        "n": int(d.size),
        "d10": float(d.mean()), "d32": float(v.sum() / (d ** 2).sum()),
        "dv10": dv10, "dv50": dv50, "dv90": dv90,
        "span": (dv90 - dv10) / dv50 if None not in (dv10, dv50, dv90) and dv50 else None,
        "vol_pct_below_small": float(100 * v[d < SMALL_UM].sum() / v.sum()),
        "vol_pct_above_large": float(100 * v[d > LARGE_UM].sum() / v.sum()),
        "edges": [e if math.isfinite(e) else None for e in edges],
        "pct_count": (100 * counts / d.size).tolist(),
        "pct_volume": (100 * vol / vol.sum()).tolist(),
    }


def steadiness_part(frames: list[str], counts: dict, liquid: dict, fs: float | None) -> dict:
    """`counts` = all droplets per frame (in + out of focus); `liquid` = total_liquid_px."""
    c = np.array([counts.get(f, 0) for f in frames], float)
    out: dict = {"frames": len(frames), "fs_hz": fs, "count_series": c.tolist()}
    if c.size and c.mean() > 0:
        out["count_mean"] = float(c.mean())
        out["count_cv"] = float(c.std() / c.mean())
        out["thin_pct"] = float(100 * (c < 0.5 * c.mean()).mean())
        c0 = c - c.mean()
        lags = {}
        for L in (1, 2, 4, 8, 16):
            if c.size > L + 2 and c0[:-L].std() > 0 and c0[L:].std() > 0:
                lags[str(L)] = float(np.corrcoef(c0[:-L], c0[L:])[0, 1])
        out["lag_r"] = lags
        out["lag1"] = lags.get("1")
        if fs and c.size >= 32:
            from scipy import signal
            nper = int(min(256, c.size))          # the 2026-10-07 pulsing analysis: 256 at 508 frames
            f, P = signal.welch(c0, fs=fs, nperseg=nper, noverlap=nper // 2, detrend="linear")
            b = f > 0.1
            if P[b].sum() > 0:
                lo = b & (f < SLOW_HZ)
                out["slow_share_pct"] = float(100 * P[lo].sum() / P[b].sum())
                out["slow_peak_hz"] = float(f[lo][np.argmax(P[lo])]) if lo.any() else None
                out["slow_share_noise_pct"] = float(100 * (SLOW_HZ - 0.1) / (fs / 2 - 0.1))
            out["spectrum_hz"] = f.tolist()
            out["spectrum"] = P.tolist()
    li = np.array([liquid[f] for f in frames if f in liquid], float)
    if li.size and li.mean() > 0:
        out["liquid_series"] = li.tolist()
        out["liquid_cv"] = float(li.std() / li.mean())
        out["liquid_max_px"] = float(li.max())
    return out


def pressure_part(xlsx: Path, frames: list[str], counts: dict, fs: float | None) -> dict | None:
    import openpyxl
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    if "Pressure" not in wb.sheetnames:
        return None
    rows = [r for r in wb["Pressure"].iter_rows(min_row=2, values_only=True) if r and r[0] is not None]
    if not rows:
        return None
    t = np.array([_num(r[0]) for r in rows], float)
    p = np.array([_num(r[1]) if len(r) > 1 else None for r in rows], float)
    q = np.array([(_num(r[2]) or 0.0) if len(r) > 2 else 0.0 for r in rows], float)
    t0 = next((_num(r[3]) for r in rows if len(r) > 3 and _num(r[3]) is not None), None)
    t1 = next((_num(r[4]) for r in rows if len(r) > 4 and _num(r[4]) is not None), None)
    out = {"peak_bar_log": float(np.nanmax(p)) if np.isfinite(p).any() else None,
           "camera_window": [t0, t1]}
    if t0 is None or t1 is None:
        out["note"] = "no camera window recorded in the pressure log"
        return out
    _, iu = np.unique(t, return_index=True)
    t, p, q = t[iu], p[iu], q[iu]
    m = (t >= t0) & (t <= t1) & np.isfinite(p)
    if m.sum() < 3:
        out["note"] = "too few pressure samples inside the camera window"
        return out
    tp, pp, qq = t[m] - t0, p[m], q[m]
    out.update({"t": tp.tolist(), "pressure": pp.tolist(), "flow": qq.tolist(),
                "pressure_mean": float(pp.mean()), "pressure_cv_pct": float(100 * pp.std() / pp.mean()),
                "gas_flow_cv_pct": float(100 * qq.std() / qq.mean()) if qq.mean() > 0 else None})
    # droplets per frame binned onto the pressure samples (frame times assume the cine
    # starts at cam_start, as in the 2026-10-07 pulsing analysis)
    if fs and frames:
        n = np.array([int(re.search(r"_n(-?\d+)$", f).group(1)) for f in frames
                      if re.search(r"_n(-?\d+)$", f)], float)
        if n.size == len(frames):
            # n counts CINE frames; fs is analysed frames per second, so the cine's own rate is
            # fs x (cine frames between analysed frames)
            spacing = float(np.median(np.diff(n))) if n.size > 1 else 1.0
            tf = (n - n.min()) / (fs * spacing)
            c = np.array([counts.get(f, 0) for f in frames], float)
            edges = tp[:-1] + np.diff(tp) / 2
            idx = np.digitize(tf, edges)
            binned = np.array([c[idx == i].mean() if (idx == i).any() else np.nan
                               for i in range(len(tp))])
            ok = np.isfinite(binned)
            if ok.sum() >= 5 and binned[ok].std() > 0 and pp[ok].std() > 0:
                out["r_pressure_droplets"] = float(np.corrcoef(pp[ok], binned[ok])[0, 1])
    return out


def mm2_per_px(path: Path) -> float:
    """The area of one pixel for the run `path` belongs to, from its optics profile (config/optics/)."""
    from . import pipeline_spec as spec
    return spec.rdc_module("optics").mm2_per_px(path)


def pieces_part(components_csv: Path, n_frames: int) -> dict:
    rows = _read_csv(components_csv)
    a = np.array([_num(r.get("area_px")) for r in rows], float)
    a = a[np.isfinite(a)]
    if a.size == 0:
        return {"n": 0, "per_frame": 0.0}
    mm2 = mm2_per_px(components_csv)
    return {"n": int(a.size), "per_frame": float(a.size / max(n_frames, 1)),
            "median_px": float(np.median(a)), "largest_px": float(a.max()),
            "median_mm2": float(np.median(a) * mm2), "largest_mm2": float(a.max() * mm2),
            "p90_px": float(np.percentile(a, 90)),
            "log10_area_hist": np.histogram(np.log10(a), bins=np.arange(1, 7.01, 0.25))[0].tolist()}


def background_part(bg_path: Path) -> dict | None:
    import cv2
    bg = cv2.imread(str(bg_path), cv2.IMREAD_UNCHANGED)
    if bg is None:
        return None
    bg = bg.astype(np.float32)
    ratio = bg / np.maximum(cv2.GaussianBlur(bg, (0, 0), 25), 1)
    return {"bg_median": float(np.median(bg)), "dark_speck_pct": float(100 * (ratio < 0.75).mean())}


def detections_part(pred_path: Path, thr: float, n_frames: int) -> dict:
    preds = json.loads(Path(pred_path).read_text(encoding="utf-8"))
    kept = [d for d in preds if d.get("score", 0) >= thr]
    per_cls: dict[int, int] = {}
    spots: dict[tuple, set] = {}
    frames = set()
    for d in kept:
        per_cls[d["category_id"]] = per_cls.get(d["category_id"], 0) + 1
        frames.add(d["image_id"])
        if d["category_id"] == 1:
            x, y, w, h = d["bbox"]
            k = (round((x + w / 2) / STATIC_GRID_PX), round((y + h / 2) / STATIC_GRID_PX))
            spots.setdefault(k, set()).add(d["image_id"])
    nf = max(n_frames, len(frames), 1)
    static = sorted(((len(v), k) for k, v in spots.items() if len(v) >= STATIC_FRAC * nf), reverse=True)
    return {"detections_per_frame": len(kept) / nf,
            "per_class_per_frame": {str(k): v / nf for k, v in sorted(per_cls.items())},
            "static_spots": len(static),
            "static_spots_xy": [[k[0] * STATIC_GRID_PX, k[1] * STATIC_GRID_PX, n / nf]
                                for n, k in static[:10]]}


# ---- assembling one run --------------------------------------------------------------------------

def _inputs(run_dir: Path, meas_dir: Path, cl_dir: Path) -> dict[str, Path]:
    raw = run_dir / "shadowgraph" / "raw"
    return {"droplet_sizes": meas_dir / "droplet_sizes.csv", "summary": meas_dir / "summary.json",
            "per_frame_cl": cl_dir / "classical_per_frame.csv",
            "components": cl_dir / "classical_components.csv",
            "run_summary": run_dir / "run_summary.xlsx",
            "extraction": raw / "extraction_metadata.json",
            "background": raw / "background_median.tiff",
            "predictions": run_dir / "shadowgraph" / "analysis" / "predictions.json"}


def _signature(paths: dict[str, Path], extra: dict) -> str:
    h = hashlib.sha256(json.dumps(extra, sort_keys=True).encode())
    for k in sorted(paths):
        p = paths[k]
        st = p.stat() if p.exists() else None
        h.update(f"{k}:{st.st_size if st else -1}:{st.st_mtime_ns if st else -1}".encode())
    return h.hexdigest()[:20]


def compute(run_dir: Path, thr: float, *, rec: dict | None = None, context=None,
            meas_dir: Path | None = None, cl_dir: Path | None = None,
            bin_width: float = 25.0, bin_max: float = 200.0, detections: bool = True) -> dict:
    """All parts for one run. `rec` is taguchi_analysis.load_run's record when the caller
    already has it (its per-frame in-focus diameters and frame list are reused);
    `context` is covariates.RunContext (liquid mL/min, peak pressure)."""
    from . import size_bins
    run_dir = Path(run_dir)
    an = run_dir / "shadowgraph" / "analysis"
    meas_dir = Path(meas_dir or (rec or {}).get("meas_dir") or an / f"droplets_{thr:.2f}")
    cl_dir = Path(cl_dir or (rec or {}).get("cl_dir") or an / f"liquid_{thr:.2f}")
    paths = _inputs(run_dir, meas_dir, cl_dir)
    out: dict = {"schema": SCHEMA, "run": run_dir.name, "path": str(run_dir),
                 "captured": capture_time(run_dir), "missing": {}}

    def part(name, fn):
        try:
            v = fn()
            if v is None:
                out["missing"].setdefault(name, "input not found")
            return v
        except FileNotFoundError as e:
            out["missing"][name] = f"missing: {Path(e.filename).name if e.filename else e}"
        except Exception as e:                             # noqa: BLE001 -- recorded, never fatal
            out["missing"][name] = f"{type(e).__name__}: {e}"
        return None

    # per-frame inputs
    if rec is not None:
        frames = list(rec["frames"])
        dia = {f: np.asarray(rec["dia"].get(f, []), float) for f in frames}
        liquid = {f: _num(rec["cpf"][f].get("total_liquid_px")) or 0.0 for f in frames}
    else:
        frames, dia, liquid = [], {}, {}
        if paths["per_frame_cl"].exists():
            for r in _read_csv(paths["per_frame_cl"]):
                frames.append(r["frame"])
                liquid[r["frame"]] = _num(r.get("total_liquid_px")) or 0.0
        if paths["droplet_sizes"].exists():
            for r in _read_csv(paths["droplet_sizes"]):
                if r.get("in_focus") == "1":
                    dia.setdefault(r["frame"], []).append(float(r["diameter_um"]))
        frames = sorted(frames)
    counts: dict[str, int] = {}
    if paths["droplet_sizes"].exists():
        for r in _read_csv(paths["droplet_sizes"]):
            counts[r["frame"]] = counts.get(r["frame"], 0) + 1
    fs = None
    try:
        meta = json.loads(paths["extraction"].read_text(encoding="utf-8"))
        if meta.get("frame_rate_fps") and meta.get("stride"):
            fs = float(meta["frame_rate_fps"]) / float(meta["stride"])
    except (OSError, ValueError):
        pass
    out["fs_hz"] = fs

    edges = size_bins.make_edges(bin_width, bin_max)
    out["size"] = part("size", lambda: size_part(dia, edges) if dia else None)
    out["steadiness"] = part("steadiness", lambda: steadiness_part(frames, counts, liquid, fs)
                             if frames else None)
    out["pressure"] = part("pressure", lambda: pressure_part(paths["run_summary"], frames, counts, fs))
    out["pieces"] = part("pieces", lambda: pieces_part(paths["components"], len(frames)))

    q: dict = {}
    try:
        s = json.loads(paths["summary"].read_text(encoding="utf-8"))
        prov = s.get("provenance", {})
        n_in, n_oof = s.get("droplets_in_focus"), s.get("droplets_out_of_focus")
        moved = prov.get("sharpness_reclassified")
        q.update({"sizer_version": prov.get("sizer_version"), "in_focus": n_in, "out_of_focus": n_oof,
                  "moved_by_rule": moved,
                  "moved_pct_of_in_focus": (100 * moved / (n_in + moved)
                                            if moved is not None and n_in is not None and (n_in + moved)
                                            else None),
                  "method_counts": prov.get("diameter_method_counts")})
    except (OSError, ValueError):
        out["missing"]["quality.summary"] = "summary.json not readable"
    bgq = part("quality.background", lambda: background_part(paths["background"])
               if paths["background"].exists() else None)
    if bgq:
        q.update(bgq)
    if detections:
        dq = part("quality.detections", lambda: detections_part(paths["predictions"], thr, len(frames))
                  if paths["predictions"].exists() else None)
        if dq:
            q.update(dq)
    out["quality"] = q or None

    atom = (rec or {}).get("atom")
    ml = getattr(context, "liquid_ml_min", None)
    out["throughput"] = {"liquid_ml_min": ml, "atomised_pct": atom,
                         "af_flow_ml_min": ml * atom / 100 if ml is not None and atom is not None else None}
    peak_ctx = getattr(context, "peak_pressure_bar", None)
    out["peak_bar"] = peak_ctx if peak_ctx is not None else (out.get("pressure") or {}).get("peak_bar_log")
    return _jsonable(out)


def load_or_compute(run_dir: Path, thr: float, *, cache_dir: Path | None = None, **kw) -> dict:
    """compute(), cached as <cache_dir>/<run>.json and recomputed whenever an input changes."""
    if cache_dir is None:
        return compute(run_dir, thr, **kw)
    run_dir = Path(run_dir)
    rec = kw.get("rec") or {}
    an = run_dir / "shadowgraph" / "analysis"
    paths = _inputs(run_dir, Path(rec.get("meas_dir") or an / f"droplets_{thr:.2f}"),
                    Path(rec.get("cl_dir") or an / f"liquid_{thr:.2f}"))
    ctx = kw.get("context")
    extra = {"schema": SCHEMA, "thr": thr, "bins": [kw.get("bin_width", 25.0), kw.get("bin_max", 200.0)],
             "det": kw.get("detections", True),
             "ctx": [getattr(ctx, "liquid_ml_min", None), getattr(ctx, "peak_pressure_bar", None)],
             "atom": rec.get("atom")}
    sig = _signature(paths, extra)
    cache = Path(cache_dir) / f"{run_dir.name}.json"
    try:
        hit = json.loads(cache.read_text(encoding="utf-8"))
        if hit.get("_signature") == sig:
            return hit
    except (OSError, ValueError):
        pass
    m = compute(run_dir, thr, **kw)
    m["_signature"] = sig
    cache.parent.mkdir(parents=True, exist_ok=True)
    tmp = cache.with_name(f".{cache.name}.tmp")
    tmp.write_text(json.dumps(m), encoding="utf-8")
    fsops.replace(tmp, cache)              # retries Windows sharing violations, like every writer
    return m


# ---- the analysed values ------------------------------------------------------------------------

def response_values(m: dict, rec: dict | None = None, n_boot: int = 0, seed: int = 0,
                    band=DEFAULT_BAND_UM) -> dict:
    """{key: value} for EXTRA_RESPONSES, plus `boot_<key>` replicate arrays (centred on the
    point estimate, like taguchi_analysis.bootstrap) where a frame bootstrap is meaningful:
    Dv50 and frames-below-half by resampling frames at the run's ci_stride; atomised liquid
    flow from the run's own atomised-fraction bootstrap. Lag-1 and peak pressure carry no CI."""
    st = m.get("steadiness") or {}
    vals = {"dv50": (m.get("size") or {}).get("dv50"),
            "af_flow": (m.get("throughput") or {}).get("af_flow_ml_min"),
            "thin_pct": st.get("thin_pct"),
            "lag1": st.get("lag1"),
            "peak_bar": m.get("peak_bar"),
            "atom_vol": None, "af_flow_vol": None, "band_per_frame": None, "rate": None, "rate_band": None}
    bc = band_counts(rec, band or DEFAULT_BAND_UM)
    if bc is not None and bc.size:
        vals["band_per_frame"] = float(bc.mean())
    vols = _vol_per_frame(rec)
    if vols is not None and vols[1].sum() > 0:
        vals["atom_vol"] = float(100.0 * vols[0].sum() / vols[1].sum())
        ml = (m.get("throughput") or {}).get("liquid_ml_min")
        vals["af_flow_vol"] = ml * vals["atom_vol"] / 100 if ml is not None else None
    rp = _rate_parts(rec, band or DEFAULT_BAND_UM)
    ml = (m.get("throughput") or {}).get("liquid_ml_min")
    if vals["atom_vol"] is not None and ml is not None and rp is not None and rp[0].sum() > 0:
        n_f, nb_f, s3_f = rp
        r = float(mass_balance_rate(ml, vals["atom_vol"] / 100, n_f.sum(), s3_f.sum()))
        if np.isfinite(r):
            vals["rate"] = r
            vals["rate_band"] = r * float(nb_f.sum() / n_f.sum())
    out = dict(vals)
    if not (n_boot and rec):
        return out
    # numpy on macOS (Accelerate BLAS) raises spurious divide/overflow warnings from matmul on
    # perfectly finite inputs; replicates are checked finite where they are used
    with np.errstate(all="ignore"):
        rng = np.random.default_rng(seed)
        frames = list(rec["frames"])
        pool = np.arange(len(frames))[::max(1, int(rec.get("ci_stride") or 1))]
        n = pool.size
        if n >= 2:
            W = np.zeros((n_boot, len(frames)))
            draws = pool[rng.integers(0, n, (n_boot, n))]
            for b in range(n_boot):
                np.add.at(W[b], draws[b], 1.0)
            # Dv50 from per-frame volume histograms on 1 um bins (exact enough: < 0.5 um)
            if vals["dv50"] is not None:
                d_all = [np.asarray(rec["dia"].get(f, []), float) for f in frames]
                top = max((d.max() for d in d_all if d.size), default=0.0)
                edges = np.arange(0.0, top + 2.0, 1.0)
                H = np.vstack([np.histogram(d, bins=edges, weights=d ** 3)[0] if d.size
                               else np.zeros(edges.size - 1) for d in d_all])
                V = W @ H
                cum = np.cumsum(V, axis=1)
                tot = cum[:, -1:]
                ok = tot[:, 0] > 0
                frac = np.where(tot > 0, cum / np.where(tot > 0, tot, 1), 0)
                mids = edges[1:]
                reps = np.array([np.interp(0.5, frac[i], mids) if ok[i] else np.nan for i in range(n_boot)])
                base = np.interp(0.5, np.cumsum(H[pool].sum(0)) / max(H[pool].sum(), 1e-12), mids)
                out["boot_dv50"] = (reps - base + vals["dv50"]).tolist()
            counts = np.array(st.get("count_series") or [], float)
            if vals["thin_pct"] is not None and counts.size == len(frames):
                thin = (counts < 0.5 * counts.mean()).astype(float)
                reps = 100 * (W @ thin) / W.sum(axis=1)
                base = 100 * thin[pool].mean()
                out["boot_thin_pct"] = (reps - base + vals["thin_pct"]).tolist()
        if vals["af_flow"] is not None and rec.get("boot_atom") is not None and rec.get("atom"):
            out["boot_af_flow"] = (np.asarray(rec["boot_atom"]) * vals["af_flow"] / rec["atom"]).tolist()
        if bc is not None and bc.size == len(frames) and n >= 2:
            reps = (W @ bc) / W.sum(axis=1)
            out["boot_band_per_frame"] = (reps - bc[pool].mean() + vals["band_per_frame"]).tolist()
        if vals["atom_vol"] is not None and n >= 2:
            dv, tv = vols
            reps = 100.0 * (W @ dv) / np.where(W @ tv > 0, W @ tv, np.nan)
            base = 100.0 * dv[pool].sum() / tv[pool].sum() if tv[pool].sum() > 0 else vals["atom_vol"]
            out["boot_atom_vol"] = (reps - base + vals["atom_vol"]).tolist()
            if vals["af_flow_vol"] is not None and vals["atom_vol"]:
                out["boot_af_flow_vol"] = (np.asarray(out["boot_atom_vol"]) * vals["af_flow_vol"]
                                           / vals["atom_vol"]).tolist()
            if vals["rate"] is not None and rp is not None:
                n_f, nb_f, s3_f = rp
                frac = (W @ dv) / np.where(W @ tv > 0, W @ tv, np.nan)
                reps = mass_balance_rate(ml, frac, W @ n_f, W @ s3_f)
                base = float(mass_balance_rate(ml, dv[pool].sum() / tv[pool].sum(), n_f[pool].sum(),
                                               s3_f[pool].sum())) if tv[pool].sum() > 0 else vals["rate"]
                out["boot_rate"] = (reps - base + vals["rate"]).tolist()
                share = (W @ nb_f) / np.where(W @ n_f > 0, W @ n_f, np.nan)
                share0 = nb_f[pool].sum() / n_f[pool].sum() if n_f[pool].sum() > 0 else np.nan
                out["boot_rate_band"] = (reps * share - base * share0 + vals["rate_band"]).tolist()
    return out
