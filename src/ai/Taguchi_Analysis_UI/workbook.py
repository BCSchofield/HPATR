"""workbook.py -- taguchi_analysis.xlsx: one sheet per table, drawn from tables.py.

pandas + openpyxl (xlsxwriter is not installed). Written to a temp file and moved into
place, so a half-written workbook never replaces a good one. On Windows a workbook that is
open in Excel cannot be replaced; that is reported as a plain instruction, not a traceback.
"""
from __future__ import annotations

import csv
import os
from pathlib import Path

import pandas as pd

from . import fsops, tables

SHEETS = (                               # (sheet name <= 31 chars, table key, always written)
    ("Per-run responses", "per_run", True),
    ("Design matrix", "design_matrix", True),
    ("Main effects", "main_effects", False),
    ("ANOVA", "anova", False),
    ("S-N ratios", "sn_ratios", False),
    ("S-N by condition", "sn_by_condition", False),
    ("Size bins (count)", "size_bins_count", False),
    ("Size bins (volume)", "size_bins_volume", False),
    ("Size bins count by condition", "size_bins_count_by_condition", False),
    ("Size bins volume by condition", "size_bins_volume_by_condition", False),
    ("GLR context", "glr_context", False),
    ("GLR trends", "glr_trends", False),
    ("Silicone film at exit", "film_at_exit", False),
    ("Replicate outliers", "replicate_outliers", False),
    ("Condition summary", "condition_summary", False),
    ("Size percentiles", "size_percentiles", False),
    ("Size percentiles by condition", "size_percentiles_by_condition", False),
    ("Steadiness", "steadiness", False),
    ("Drift", "drift", False),
    ("Residuals", "residuals", False),
    ("Response correlations", "response_correlations", False),
    ("Liquid pieces", "liquid_pieces", False),
    ("Data quality", "data_quality", False),
    ("Skipped runs", "skipped_runs", False),
)
P_COLUMNS = {"p", "q (BH)"}
MAX_WIDTH = 60


class WorkbookError(RuntimeError):
    pass


def read_timings(output_dir: Path) -> list[dict]:
    """The batch's own timings.csv (batch_runs schema), unchanged, for the Timings sheet."""
    p = Path(output_dir) / "_job" / "timings.csv"
    try:
        with open(p, newline="", encoding="utf-8") as f:
            return list(csv.DictReader(f))
    except OSError:
        return []


def _style(ws, frame: pd.DataFrame) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    head_fill = PatternFill("solid", fgColor="E5E5EA")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = head_fill
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"
    for i, col in enumerate(frame.columns, start=1):
        letter = ws.cell(row=1, column=i).column_letter
        longest = max([len(str(col))] + [len(f"{v:.4g}" if isinstance(v, float) else str(v))
                                         for v in frame[col].head(200)])
        ws.column_dimensions[letter].width = min(MAX_WIDTH, max(8, longest + 2))
        if col in P_COLUMNS:
            for row in range(2, len(frame) + 2):
                ws.cell(row=row, column=i).number_format = "0.0000"
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, str) and cell.value[:1] in "=+-@":
                cell.data_type = "s"                   # text that looks like a formula stays text


FILM_SHEET = "Silicone film at exit"
DIAGRAM = "fig_orifice_model.png"


def annotate_film(ws, n_cols: int, figures_dir: Path | None) -> None:
    """To the right of the per-run table on the film sheet: the model diagram (when it was drawn) and, below
    it, the maths and assumptions in plain text (orifice.EXPLANATION)."""
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter
    from . import orifice
    col = n_cols + 2
    letter = get_column_letter(col)
    row = 1
    png = Path(figures_dir) / DIAGRAM if figures_dir else None
    if png is not None and png.is_file():
        from openpyxl.drawing.image import Image
        img = Image(str(png))
        scale = 820 / img.width
        img.width, img.height = int(img.width * scale), int(img.height * scale)
        ws.add_image(img, f"{letter}1")
        row = int(img.height / 20) + 3                  # default rows are 20 px
    ws.column_dimensions[letter].width = 120            # about the image's width; longer lines run on to the right
    for i, line in enumerate(orifice.EXPLANATION):
        cell = ws.cell(row=row + i, column=col, value=line)
        cell.alignment = Alignment(wrap_text=False, vertical="top")   # wrapped text in default-height rows
        if i == 0 or line.endswith(":"):                             # shows one line in Excel
            cell.font = Font(bold=True)


def write(path: Path, results, bins, prov, timings: list[dict] | None = None,
          figures_dir: Path | None = None) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    data = tables.all_tables(results, bins)
    sheets: list[tuple[str, pd.DataFrame]] = []
    for name, key, always in SHEETS:
        rows = data.get(key) or []
        if rows or always:
            sheets.append((name, pd.DataFrame(rows)))
    if timings:
        sheets.append(("Timings", pd.DataFrame(timings)))
    sheets.append(("Provenance", pd.DataFrame(prov.rows(), columns=["Item", "Value"])))

    tmp = path.with_name(f".{path.stem}.{os.getpid()}.tmp.xlsx")
    try:
        with pd.ExcelWriter(tmp, engine="openpyxl") as xw:
            for name, frame in sheets:
                frame.to_excel(xw, sheet_name=name[:31], index=False)
                _style(xw.sheets[name[:31]], frame)
                if name == FILM_SHEET:
                    annotate_film(xw.sheets[name], len(frame.columns), figures_dir)
        try:
            fsops.replace(tmp, path)            # retried for a few seconds, then: Excel has it open
        except PermissionError:
            raise WorkbookError(f"cannot replace {path.name}: close it in Excel and analyse again") from None
    finally:
        tmp.unlink(missing_ok=True)
    return path
